"""System 06 対検査レビューの指摘書（Word／Excel）。データ源は report.md と同じ（summary・gate・findings・decisions・trail・設定）。"""
from __future__ import annotations

import io
from datetime import datetime

from web.services.docverify import pair_cases as pc
from web.services.docverify import pair_review as pr

FONT = "游ゴシック"
COLUMNS = ("番号", "種別", "重大度", "状態", "要件の該当箇所", "設計の該当箇所", "区分", "採否者", "採否日時")
_ACTIONS = {"approved": "承認", "decided": "採否を記録", "revoked": "採否を取り消し"}


def _ref(ref: dict | None) -> str:
    if not ref:
        return "未確認"
    sec = ref.get("section") or "未確認"
    return f"§{sec}" + (f"（{ref['title']}）" if ref.get("title") else "")


def _quote(ref: dict | None) -> str:
    return ((ref or {}).get("quote") or (ref or {}).get("text") or "").strip() or "引用なし（未確認）"


def _who(user: str | None) -> str:
    return "未ログインの利用者" if user == pr.NONE_USER else (user or "未記録")


def _kind(f: dict) -> str:
    """区分: 両文書の引用があれば仕様、自動検査の検出は観測、それ以外は推論。"""
    if f.get("req_ref") and f.get("design_ref"):
        return "仕様"
    return "観測" if f.get("source") == "ai" else "推論"


def collect(iid: str) -> dict:
    from web.services.docverify import settings_store as ss
    d = pr._existing(iid)
    s, g = pc.summary(iid), pr.gate(iid)
    findings, trail = pr._findings(d), pr.trail(iid)
    last = {}
    for dec in pr._decisions(d):
        if dec.get("revoke"):
            last.pop(dec.get("finding_id"), None)
        else:
            last[dec["finding_id"]] = dec
    approvals = [t for t in trail if t.get("action") == "approved"]
    rows = []
    for i, f in enumerate(findings, 1):
        dec = last.get(f["id"]) or {}
        rows.append({"no": i, "id": f["id"], "title": f.get("title") or "", "type": f.get("type") or "",
                     "severity": f.get("severity") or "未設定", "status": f.get("status") or "未判定",
                     "req": _ref(f.get("req_ref")), "design": _ref(f.get("design_ref")), "kind": _kind(f),
                     "by": _who(dec.get("by")) if dec else "未判定", "at": dec.get("at") or "", "reason": dec.get("reason") or "",
                     "changed": bool(f.get("changed_in_current")), "f": f})
    return {"iid": pr.resolve_iid(iid), "s": s, "g": g, "rows": rows, "trail": trail,
            "approver": _who(approvals[-1]["user"]) if approvals else "未承認",
            "approved_at": approvals[-1]["at"] if approvals else "",
            "cfg": ss.effective(), "today": datetime.now().strftime("%Y-%m-%d")}


def _counts(scope: dict) -> list[tuple[str, list[str]]]:
    missing = [m.get("req_id", "") for m in scope["missing_links"]]
    return [("未解消の重大", scope["unresolved_critical"]), ("要確認", scope["needs_review"]),
            ("欠落（設計に対応の無い要件）", missing), ("未検証", scope["unverified"])]


def _names(ids: list[str]) -> str:
    return "、".join(ids) if ids else "なし"


def _row_values(r: dict) -> list:
    return [r["id"], r["type"], r["severity"], r["status"] + ("（今回の変更分）" if r["changed"] else ""),
            r["req"], r["design"], r["kind"], r["by"], r["at"]]


def _criteria(c: dict) -> list[tuple[str, str]]:
    labels, mo = c["cfg"]["labels"], c["cfg"]["exit"]["max_open"]
    out = [(f"未対応の{labels.get(k, k)}の上限", "上限なし" if v is None else f"{v} 件") for k, v in mo.items()]
    rule = c["cfg"]["exit"]["changed_rule"]
    out.append(("今回の変更分の合格条件", "未対応 0 件" if rule == "zero" else "全体と同じ上限"))
    out.append(("検査する種別", "、".join(t for t, on in c["cfg"]["types"].items() if on) or "なし"))
    out.append(("固定の条件", "要確認・設計に対応の無い要件が 1 件でもあれば不合格。未判定は解消に数えない。"))
    return out


# ---------- Word ----------
def _font(doc) -> None:
    from docx.oxml.ns import qn
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[name]
        st.font.name = FONT
        rpr = st.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, header):
        cell.text = str(h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for r in rows:
        for cell, v in zip(t.add_row().cells, r):
            cell.text = str(v)
    return t


def _summary_section(doc, c) -> None:
    doc.add_heading("1. 判定の要約", 1)
    g = c["g"]
    assessment = c["s"].get("assessment") or {}
    doc.add_paragraph(f"検査範囲: 基準 {assessment.get('baseline_units', 0)} 要素・確認対象 {assessment.get('target_units', 0)} 要素・比較 {assessment.get('compared_pairs', 0)} 組")
    if g["overall"].get("ai_scope"):
        doc.add_paragraph(g["overall"]["ai_scope"])
    for label, scope, verdict in (("全体", g["overall"], c["s"]["verdict"]),
                                  ("今回の変更分", g["changed"], g["changed"].get("verdict", "判定不能"))):
        doc.add_heading(f"{label}: {verdict}", 2)
        _table(doc, ("項目", "件数", "該当の指摘"), [(n, f"{len(ids)} 件", _names(ids)) for n, ids in _counts(scope)])
    doc.add_heading("判定の根拠", 2)
    for line in g["reasons"]:
        doc.add_paragraph(line, style="List Bullet")


def _detail_section(doc, c) -> None:
    doc.add_heading("3. 指摘の詳細", 1)
    for r in c["rows"]:
        f = r["f"]
        doc.add_heading(f"{r['id']} {r['title']}".strip(), 2)
        doc.add_paragraph(f"種別: {r['type']} ／ 重大度: {r['severity']} ／ 状態: {r['status']} ／ 区分: {r['kind']}")
        doc.add_paragraph("本文の引用", style="Heading 3")
        doc.add_paragraph(f"{c['s']['baseline_name']} {r['req']}: {_quote(f.get('req_ref'))}", style="Intense Quote")
        doc.add_paragraph(f"{c['s']['target_name']} {r['design']}: {_quote(f.get('design_ref'))}", style="Intense Quote")
        doc.add_paragraph("説明", style="Heading 3")
        doc.add_paragraph(f.get("description") or "未確認")
        doc.add_paragraph("確認手順", style="Heading 3")
        for step in (f"{c['s']['baseline_name']} {r['req']} を開き、上の引用を読む。",
                     f"{c['s']['target_name']} {r['design']} を開き、同じ事柄の記述と比べる。",
                     f"「{r['type']}」に当たるかを判断し、対検査画面で採否を記録する。"):
            doc.add_paragraph(step, style="List Number")
        doc.add_paragraph("証跡", style="Heading 3")
        doc.add_paragraph(f"{r['status']}: {r['by']} が {r['at']} に記録。理由: {r['reason'] or '未記入'}"
                          if r["at"] else "採否: 未確認（まだ誰も判定していません）")


def to_docx(iid: str) -> bytes:
    from docx import Document
    c = collect(iid)
    s = c["s"]
    doc = Document()
    _font(doc)
    doc.add_paragraph("対検査レビュー 指摘書", style="Title")
    _table(doc, ("項目", "内容"), [
        ("レビュー名", s["name"]), ("対象文書と版", f"{s['baseline_name']} 第 {s['req_version']} 版 ⇔ {s['target_name']} 第 {s['design_version']} 版"),
        ("実施日", c["today"]), ("作成者", s["assignee"] or "未記録"),
        ("承認者", c["approver"] + (f"（{c['approved_at']}）" if c["approved_at"] else "")), ("判定", s["verdict"])])
    doc.add_page_break()
    _summary_section(doc, c)
    doc.add_heading("2. 指摘一覧", 1)
    _table(doc, COLUMNS, [_row_values(r) for r in c["rows"]])
    doc.add_paragraph("区分の読み方: 「仕様」は両文書の引用がある事実、「観測」は自動検査の検出、「推論」は指摘者の解釈です。")
    _detail_section(doc, c)
    doc.add_heading("4. 証跡（時系列）", 1)
    _table(doc, ("日時", "記録者", "操作", "対象", "理由"),
           [(t.get("at", ""), _who(t.get("user")), _ACTIONS.get(t.get("action"), t.get("action", "")),
             t.get("target", ""), t.get("reason", "")) for t in c["trail"]] or [("記録なし", "", "", "", "")])
    doc.add_heading("付録 判定基準（設定の値）", 1)
    _table(doc, ("基準", "値"), _criteria(c))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- Excel ----------
def _sheet(wb, title, header, rows, freeze=True):
    from openpyxl.styles import Border, Font, PatternFill, Side
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for r in rows:
        ws.append([("'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value) for value in r])
    thin = Side(style="thin")
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=cell.row == 1)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="E8EEF5")
    for i, h in enumerate(header):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(10, min(40, len(str(h)) * 3))
    if freeze:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    return ws


def to_xlsx(iid: str) -> bytes:
    from openpyxl import Workbook
    c = collect(iid)
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "指摘一覧", COLUMNS, [_row_values(r) for r in c["rows"]])
    _sheet(wb, "証跡", ("日時", "記録者", "操作", "対象", "理由"),
           [(t.get("at", ""), _who(t.get("user")), _ACTIONS.get(t.get("action"), t.get("action", "")),
             t.get("target", ""), t.get("reason", "")) for t in c["trail"]])
    s, g = c["s"], c["g"]
    rows = [("レビュー名", s["name"], ""), ("対象文書と版", f"{s['baseline_name']} 第 {s['req_version']} 版 ⇔ {s['target_name']} 第 {s['design_version']} 版", ""),
            ("判定（全体）", s["verdict"], ""), ("判定（今回の変更分）", g["changed"].get("verdict", "判定不能"), ""),
            ("作成者", s["assignee"] or "未記録", ""), ("承認者", c["approver"], c["approved_at"])]
    for label, scope in (("全体", g["overall"]), ("今回の変更分", g["changed"])):
        rows += [(f"{label}: {n}", f"{len(ids)} 件", _names(ids)) for n, ids in _counts(scope)]
    assessment = s.get("assessment") or {}
    rows += [("検査目的", s.get("purpose", ""), ""), ("比較実行数", assessment.get("compared_pairs", 0), "")]
    if g["overall"].get("ai_scope"):
        rows += [("AI 検査", g["overall"]["ai_scope"], "")]
    rows += [("判定の根拠", " ".join(g["reasons"]), "")] + [(k, v, "") for k, v in _criteria(c)]
    _sheet(wb, "判定", ("項目", "値", "該当"), rows, freeze=False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
