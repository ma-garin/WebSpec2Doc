"""影響度分析（System 08）。変更抽出・クエリ生成・影響度と優先度・API の受入基準。"""
from __future__ import annotations

import io

import pytest

import app as appmod
from web import config
from web.services.impact import analyze

OLD = "# 会員登録 画面仕様\n## 入力項目\n- 電話番号: 任意。半角数字 10〜13 桁。\n## 確認画面\n- 電話番号が未入力の場合は「（未入力）」と表示する。\n"
NEW = OLD.replace("任意。半角数字 10〜13 桁", "必須。半角数字 10〜11 桁")
DESIGN = ("# 会員管理 基本設計\n## 会員情報の更新\n- プロフィール更新時は、会員登録と同じ入力チェック（電話番号の桁数を含む）を行う。\n"
          "## 一括取込\n- 管理画面から CSV で会員を一括取込する。電話番号などの項目チェックは画面登録と共通の検証部品を用いる。\n"
          "## 退会\n- 退会時は会員情報を論理削除する。\n")
TESTS = "# 既存テストケース\n- TC-REG-031: 電話番号を空欄で登録できること。\n- TC-REG-040: 氏名 31 文字でエラーになること。\n"
DOCS = {"画面仕様書_会員登録_旧.md": OLD, "画面仕様書_会員登録_新.md": NEW,
        "基本設計書_会員管理.md": DESIGN, "テストケース_会員登録.md": TESTS}


def _passages(text: str) -> list[dict]:
    return [{"id": f"P{i}", "text": t, "location": f"{i}行"} for i, t in enumerate(text.strip().split("\n"), start=1)]


def _docs():
    roles = {"画面仕様書_会員登録_旧.md": ("old", "画面仕様書"), "画面仕様書_会員登録_新.md": ("new", "画面仕様書"),
             "基本設計書_会員管理.md": ("store", "設計書"), "テストケース_会員登録.md": ("store", "テストケース")}
    return [{"name": n, "role": r, "store": s} for n, (r, s) in roles.items()]


def _run(order="電話番号を必須にする"):
    return analyze.run(order, _docs(), lambda n: _passages(DOCS[n]))


def test_change_is_extracted_from_old_and_new_versions():
    """Given 変更前と変更後の仕様書 When 分析 Then 変わった 1 行だけが変更前→変更後つきで抜き出される。"""
    changes = _run()["changes"]
    assert len(changes) == 1
    c = changes[0]
    assert c["kind"] == "変更" and "任意" in c["before"] and "必須" in c["after"]
    assert c["location"].startswith("画面仕様書_会員登録_新.md")


def test_without_documents_the_order_becomes_the_change():
    """Given 文書なし・指示のみ When 分析 Then 指示の各文が変更になり、文書不足が notes で知らされる。"""
    result = analyze.run("電話番号を必須にする。桁数を11桁までにする。", [], lambda n: [])
    assert [c["kind"] for c in result["changes"]] == ["指示", "指示"]
    assert any("変更後の文書がありません" in n for n in result["notes"])
    assert all(r["level_name"] == "未確認" for r in result["impacts"])


def test_query_uses_subject_and_changed_terms_with_paraphrases():
    """Given 任意→必須・13桁→11桁 When クエリ生成 Then 主題「電話番号」と変わった語・言い換え（空欄・桁数）を持つ。"""
    q = _run()["queries"][0]
    assert q["subject"][0] == "電話番号"
    assert {"必須", "任意", "11桁", "13桁", "空欄", "桁数"} <= set(q["diff"])


def test_shared_component_with_harmful_operation_is_p1():
    """Given 共通の検証部品を使う一括取込 When 判定 Then 影響度=致命・優先度=P1 で、根拠の文書と引用が付く。"""
    top = _run()["impacts"][0]
    assert top["priority"] == 1 and top["level_name"] == "致命"
    assert top["doc"] == "基本設計書_会員管理.md" and "一括取込" in top["place"]
    assert "共通の検証部品" in top["quote"]


def test_existing_test_case_on_changed_rule_needs_expected_value_update():
    """Given 「空欄で登録できること」のテストケース When 判定 Then P2 で、期待値の書き換えをテストとして示す。"""
    row = next(r for r in _run()["impacts"] if "TC-REG-031" in r["quote"])
    assert row["priority"] == 2 and "期待値" in row["test"]


def test_unrelated_passages_are_not_reported():
    """Given 電話番号に触れない記述（退会・氏名） When 判定 Then 影響範囲に出ない。"""
    quotes = " ".join(r["quote"] for r in _run()["impacts"])
    assert "退会" not in quotes and "TC-REG-040" not in quotes


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    return appmod.app.test_client()


def test_api_end_to_end(client, monkeypatch):
    """作成 → 取り込み → 分析 → CSV → 削除 が API で通り、結果が保存される。"""
    from web.services.impact import llm
    monkeypatch.setattr(llm, "judge", lambda c, k: None)
    jid = client.post("/api/impact/jobs", json={"name": "電話番号"}).get_json()["job"]["id"]
    files = [(io.BytesIO(t.encode()), n) for n, t in DOCS.items()]
    r = client.post(f"/api/impact/jobs/{jid}/docs", data={"files": files}, content_type="multipart/form-data")
    assert r.status_code == 200
    roles = {d["name"]: d["role"] for d in r.get_json()["job"]["docs"]}
    assert roles["画面仕様書_会員登録_旧.md"] == "old" and roles["画面仕様書_会員登録_新.md"] == "new"
    r = client.post(f"/api/impact/jobs/{jid}/analyze", json={"order": "電話番号を必須にする"})
    job = r.get_json()["job"]
    assert r.status_code == 200 and job["analyzed_at"] and job["impacts"][0]["priority"] == 1
    assert client.get(f"/api/impact/jobs/{jid}").get_json()["job"]["impacts"] == job["impacts"]
    csv = client.get(f"/api/impact/jobs/{jid}/export.csv").get_data(as_text=True)
    assert csv.lstrip("﻿").startswith("ID,優先度,影響度") and "P1" in csv
    assert client.delete(f"/api/impact/jobs/{jid}").status_code == 200
    assert client.get(f"/api/impact/jobs/{jid}").status_code == 404


def test_api_errors_tell_the_next_action(client):
    """空の分析・取り込めない形式・存在しない分析は、次の行動を添えたエラーになる。"""
    jid = client.post("/api/impact/jobs", json={"name": "x"}).get_json()["job"]["id"]
    r = client.post(f"/api/impact/jobs/{jid}/analyze", json={"order": ""})
    assert r.status_code == 400 and "取り込んでください" in r.get_json()["error"]
    r = client.post(f"/api/impact/jobs/{jid}/docs", data={"files": [(io.BytesIO(b"x"), "a.exe")]},
                    content_type="multipart/form-data")
    assert r.status_code == 400 and "いずれかを選んでください" in r.get_json()["error"]
    assert client.get("/api/impact/jobs/000000000000").status_code == 404
    assert client.post("/api/impact/jobs", json={"name": ""}).status_code == 400


def test_llm_verdicts_replace_rule_and_drop_unrelated():
    """Given LLM が 1 件を無関係・1 件を致命と判定 When 分析 Then 無関係は消え、致命は P1・判定=AI になる。"""
    def fake(changes, candidates):
        out = []
        for c in candidates:
            if "TC-REG-031" in c["text"]:
                out.append({"key": c["key"], "related": False, "level": 1, "why": "", "test": ""})
            elif "確認画面" in c["place"] or "未入力" in c["text"]:
                out.append({"key": c["key"], "related": True, "level": 4, "why": "必須化で表示が変わる", "test": "確認画面を開く"})
        return out
    result = analyze.run("電話番号を必須にする", _docs(), lambda n: _passages(DOCS[n]), fake)
    assert result["engine"]["mode"] == "llm"
    assert not any("TC-REG-031" in r["quote"] for r in result["impacts"])
    row = next(r for r in result["impacts"] if "未入力" in r["quote"])
    assert row["method"] == "llm" and row["priority"] == 1 and row["why"] == "必須化で表示が変わる"


def test_llm_unavailable_falls_back_to_rule_and_says_so():
    """Given LLM に接続できない When 分析 Then 規則の判定のまま、理由が engine に残る。"""
    result = analyze.run("電話番号を必須にする", _docs(), lambda n: _passages(DOCS[n]), lambda c, k: None)
    assert result["engine"]["mode"] == "rule" and "LLM" in result["engine"]["reason"]
    assert result["impacts"][0]["priority"] == 1


def _job_with_docs(client):
    jid = client.post("/api/impact/jobs", json={"name": "電話番号"}).get_json()["job"]["id"]
    files = [(io.BytesIO(t.encode()), n) for n, t in DOCS.items()]
    client.post(f"/api/impact/jobs/{jid}/docs", data={"files": files}, content_type="multipart/form-data")
    return jid


def test_review_survives_reanalysis_and_reaches_export(client, monkeypatch):
    """Given 人が採用・P を直した When 再分析・書き出し Then 判断が残り、Excel に人の判断と直した P が出る。"""
    from openpyxl import load_workbook
    from web.services.impact import llm
    monkeypatch.setattr(llm, "judge", lambda c, k: None)
    jid = _job_with_docs(client)
    job = client.post(f"/api/impact/jobs/{jid}/analyze", json={"order": "電話番号を必須にする"}).get_json()["job"]
    iid = job["impacts"][1]["id"]
    r = client.patch(f"/api/impact/jobs/{jid}/impacts/{iid}", json={"status": "accepted", "priority": 1, "note": "要件側に確認済み"})
    assert r.status_code == 200
    job = client.post(f"/api/impact/jobs/{jid}/analyze", json={}).get_json()["job"]
    kept = [x for x in job["impacts"] if x.get("review")]
    assert len(kept) == 1 and kept[0]["review"]["note"] == "要件側に確認済み"
    wb = load_workbook(io.BytesIO(client.get(f"/api/impact/jobs/{jid}/export.xlsx").get_data()))
    rows = list(wb["影響範囲"].iter_rows(values_only=True))
    assert rows[0][:5] == ("ID", "優先度", "影響度", "判定", "人の判断")
    assert any(r[4] == "採用" and r[1] == "P1" and r[5] == "要件側に確認済み" for r in rows[1:])
    assert client.patch(f"/api/impact/jobs/{jid}/impacts/{iid}", json={"status": "maybe"}).status_code == 400


def test_library_docs_are_searched_by_every_analysis(client, monkeypatch):
    """Given 設計書を文書庫に入れた When 文書を持たない別の分析 Then 文書庫の設計書から影響が見つかる。"""
    from web.services.impact import llm
    monkeypatch.setattr(llm, "judge", lambda c, k: None)
    r = client.post("/api/impact/library", data={"files": [(io.BytesIO(DESIGN.encode()), "基本設計書_会員管理.md")]},
                    content_type="multipart/form-data")
    assert r.status_code == 200 and r.get_json()["library"]["docs"][0]["store"] == "設計書"
    jid = client.post("/api/impact/jobs", json={"name": "別"}).get_json()["job"]["id"]
    job = client.post(f"/api/impact/jobs/{jid}/analyze", json={"order": "電話番号の桁数を11桁までにする。"}).get_json()["job"]
    assert any(x["doc"] == "基本設計書_会員管理.md" for x in job["impacts"])
    assert client.delete("/api/impact/library/基本設計書_会員管理.md").get_json()["library"]["docs"] == []


def test_job_doc_can_move_to_library(client):
    """Given 分析の検索先の文書 When 文書庫へ Then 分析から消え、文書庫に入る。変更後の文書は入れられない。"""
    jid = _job_with_docs(client)
    r = client.post(f"/api/impact/jobs/{jid}/docs/基本設計書_会員管理.md/library")
    assert r.status_code == 200
    assert [d["name"] for d in r.get_json()["library"]["docs"]] == ["基本設計書_会員管理.md"]
    assert all(d["name"] != "基本設計書_会員管理.md" for d in r.get_json()["job"]["docs"])
    assert client.post(f"/api/impact/jobs/{jid}/docs/画面仕様書_会員登録_新.md/library").status_code == 400


def test_sample_creates_an_analyzed_job(client, monkeypatch):
    """見本で試す: 文書なしで 1 クリック → 分析済みの見本ができ、P1 がある。"""
    from web.services.impact import llm
    monkeypatch.setattr(llm, "judge", lambda c, k: None)
    r = client.post("/api/impact/sample")
    job = r.get_json()["job"]
    assert r.status_code == 201 and job["analyzed_at"] and len(job["docs"]) == 4
    assert job["impacts"][0]["priority"] == 1
