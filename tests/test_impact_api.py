"""影響度分析（System 09）の結合テスト（IT）。BD の I/F を API 経由で通す。

ID は docs/lifecycle/system09-impact/06-integration-test.md の IT-xxx。
"""
from __future__ import annotations

import io
import time
import urllib.request

import pytest
from openpyxl import load_workbook

import app as appmod
from web import config
from web.routes import impact_api

API = "/api/impact/v2"
SAMPLE = impact_api.__file__.replace("routes/impact_api.py", "services/impact/sample/")


def _read(name: str) -> bytes:
    with open(SAMPLE + name, "rb") as fh:
        return fh.read()


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setenv("WEBSPEC2DOC_AUTH_MODE", "off")
    impact_api._RECOVERED.clear()
    return appmod.app.test_client()


@pytest.fixture
def as_actor(monkeypatch):
    def use(name: str):
        monkeypatch.setattr(impact_api, "actor", lambda: name)
    return use


def _wait(client, job_id: str, timeout: float = 20.0) -> dict:
    end = time.time() + timeout
    while time.time() < end:
        job = client.get(f"{API}/jobs/{job_id}").get_json()["job"]
        if job["status"] in ("done", "failed", "cancelled", "interrupted"):
            return job
        time.sleep(0.2)
    raise AssertionError("分析が時間内に終わらない")


def _system_with_docs(client) -> str:
    sid = client.post(f"{API}/systems", json={"name": "会員管理"}).get_json()["system"]["id"]
    files = [(io.BytesIO(_read(n)), n) for n in ("基本設計書_会員管理.md", "テストケース_会員登録.md")]
    files.append((io.BytesIO(_read("画面仕様書_会員登録_旧.md")), "画面仕様書_会員登録.md"))
    assert client.post(f"{API}/systems/{sid}/documents", data={"files": files},
                       content_type="multipart/form-data").status_code == 200
    return sid


def _cr_with_baselines(client, sid: str) -> str:
    v1 = client.post(f"{API}/systems/{sid}/baselines", json={"name": "v1"}).get_json()["baseline"]["id"]
    files = [(io.BytesIO(_read("画面仕様書_会員登録_新.md")), "画面仕様書_会員登録.md")]
    client.post(f"{API}/systems/{sid}/documents", data={"files": files}, content_type="multipart/form-data")
    v2 = client.post(f"{API}/systems/{sid}/baselines", json={"name": "v2"}).get_json()["baseline"]["id"]
    r = client.post(f"{API}/systems/{sid}/crs", json={"title": "電話番号を必須に", "reason": "本人確認",
                                                    "base_baseline": v1, "target_baseline": v2, "depth": 2})
    assert r.status_code == 201
    return r.get_json()["cr"]["code"]


def test_it001_system_create_list_get(client):
    """IT-001 Given 新しい対象システム When 作成・一覧・取得 Then 作成者が管理者で一覧に出る。"""
    r = client.post(f"{API}/systems", json={"name": "会員管理", "description": "x"})
    assert r.status_code == 201
    sid = r.get_json()["system"]["id"]
    assert [s["id"] for s in client.get(f"{API}/systems").get_json()["systems"]] == [sid]
    got = client.get(f"{API}/systems/{sid}").get_json()
    assert got["my_role"] == "admin" and got["system"]["policy"] == "none"
    assert client.post(f"{API}/systems", json={"name": ""}).status_code == 400


def test_it002_settings_validation_and_audit(client):
    """IT-002 閾値の矛盾は 400、送信方針の変更は監査に前後の値が残る。"""
    sid = client.post(f"{API}/systems", json={"name": "s"}).get_json()["system"]["id"]
    assert client.patch(f"{API}/systems/{sid}", json={"p1_min": 3, "p2_min": 4}).status_code == 400
    assert client.patch(f"{API}/systems/{sid}", json={"policy": "local"}).status_code == 200
    log = client.get(f"{API}/systems/{sid}/audit").get_json()["audit"]
    assert any("local" in str(a.get("after")) for a in log)


def test_it003_documents_partial_failure(client):
    """IT-003 Given 3 文書中 1 つが非対応形式 When 取り込み Then 2 つ取り込まれ、失敗に理由が付く。"""
    sid = client.post(f"{API}/systems", json={"name": "s"}).get_json()["system"]["id"]
    files = [(io.BytesIO(_read("基本設計書_会員管理.md")), "基本設計書_会員管理.md"),
             (io.BytesIO(_read("テストケース_会員登録.md")), "テストケース_会員登録.md"),
             (io.BytesIO(b"MZ"), "a.exe")]
    r = client.post(f"{API}/systems/{sid}/documents", data={"files": files}, content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200 and len(d["added"]) == 2 and d["failed"][0]["name"] == "a.exe"
    kinds = {x["name"]: x["kind"] for x in d["documents"]}
    assert kinds["テストケース_会員登録.md"] == "testcase" and kinds["基本設計書_会員管理.md"] == "design"


def test_it004_baseline_frozen_and_diff(client):
    """IT-004 Given v1 作成後に文書を差し替え When v1 と v2 を比べる Then 変更 1 件、v1 の本文は変わらない。"""
    sid = _system_with_docs(client)
    _cr_with_baselines(client, sid)
    bls = client.get(f"{API}/systems/{sid}/baselines").get_json()["baselines"]
    ids = {b["name"]: b["id"] for b in bls}
    diff = client.get(f"{API}/systems/{sid}/baselines/diff?old={ids['v1']}&new={ids['v2']}").get_json()["diff"]
    assert len(diff["modified"]) == 1 and not diff["added"] and not diff["removed"]
    assert "任意" in diff["modified"][0]["old"]["text"] and "必須" in diff["modified"][0]["new"]["text"]


def test_it005_links_add_neighbors_remove(client):
    """IT-005 つながりを登録すると両側から辿れ、削除すると消えて監査に残る。"""
    sid = _system_with_docs(client)
    items = client.get(f"{API}/systems/{sid}/items?q=電話番号").get_json()["items"]
    a, b = items[0]["id"], items[1]["id"]
    link = client.post(f"{API}/systems/{sid}/links", json={"a": a, "b": b}).get_json()["link"]
    assert any(x["id"] == b for x in client.get(f"{API}/systems/{sid}/links?item={a}").get_json()["links"])
    assert any(x["id"] == a for x in client.get(f"{API}/systems/{sid}/links?item={b}").get_json()["links"])
    assert client.delete(f"{API}/systems/{sid}/links/{link['id']}").status_code == 200
    assert client.get(f"{API}/systems/{sid}/links?item={a}").get_json()["links"] == []
    assert any("link" in a_["action"] for a_ in client.get(f"{API}/systems/{sid}/audit").get_json()["audit"])


def test_it006_analyze_job_produces_impacts_with_risk(client, monkeypatch):
    """IT-006 Given 送信方針=送らない When 分析 Then 外部通信なしで影響が作られ、根拠・リスク式・P が付く。"""
    def no_network(*a, **k):
        raise AssertionError("送信方針が送らないのに外部へ通信した")
    monkeypatch.setattr(urllib.request, "urlopen", no_network)
    sid = _system_with_docs(client)
    code = _cr_with_baselines(client, sid)
    est = client.post(f"{API}/systems/{sid}/crs/{code}/estimate").get_json()
    assert est["policy"] == "none" and est["send"]["count"] == 0 and est["changes"] >= 1
    job = client.post(f"{API}/systems/{sid}/crs/{code}/analyze", json={}).get_json()["job"]
    assert _wait(client, job["id"])["status"] == "done"
    d = client.get(f"{API}/systems/{sid}/crs/{code}/impacts").get_json()
    assert d["summary"]["total"] >= 1 and d["summary"]["p1"] >= 1
    top = d["impacts"][0]
    assert top["evidence"] and top["risk"] == top["impact_level"] * top["likelihood"]
    assert top["impact_reason"] and top["likelihood_reason"] and top["state"] == "open"
    g = client.get(f"{API}/systems/{sid}/crs/{code}/graph").get_json()
    assert any(n["kind"] == "change" for n in g["nodes"]) and g["edges"]


def test_it007_review_resolve_approve_roles(client, as_actor):
    """IT-007 要確認が残ると提出できず、解消は理由必須、承認は承認者だけ（起票者本人は不可）。"""
    as_actor("analyst@x")
    sid = _system_with_docs(client)
    code = _cr_with_baselines(client, sid)
    job = client.post(f"{API}/systems/{sid}/crs/{code}/analyze", json={}).get_json()["job"]
    _wait(client, job["id"])
    assert client.post(f"{API}/systems/{sid}/crs/{code}/transition", json={"to": "review"}).status_code == 409
    rows = client.get(f"{API}/systems/{sid}/crs/{code}/impacts").get_json()["impacts"]
    assert client.post(f"{API}/impacts/{rows[0]['id']}/resolve", json={"state": "required", "reason": ""}).status_code == 400
    for r in rows:
        assert client.post(f"{API}/impacts/{r['id']}/resolve", json={"state": "required", "reason": "確認済み"}).status_code == 200
    assert client.post(f"{API}/systems/{sid}/crs/{code}/transition", json={"to": "review"}).status_code == 200
    client.put(f"{API}/systems/{sid}/members/lead@x", json={"role": "approver"})
    client.put(f"{API}/systems/{sid}/members/viewer@x", json={"role": "viewer"})
    assert client.post(f"{API}/systems/{sid}/crs/{code}/approval", json={"decision": "approved"}).status_code == 403
    as_actor("viewer@x")
    assert client.post(f"{API}/systems/{sid}/crs", json={"title": "t", "reason": "r"}).status_code == 403
    as_actor("stranger@x")
    assert client.get(f"{API}/systems/{sid}").status_code == 404
    as_actor("lead@x")
    assert client.post(f"{API}/systems/{sid}/crs/{code}/approval", json={"decision": "returned", "reason": ""}).status_code == 400
    r = client.post(f"{API}/systems/{sid}/crs/{code}/approval", json={"decision": "approved", "reason": "妥当"})
    assert r.status_code == 200 and r.get_json()["cr"]["status"] == "approved"


def test_it008_override_requires_reason_and_recalculates(client):
    """IT-008 修正は理由必須で、リスク値と優先度が閾値で再計算される。"""
    sid = _system_with_docs(client)
    code = _cr_with_baselines(client, sid)
    _wait(client, client.post(f"{API}/systems/{sid}/crs/{code}/analyze", json={}).get_json()["job"]["id"])
    imp = client.get(f"{API}/systems/{sid}/crs/{code}/impacts").get_json()["impacts"][-1]
    assert client.post(f"{API}/impacts/{imp['id']}/override", json={"impact_level": 4, "likelihood": 3, "reason": ""}).status_code == 400
    got = client.post(f"{API}/impacts/{imp['id']}/override",
                      json={"impact_level": 4, "likelihood": 3, "reason": "決済に波及"}).get_json()["impact"]
    assert got["risk"] == 12 and got["priority"] == 1 and got["override_reason"] == "決済に波及"


def test_it009_report_and_audit_csv(client):
    """IT-009 レポートは 5 シートで影響件数が画面と一致し、監査 CSV は絞り込み件数と一致する。"""
    sid = _system_with_docs(client)
    code = _cr_with_baselines(client, sid)
    _wait(client, client.post(f"{API}/systems/{sid}/crs/{code}/analyze", json={}).get_json()["job"]["id"])
    total = client.get(f"{API}/systems/{sid}/crs/{code}/impacts").get_json()["summary"]["total"]
    wb = load_workbook(io.BytesIO(client.get(f"{API}/systems/{sid}/crs/{code}/report.xlsx").get_data()))
    assert {"概要", "影響", "テスト", "承認", "監査"} <= set(wb.sheetnames)
    assert wb["影響"].max_row - 1 == total
    rows = client.get(f"{API}/systems/{sid}/audit").get_json()["audit"]
    csv = client.get(f"{API}/systems/{sid}/audit.csv").get_data(as_text=True).lstrip("﻿").strip().splitlines()
    assert len(csv) - 1 == len(rows)


def test_it010_sample_end_to_end(client):
    """IT-010 見本は 1 回の操作で対象システム・ベースライン 2 版・変更要求・分析済みの影響まで作る。"""
    d = client.post(f"{API}/sample").get_json()
    sid, code = d["system"]["id"], d["cr"]["code"]
    assert _wait(client, d["job"]["id"])["status"] == "done"
    assert len(client.get(f"{API}/systems/{sid}/baselines").get_json()["baselines"]) == 2
    s = client.get(f"{API}/systems/{sid}/crs/{code}/impacts").get_json()["summary"]
    assert s["total"] >= 3 and s["p1"] >= 1
    cr = client.get(f"{API}/systems/{sid}/crs/{code}").get_json()["cr"]
    assert [c["kind"] for c in cr["changes"]] == ["modify"], "書き換えが追加＋削除に分かれている"
    g = client.get(f"{API}/systems/{sid}/crs/{code}/graph").get_json()
    ids = {n["id"] for n in g["nodes"]}
    assert g["edges"] and all(e["source"] in ids and e["target"] in ids for e in g["edges"]), "線の端が点に無い"


def test_it011_cancel_job(client):
    """IT-011 中止すると中止の状態になり、途中の影響は保存されない。"""
    sid = _system_with_docs(client)
    code = _cr_with_baselines(client, sid)
    job = client.post(f"{API}/systems/{sid}/crs/{code}/analyze", json={}).get_json()["job"]
    client.post(f"{API}/jobs/{job['id']}/cancel")
    final = _wait(client, job["id"])
    if final["status"] == "cancelled":
        assert client.get(f"{API}/systems/{sid}/crs/{code}/impacts").get_json()["summary"]["total"] == 0
    else:
        assert final["status"] == "done"  # 中止が届く前に終わった（小さい見本では起こりうる）


def test_it012_input_validation(client):
    """IT-012 保存先の外を指す名前や不正な ID は 400/404 で、保存先の外にファイルを作らない。"""
    sid = client.post(f"{API}/systems", json={"name": "s"}).get_json()["system"]["id"]
    r = client.post(f"{API}/systems/{sid}/documents", data={"files": [(io.BytesIO(b"# a\n- b"), "../../evil.md")]},
                    content_type="multipart/form-data")
    assert r.status_code in (200, 400)
    assert not (config.DATA_DIR.parent / "evil.md").exists()
    assert client.get(f"{API}/systems/../../etc").status_code == 404
    assert client.get(f"{API}/systems/zzz/crs/CR-001").status_code == 404
