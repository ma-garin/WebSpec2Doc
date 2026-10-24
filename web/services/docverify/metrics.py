"""終了判定の集計と Excel 出力。"""
from __future__ import annotations

import io
from collections import Counter

_OPEN = {"adopted"}  # 候補（未検討）は未対応に数えず、別の理由として出す


def summary(insp: dict) -> dict:
    fs = insp.get("findings", []) or []
    by_status = dict(Counter(f.get("status") for f in fs))
    from web.services.docverify.settings import normalize_severity
    sev = lambda f: normalize_severity(f.get("severity"))  # noqa: E731
    by_severity = dict(Counter(sev(f) for f in fs if f.get("status") != "rejected"))
    adopted = sum(1 for f in fs if f.get("status") in ("adopted", "fixed", "verified"))
    n = len(insp.get("passages", []) or [])
    open_critical = sum(1 for f in fs if sev(f) == "critical" and f.get("status") in _OPEN)
    open_major = sum(1 for f in fs if sev(f) == "major" and f.get("status") in _OPEN)
    ec = insp.get("exit_criteria") or {}
    max_c = ec.get("max_open_critical", 0)
    max_w = ec.get("max_open_major", ec.get("max_open_warning", 5))
    reasons = []
    if not insp.get("document"):
        reasons.append("文書がまだ取り込まれていません。")
    pending = sum(1 for f in fs if f.get("status") == "candidate")
    if pending:
        reasons.append(f"未検討の候補が {pending} 件あります。「指摘にする」か「却下」を選んでください。")
    if open_critical > max_c:
        reasons.append(f"未対応の重大な指摘が {open_critical} 件あります（上限 {max_c} 件）。")
    if open_major > max_w:
        reasons.append(f"未対応の主要な指摘が {open_major} 件あります（上限 {max_w} 件）。")
    return {"by_status": by_status, "by_severity": by_severity, "quality": quality(insp),
            "density": round(adopted / n, 3) if n else 0.0,
            "open_critical": open_critical, "open_major": open_major,
            "open_warning": open_major,  # v1 画面との互換
            "exit_ok": not reasons, "reasons": reasons,
            "history": history(insp), "unread_lines": unread_lines(insp), "scope": scope(insp)}


_ITEMS = ("necessary", "unambiguous", "complete", "consistent", "feasible", "verifiable", "traceable", "singular")
_LIVE = {"candidate", "adopted"}  # 却下・修正済みは減点しない


def half_up(x: float) -> int:
    """四捨五入（0 以上の値）。round() は偶数丸めのため使わない。"""
    return int(x + 0.5)


def quality(insp: dict) -> dict:
    """29148 の特性ごとに「未解決の指摘が無い断片の割合」を 0〜100 で返す。指摘の無い特性は検査対象外と区別しない。"""
    passages = insp.get("passages", []) or []
    n = len(passages)
    live_all = [f for f in insp.get("findings", []) or [] if f.get("status") in _LIVE]
    live = [f for f in live_all if f.get("checklist_item")]  # 特性別は特性の付いた指摘だけ。総合は全指摘
    items = {}
    for item in _ITEMS:
        hit = {f.get("passage_id") for f in live if f["checklist_item"] == item}
        items[item] = {"open": sum(1 for f in live if f["checklist_item"] == item),
                       "score": half_up(100 * (1 - len(hit) / n)) if n else None}
        items[item]["label"] = ITEM_LABEL.get(item, item)
    hit_any = {f.get("passage_id") for f in live_all}
    return {"items": items, "note": list(QUALITY_NOTE), "overall": half_up(100 * (1 - len(hit_any) / n)) if n else None, "passages": n}


def _safe(v):
    s = "" if v is None else str(v)
    return "'" + s if s.startswith(("=", "+", "-", "@")) else s


def to_xlsx(insp: dict) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "指摘一覧"
    cols = ["id", "severity", "viewpoint", "message", "location", "quote", "certainty", "status", "answer", "origin"]
    ws.append(["No.", "重大度", "観点", "指摘", "場所", "原文", "確度", "状態", "回答", "起点"])
    for f in insp.get("findings", []) or []:
        ws.append([_safe(f.get(c)) for c in cols])
    s = summary(insp)
    ws2 = wb.create_sheet("集計")
    ws2.append(["項目", "値"])
    for k, v in s["by_status"].items():
        ws2.append([f"状態: {k}", v])
    for k, v in s["by_severity"].items():
        ws2.append([f"重大度: {k}", v])
    ws2.append(["断片あたりの採用指摘", s["density"]])
    ws2.append(["未対応の重大", s["open_critical"]])
    ws2.append(["未対応の主要", s["open_major"]])
    ws2.append(["品質スコア（全体）", s["quality"]["overall"]])
    for k, v in s["quality"]["items"].items():
        ws2.append([f"品質スコア: {ITEM_LABEL.get(k, k)}", v["score"]])
    ws2.append(["終了判定", "可" if s["exit_ok"] else "不可"])
    for r in s["reasons"]:
        ws2.append(["理由", r])
    for h in s["history"]:
        ws2.append([f"版 {h['version']} の品質スコア", "—" if h["overall"] is None else h["overall"]])
    _scope_sheet(wb, s["scope"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ITEM_LABEL = {"necessary": "必要", "unambiguous": "曖昧でない", "complete": "完全", "consistent": "一貫",
               "feasible": "実現可能", "verifiable": "検証可能", "traceable": "追跡可能", "singular": "単一"}
_ITEM_LABEL = ITEM_LABEL  # 旧名（互換）
_LABEL = {"critical": "重大", "major": "主要", "minor": "軽微"}


_PAIR_SEV = {"重大": "critical", "警告": "major", "提案": "minor"}


def _pair_reviews() -> list[dict]:
    """対検査（pair）の保存先から、採否を反映した指摘を読む。pair 系は inspections に含まれない。"""
    import json
    from datetime import datetime, timezone
    from web import config
    from web.services.docverify import pair_review as pr
    from web.tenancy import scoped_output_dir
    root = scoped_output_dir(config.DATA_DIR) / "docverify"
    out = []
    if not root.exists():
        return out
    for fp in sorted(root.glob("*/pair/findings.json")):
        d = fp.parent
        try:
            fs = pr._replay(pr._findings(d), pr._decisions(d))
            meta = json.loads((d / "docs" / "meta.json").read_text(encoding="utf-8")) if (d / "docs" / "meta.json").exists() else {}
        except (OSError, ValueError):
            continue
        mt = max((p.stat().st_mtime for p in (d / "findings.json", d / "decisions.json") if p.exists()), default=0)
        out.append({"id": d.parent.name, "name": meta.get("name") or meta.get("review_name") or d.parent.name,
                    "inspection_id": meta.get("inspection_id"), "findings": fs,
                    "updated_at": datetime.fromtimestamp(mt, timezone.utc).isoformat() if mt else None})
    return out


def _rate(fixed: int, valid: int):
    """分母 0 は None（0.0 を返すと「0% 解消」に見える）。"""
    return round(fixed / valid, 3) if valid else None


def overall(inspections: list[dict], pair_reviews: list[dict] | None = None) -> dict:
    """レビュー別・観点別・重大度別の件数と修正完了率。pair 系（対検査）も数える。"""
    from web.services.docverify.settings import normalize_severity
    pairs = _pair_reviews() if pair_reviews is None else pair_reviews
    pair_of = {p["inspection_id"] for p in pairs if p.get("inspection_id")}
    by_review, by_viewpoint, by_severity = [], Counter(), Counter()
    total_valid = total_fixed = total_crit = 0

    def add(rid, name, valid, fixed_n, crit_n, updated_at):
        nonlocal total_valid, total_fixed, total_crit
        total_valid += len(valid); total_fixed += fixed_n; total_crit += crit_n
        by_review.append({"review_id": rid, "review_name": name, "findings": len(valid),
                          "fix_rate": _rate(fixed_n, len(valid)), "unresolved_critical": crit_n, "updated_at": updated_at})

    for insp in inspections:
        if insp.get("id") in pair_of:
            continue
        fs = [f for f in insp.get("findings") or [] if f.get("status") != "rejected"]
        valid = [f for f in fs if f.get("status") != "candidate"]
        done = [f for f in valid if f.get("status") in ("fixed", "verified")]
        sev = [normalize_severity(f.get("severity")) for f in valid]
        by_viewpoint.update(f.get("checklist_item") or f.get("viewpoint") or "未分類" for f in valid)
        by_severity.update(sev)
        crit = sum(1 for f in valid if f not in done and normalize_severity(f.get("severity")) == "critical")
        add(insp.get("id"), insp.get("name"), valid, len(done), crit, insp.get("updated_at") or insp.get("created_at"))
    for pr_ in pairs:
        valid = [f for f in pr_["findings"] if f.get("status") != "不採用"]
        done = [f for f in valid if f.get("status") == "解消済み"]
        sev = [_PAIR_SEV.get(f.get("severity"), "") or normalize_severity(f.get("severity")) for f in valid]
        by_viewpoint.update(f.get("viewpoint") or "未分類" for f in valid)
        by_severity.update(sev)
        crit = sum(1 for f, v in zip(valid, sev) if f.get("status") != "解消済み" and v == "critical")
        add(pr_["id"], pr_["name"], valid, len(done), crit, pr_.get("updated_at"))
    return {"by_review": by_review, "by_viewpoint": dict(by_viewpoint), "by_severity": dict(by_severity),
            "findings": total_valid, "fixed": total_fixed, "unresolved_critical": total_crit,
            "fix_rate": _rate(total_fixed, total_valid)}


# ---- R1: 品質スコアの説明・版ごとの記録・検査した範囲（WS2D-DD-006 5 章・9 章・10 章） ----

QUALITY_NOTE = (
    "スコアは、未解決の指摘（候補・採用）が無い断片の割合を 0〜100 で示します。",
    "式は 100 ×（1 − 未解決の指摘がある断片数 ÷ 断片数）です。",
    "指摘が無い特性も 100 になるため、文書の内容が正しいことや、検査の網羅性は示しません。",
)
_M04 = "読めなかった要素は数えていません。文書を取り込み直すと数えます。"
_M13 = "— 文書を取り込むと表示されます。"
_M20 = "文書はまだ取り込まれていません。"
_M21 = "参加者が登録されていません。計画の画面で登録できます。"
_M22 = "観点チェックリストを使っていません。"
_M30 = "未検証。見つかった指摘は検出規則と人の確認の範囲に限られ、ほかに欠陥が無いことは示しません。"
_M31 = "未検証。記述が業務として正しいかは判断していません。"
_M33 = "未測定。規則ごとの適合率・再現率はまだ測っていません。"
_UNREAD_KINDS = (("images", "画像"), ("comments", "コメント"), ("shapes", "図形"))
_ADOPTED = {"adopted", "fixed", "verified"}


def quality_brief(insp: dict, final: bool) -> dict:
    q = quality(insp)
    return {"overall": q["overall"], "items": {k: v["score"] for k, v in q["items"].items()}, "final": bool(final)}


def freeze_current(insp: dict) -> list[dict]:
    versions = list(insp.get("versions") or [])
    if not versions:
        return []
    return [*versions[:-1], {**versions[-1], "quality": quality_brief(insp, True)}]


def _ended(insp: dict) -> bool:
    return insp.get("phase") == "closed" or insp.get("status") == "closed"


def history(insp: dict) -> list[dict]:
    versions = insp.get("versions") or []
    out = []
    for i, v in enumerate(versions):
        q = v.get("quality") or None
        final = bool(q and q.get("final"))
        row = {"version": v.get("version"), "name": v.get("name"), "uploaded_at": v.get("uploaded_at"),
               "overall": q.get("overall") if q else None, "final": final, "current": False}
        if i == len(versions) - 1 and not final and not _ended(insp):
            row = {**row, "overall": quality(insp)["overall"], "current": True}
        out.append(row)
    return out


def unread_lines(insp: dict) -> list[str]:
    doc = insp.get("document")
    if not doc:
        return []
    if doc.get("unread") is None:
        return [_M04]
    from web.services.docverify import unread
    return unread.lines(doc["unread"])


def _unread_content(insp: dict) -> str:
    if not insp.get("document"):
        return "該当なし"
    u = (insp.get("document") or {}).get("unread") or {}
    if not u.get("scanned"):
        return "読めなかった要素は数えていません（未確認）。"
    parts = [f"{label} {u.get(k, 0)}" for k, label in _UNREAD_KINDS if u.get(k, 0)]
    return f"{'・'.join(parts)} の中身は検査していません。" if parts else "該当なし"


def _checked(insp: dict) -> list[dict]:
    doc = insp.get("document")
    fs = [f for f in insp.get("findings") or [] if f.get("origin", "auto") == "auto"]
    st = Counter(f.get("status") for f in fs)
    ps = insp.get("participants") or []
    overall_ = quality(insp)["overall"]
    rows = [
        ("文書", f"{doc.get('name')}（版 {doc.get('version')}・{doc.get('format')}）" if doc else _M20),
        ("断片数", f"{len(insp.get('passages') or [])} 断片"),
        ("読めなかった要素", " ".join(unread_lines(insp))),
        ("自動検出", f"規則による候補 {len(fs)} 件（採用 {sum(st[k] for k in _ADOPTED)}・"
                     f"却下 {st['rejected']}・未検討 {st['candidate']}）"),
        ("人の確認", f"レビュアー {sum(1 for p in ps if p.get('done'))}/{len(ps)} 人が確認を終えました" if ps else _M21),
        ("観点チェックリスト", f"確認済みの観点 {sum(1 for v in (insp.get('checks') or {}).values() if v)} 項目"
                               if insp.get("checklist_id") else _M22),
        ("品質スコア", _M13 if overall_ is None else f"総合 {overall_}"),
    ]
    return [{"label": k, "value": v} for k, v in rows]


def scope(insp: dict) -> dict:
    unchecked = [("網羅性", _M30), ("意味の正しさ", _M31), ("読めなかった要素の中身", _unread_content(insp)),
                 ("スコアの限界", QUALITY_NOTE[2]), ("規則の精度", _M33)]
    return {"checked": _checked(insp), "unchecked": [{"label": k, "value": v} for k, v in unchecked]}


def _scope_sheet(wb, sc: dict) -> None:
    from openpyxl.styles import Alignment
    ws = wb.create_sheet("検査した範囲")
    ws.append([_safe(h) for h in ("区分", "項目", "内容")])
    for kind, key in (("検査した範囲", "checked"), ("検査していないこと", "unchecked")):
        for row in sc[key]:
            ws.append([_safe(kind), _safe(row["label"]), _safe(row["value"])])
    for col, width in (("A", 18), ("B", 24), ("C", 80)):
        ws.column_dimensions[col].width = width
    for cell in ws["C"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
