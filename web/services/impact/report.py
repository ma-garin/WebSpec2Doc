"""Excel レポート（概要・影響・テスト・承認・監査の 5 シート）。DD-009 / BD-016。"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from web.services.impact import db
from web.services.impact.repo import audit, crs, impacts

_STATE = {"open": "要確認", "required": "対応要", "not_required": "対応不要"}
_DECISION = {"approved": "承認", "returned": "差し戻し"}
_EVIDENCE = {"link": "つながり", "semantic": "意味検索", "keyword": "語の一致", "llm": "LLM"}


def _sheet(wb, title, head, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(head)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
        for cell in col[1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def _overview(cr, engine) -> list[list]:
    s = cr["summary"]
    return [["変更要求", cr["code"]], ["件名", cr["title"]], ["理由", cr["reason"]],
            ["状態", cr["status_label"]], ["起票者", cr["created_by"]], ["起票日時", cr["created_at"]],
            ["波及の段数", cr["depth"]], ["影響の件数", s["total"]], ["P1 / P2 / P3", f"{s['p1']} / {s['p2']} / {s['p3']}"],
            ["要確認", s["open"]], ["テストの抜け", s["test_gap"]], ["再テスト", s["retest"]],
            ["エンジン", engine.get("mode", "") if engine else "未分析"],
            ["エンジンの補足", engine.get("reason", "") if engine else ""]]


def _impact_rows(rows) -> list[list]:
    out = []
    for r in rows:
        ev = "／".join(f"{_EVIDENCE.get(e.get('type'), e.get('type'))}: {e.get('detail', '')}"
                      for e in r["evidence"] if isinstance(e, dict))
        formula = f"影響度 {r['effective_impact']} × 発生可能性 {r['effective_likelihood']} = {r['risk']}"
        out.append([r["item_code"] or r["id"], f"P{r['priority']}", formula, r["document_name"],
                    r["location"], r["text"], ev, _STATE.get(r["state"], r["state"]), r["resolve_reason"],
                    r["override_reason"]])
    return out


def build_xlsx(conn, sid, code) -> bytes:
    cr = crs.get(conn, sid, code)
    engine = db.loads((cr.get("last_job") or {}).get("engine"), {})
    rows = impacts.list(conn, cr["id"])
    wb = Workbook()
    _sheet(wb, "概要", ["項目", "値"], _overview(cr, engine), first=True)
    _sheet(wb, "影響", ["ID", "優先度", "リスク式", "文書", "場所", "項目", "根拠", "状態", "解消理由", "修正理由"],
           _impact_rows(rows))
    _sheet(wb, "テスト", ["影響項目", "優先度", "選定テスト", "テストの抜け", "観点"],
           [[r["item_code"] or r["id"], f"P{r['priority']}", "、".join(r["tests"]),
             "あり" if r.get("test_gap") else "なし", r.get("test_idea")] for r in rows])
    _sheet(wb, "承認", ["決定", "理由", "承認者", "日時"],
           [[_DECISION.get(a["decision"], a["decision"]), a["reason"], a["actor"], a["at"]] for a in cr["approvals"]])
    logs = audit.query(conn, sid, cr_code=code, limit=5000)
    _sheet(wb, "監査", ["日時", "操作者", "操作", "対象", "変更前", "変更後"],
           [[a["at"], a["actor"], a["action"], a["target"], db.dumps(a["before"]) if a["before"] is not None else "",
             db.dumps(a["after"]) if a["after"] is not None else ""] for a in reversed(logs)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
