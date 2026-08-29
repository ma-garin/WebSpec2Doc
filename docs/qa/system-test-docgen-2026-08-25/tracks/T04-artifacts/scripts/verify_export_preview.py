"""C038-C045: エクスポート（xlsx/zip/spec-ts）と /preview。"""
import io, json, urllib.parse, zipfile, sys
sys.path.insert(0, "docs/qa/system-test-docgen-2026-08-25/tracks/T04-artifacts/scripts")
from qa_common import QDOM, OUT, EV, http, check, save
from openpyxl import load_workbook

EXPECTED = ["画面一覧", "フォーム", "項目定義書", "境界値データ", "テスト設計", "テストケース", "遷移表"]
st, x1 = http(f"/api/export/spec-xlsx?domain={QDOM}")
(EV / "spec_export1.xlsx").write_bytes(x1 if st == 200 else b"")
if st == 200:
    wb = load_workbook(io.BytesIO(x1), read_only=True)
    names = wb.sheetnames
    check("T04-C038", len(names) == 7, f"status=200 sheets={names}（期待名={EXPECTED}）")
    report = json.loads((OUT / "report.json").read_text(encoding="utf-8"))
    n_screens = len(report.get("screens", []))
    ws = wb[names[0]]
    rows = sum(1 for row in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in row))
    check("T04-C039", rows - 1 == n_screens, f"シート'{names[0]}' データ行={rows - 1} / report.json screens={n_screens}")
else:
    check("T04-C038", False, f"status={st} body={x1[:200]}")
    check("T04-C039", False, "C038 失敗のため判定不能")
    names = []
st2, x2 = http(f"/api/export/spec-xlsx?domain={QDOM}")
if st2 == 200:
    names2 = load_workbook(io.BytesIO(x2), read_only=True).sheetnames
    dup = [n for n in names2 if names2.count(n) > 1]
    check("T04-C040", names2 == names and not dup, f"再取得 sheets={names2} 重複={dup}")
else:
    check("T04-C040", False, f"再取得 status={st2}")
st3, b3 = http("/api/export/spec-xlsx?domain=..%2Fetc")
check("T04-C041", st3 in (400, 404), f"不正ドメイン status={st3}")
st4, z = http(f"/download-zip?domain={QDOM}")
ok = False; det = f"status={st4}"
if st4 == 200:
    zf = zipfile.ZipFile(io.BytesIO(z))
    bad = zf.testzip()
    names_z = zf.namelist()
    ok = bad is None and any(n.endswith("report.json") for n in names_z)
    det = f"status=200 entries={len(names_z)} 破損={bad} report.json含む={any(n.endswith('report.json') for n in names_z)}"
check("T04-C042", ok, det)
st5, b5 = http(f"/api/report/{QDOM}/spec-ts")
det5 = b5[:150].decode("utf-8", "replace")
check("T04-C043", st5 in (200, 404), f"status={st5} body={det5}（404 は playwright_candidates.json 未生成の明示なら仕様どおり）")
qp = urllib.parse.quote(str(OUT / "report.json"), safe="")
qh = urllib.parse.quote(str(OUT / "report.html"), safe="")
s1, p1 = http(f"/preview?path={qp}")
s2, p2 = http(f"/preview?path={qh}")
check("T04-C044", s1 == 200 and s2 == 200 and p2.lstrip()[:1] == b"<", f"json={s1} html={s2}")
s3, _ = http(f"/preview?path={urllib.parse.quote(str(OUT / 'no_such_file.json'), safe='')}")
s4, _ = http("/preview?path=%2Fetc%2Fpasswd")
s5, p5 = http("/preview?path=" + urllib.parse.quote("output/example.com/report.json", safe=""))
save("preview_results.json", {"missing": s3, "outside_output": s4, "other_domain": s5})
check("T04-C045", s3 == 404 and s4 == 404, f"不在={s3} output外(/etc/passwd)={s4} 他ドメイン(example.com)={s5}（他ドメイン可否は observation）")
