"""System 06 R3（集合の検査・追跡・取り込み拡張・規則と語彙）の単体・結合テスト。

受入基準: WS2D-RD-006 FR-06-016〜018・025〜027・120〜123・130〜135。設計: WS2D-DD-006 R3 8 章。
"""
from __future__ import annotations

import io

import openpyxl
import pytest

import app as appmod
from web import config
from web.routes import doc_verify_r3
from web.services.docverify import ingest, ingest_r3, setcheck, trace

ORIGIN = {"Origin": "http://localhost"}
EMPTY = "確認した範囲では指摘はありません。網羅性は未検証です。"


def _doc(did, text_lines, kind=None, name=None):
    ps = [{"id": f"P{i}", "text": t, "location": f"{i}行目"} for i, t in enumerate(text_lines, 1)]
    return {"id": did, "name": name or f"{did}.txt", "kind": kind, "passages": ps, "include_hidden": True}


def _xlsx(rows, hidden_sheet=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    if hidden_sheet:
        h = wb.create_sheet("隠し")
        for r in hidden_sheet:
            h.append(r)
        h.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- 集合の検査（FR-06-120〜123） ----

def test_120_near_duplicate_is_one_pair():
    pairs = setcheck.find_pairs(setcheck.units([_doc("A", ["利用者は、検索結果を一覧画面で確認できること。"]),
                                                _doc("B", ["利用者は検索結果を一覧画面で確認できること。"])]))
    assert [p["kind"] for p in pairs] == ["duplicate"]
    assert pairs[0]["similarity"] >= setcheck.DUP_THRESHOLD


def test_120_normalizes_space_width_and_punctuation():
    assert setcheck.normalize("ＡＢＣ　は ５秒，です。") == setcheck.normalize("ABCは5秒、です")


def test_120_different_sentences_no_pair():
    us = setcheck.units([_doc("A", ["利用者は検索結果を一覧で確認できること。", "管理者はログを出力すること。"])])
    assert setcheck.find_pairs(us) == []


def test_121_numbers_with_same_unit_are_contradiction():
    us = setcheck.units([_doc("A", ["検索処理は5秒以内に応答すること。"]), _doc("B", ["検索処理は10秒以内に応答すること。"])])
    pairs = setcheck.find_pairs(us)
    assert len(pairs) == 1 and pairs[0]["kind"] == "contradiction"
    assert "秒" in pairs[0]["reason"]


def test_121_negation_vs_affirmation_is_contradiction():
    us = setcheck.units([_doc("A", ["管理者はデータを削除できる。", "管理者はデータを削除できない。"])])
    pairs = setcheck.find_pairs(us)
    assert [p["kind"] for p in pairs] == ["contradiction"]


def test_121_different_subject_not_contradiction():
    us = setcheck.units([_doc("A", ["検索処理は5秒以内に応答すること。", "登録処理は10秒以内に応答すること。"])])
    assert [p["kind"] for p in setcheck.find_pairs(us)] == []


def _insp_with_pair():
    docs = [_doc("A", ["検索処理は5秒以内に応答すること。"]), _doc("B", ["検索処理は10秒以内に応答すること。"])]
    return {"findings": [{"id": "F3", "status": "candidate"}], "set_check": setcheck.run(docs)}


def test_122_decide_contradiction_adds_finding():
    insp = _insp_with_pair()
    pid = insp["set_check"]["pairs"][0]["id"]
    out = setcheck.decide(insp, pid, "contradiction", "山田")
    assert len(out["findings"]) == 2 and out["findings"][-1]["type"] == "矛盾"
    assert out["findings"][-1]["id"] == "F4" and out["findings"][-1]["status"] == "adopted"
    assert len(insp["findings"]) == 1  # 入力は変えない
    again = setcheck.decide(out, pid, "ok", "山田")
    assert len(again["findings"]) == 2 and again["findings"][-1]["status"] == "rejected"
    assert again["set_check"]["decisions"][pid]["decision"] == "ok"


def test_122_invalid_decision_and_unknown_pair():
    insp = _insp_with_pair()
    with pytest.raises(ValueError):
        setcheck.decide(insp, insp["set_check"]["pairs"][0]["id"], "maybe", "x")
    with pytest.raises(KeyError):
        setcheck.decide(insp, "Snope", "ok", "x")


def test_122_decision_carried_over_on_rerun():
    insp = _insp_with_pair()
    pid = insp["set_check"]["pairs"][0]["id"]
    out = setcheck.decide(insp, pid, "duplicate", "x")
    docs = [_doc("A", ["検索処理は5秒以内に応答すること。"]), _doc("B", ["検索処理は10秒以内に応答すること。"])]
    assert setcheck.run(docs, out["set_check"])["decisions"][pid]["decision"] == "duplicate"


def test_123_zero_candidates_shows_coverage_unverified():
    v = setcheck.view(setcheck.run([_doc("A", ["利用者はログインできること。"])]))
    assert v["pairs"] == [] and v["empty_message"] == EMPTY
    assert v["coverage"] == "網羅性: 未検証" and "網羅性: 未検証" in v["unchecked"]


# ---- 追跡表（FR-06-131〜135） ----

def _trace_docs():
    return [_doc("D1", ["REQ-001 ログインできる。", "REQ-002 検索できる。", "REQ-003 出力できる。"], "requirement", "要件.txt"),
            _doc("D2", ["設計 1 は REQ-001 を満たす。", "設計 2 は REQ-002 を満たす。"], "design", "設計.txt")]


def test_131_132_trace_rows_rate_and_orphans():
    t = trace.build(_trace_docs(), "REQ")
    assert [r["id"] for r in t["rows"]] == ["REQ-001", "REQ-002", "REQ-003"]
    assert t["rows"][2]["design"] == [] and t["rows"][0]["design"]
    assert t["complete"] == {"covered": 2, "total": 3, "rate": 67}
    assert t["orphans"] == [] and t["coverage"] == "網羅性: 未検証"


def test_132_orphan_downstream_id():
    docs = [*_trace_docs(), _doc("D3", ["試験 9 は REQ-009 を確かめる。"], "test", "試験.txt")]
    t = trace.build(docs, "REQ")
    assert [o["id"] for o in t["orphans"]] == ["REQ-009"]
    assert t["complete"]["covered"] == 0  # テストに対応がある上流 ID は無い


def test_134_format_picks_only_that_form():
    ids = trace.extract([{"text": "FR-001 と NFR-002 と REQ-003 と FR-010", "location": "1行目"}], "FR")
    assert [x["id"] for x in ids] == ["FR-001", "FR-010"]
    with pytest.raises(ValueError):
        trace.pattern("XX")


def test_135_document_without_ids_is_excluded_with_reason():
    docs = [_trace_docs()[0], _doc("D9", ["設計の説明だけで ID が無い。"], "design", "設計2.txt")]
    t = trace.build(docs, "REQ")
    assert t["unreadable"][0]["message"] == "この文書から ID を読み取れませんでした。ID の書式を確かめてください。"
    assert [c["kind"] for c in t["columns"]] == ["requirement", "design"]


def test_133_xlsx_has_same_rows():
    t = trace.build(_trace_docs(), "REQ")
    ws = openpyxl.load_workbook(io.BytesIO(trace.to_xlsx(t)))["追跡表"]
    assert ws.max_row - 1 == len(t["rows"]) == 3
    assert ws.cell(row=4, column=3).value == "（対応なし）"


# ---- 取り込み拡張（FR-06-016〜018） ----

def test_016_table_row_is_one_passage(tmp_path):
    data = _xlsx([["項目名", "型", "必須"], ["氏名", "文字列", "○"]])
    _, ps = ingest_r3.ingest_r3("項目.xlsx", data, tmp_path)
    assert any(p["text"] == "項目名: 氏名／型: 文字列／必須: ○" for p in ps)


def test_017_hidden_sheet_is_marked(tmp_path):
    data = _xlsx([["表に見える文です。"]], hidden_sheet=[["隠れた文です。"]])
    doc, ps = ingest_r3.ingest_r3("隠し.xlsx", data, tmp_path)
    marks = {p["text"]: p["hidden"] for p in ps}
    assert marks["隠れた文です。"] is True and marks["表に見える文です。"] is False
    assert doc["hidden_passages"] == 1


def test_018_encrypted_office_is_not_called_broken(tmp_path):
    data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600
    with pytest.raises(ingest.IngestError) as e:
        ingest_r3.ingest_r3("鍵付き.xlsx", data, tmp_path)
    assert "パスワードを外してから" in str(e.value) and "破損" not in str(e.value) and "壊れ" not in str(e.value)


# ---- 規則・語彙・要求属性（FR-06-025〜027） ----

def test_026_missing_verification_column():
    out = setcheck.missing_attributes([{"headers": ["ID", "要求", "優先度", "根拠"], "location": "S 1行目", "passage_id": "P1"}])
    assert [f["message"] for f in out] == ["検証方法がありません。要求一覧に「検証方法」の列を足してください。"]


def test_026_full_columns_no_candidate_and_non_requirement_table_ignored():
    assert setcheck.missing_attributes([{"headers": ["ID", "要求", "優先度", "根拠", "検証方法"]}]) == []
    assert setcheck.missing_attributes([{"headers": ["項目名", "型", "必須"]}]) == []


def test_027_term_variant():
    ps = [{"id": "P1", "text": "顧客は登録できる。", "location": "1行目"}, {"id": "P2", "text": "お客様は検索できる。", "location": "2行目"}]
    out = setcheck.term_variants(ps, [{"canonical": "お客様", "variants": ["顧客"]}])
    assert len(out) == 1 and "用語ゆれ" in out[0]["message"]
    assert setcheck.term_variants(ps[:1], [{"canonical": "お客様", "variants": ["顧客"]}]) == []


def test_025_forbidden_and_allowed_words():
    ps = [{"id": "P1", "text": "画面は適宜更新すること。", "location": "1行目"}]
    found = [{"rule": "REQ.OPEN_ENDED", "quote": "適宜"}]
    out = setcheck.apply_rules(found, ps, {"forbidden": ["更新"], "allowed": ["適宜"], "disabled_rules": []})
    assert [f["rule"] for f in out] == ["ORG.FORBIDDEN"]


# ---- 結合（Flask テストクライアント） ----

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setenv("FLASK_TESTING", "1")
    monkeypatch.setenv("WEBSPEC2DOC_ALLOW_LOCAL", "1")
    monkeypatch.setenv("WEBSPEC2DOC_AUTH_MODE", "off")
    app = appmod.app
    if "doc_verify_r3" not in app.blueprints:  # 結線（web/__init__.py）は本線が行う。テストでは自分で登録する
        got = getattr(app, "_got_first_request", False)
        app._got_first_request = False
        app.register_blueprint(doc_verify_r3.bp)
        app._got_first_request = got
    return app.test_client()


def _create(client) -> str:
    r = client.post("/api/docverify/inspections", json={"name": "R3", "goal": "確認"}, headers=ORIGIN)
    assert r.status_code == 201
    return r.get_json()["inspection"]["id"]


def _upload(client, iid, name, data, kind, include_hidden="1"):
    return client.post(f"/api/docverify/inspections/{iid}/r3/documents", headers=ORIGIN,
                       data={"kind": kind, "include_hidden": include_hidden, "file": (io.BytesIO(data), name)},
                       content_type="multipart/form-data")


def test_130_register_two_documents_with_kinds(client):
    iid = _create(client)
    assert _upload(client, iid, "要件.txt", "REQ-001 ログインできる。\n".encode(), "requirement").status_code == 201
    r = _upload(client, iid, "設計.txt", "設計は REQ-001 を満たす。\n".encode(), "design")
    docs = r.get_json()["documents"]
    assert [(d["name"], d["kind_label"]) for d in docs] == [("要件.txt", "要件"), ("設計.txt", "設計")]
    assert all("passages" not in d for d in docs)
    assert r.get_json()["trace"]["complete"]["rate"] == 100


def test_130_bad_kind_and_missing_inspection(client):
    iid = _create(client)
    assert _upload(client, iid, "a.txt", "文。".encode(), "other").status_code == 400
    assert client.get("/api/docverify/inspections/000000000000/r3").status_code == 404


def test_018_api_encrypted_message(client):
    iid = _create(client)
    r = _upload(client, iid, "鍵付き.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600, "design")
    assert r.status_code == 400 and "パスワードを外してから" in r.get_json()["error"]


def test_025_disabled_rule_suppresses_candidate(client):
    iid = _create(client)
    text = "システムは必ず応答すること。\n".encode()
    before = _upload(client, iid, "a.txt", text, "requirement").get_json()["documents"][0]["candidates"]
    assert any(c["rule"].endswith("UNIVERSAL") for c in before)
    assert client.put("/api/docverify/r3/settings", json={"disabled_rules": ["UNIVERSAL"]}, headers=ORIGIN).status_code == 200
    after = _upload(client, iid, "b.txt", text, "requirement").get_json()["documents"][1]["candidates"]
    assert not any(c["rule"].endswith("UNIVERSAL") for c in after)
    assert client.put("/api/docverify/r3/settings", json={"disabled_rules": ["NOPE"]}, headers=ORIGIN).status_code == 400


def test_026_027_api_candidates(client):
    iid = _create(client)
    client.put("/api/docverify/r3/settings", json={"synonyms": [{"canonical": "お客様", "variants": ["顧客"]}]}, headers=ORIGIN)
    data = _xlsx([["ID", "要求", "優先度", "根拠"], ["REQ-001", "顧客は登録できること。", "高", "契約"],
                  ["REQ-002", "お客様は検索できること。", "中", "契約"]])
    c = _upload(client, iid, "要求.xlsx", data, "requirement").get_json()["documents"][0]["candidates"]
    msgs = [x["message"] for x in c]
    assert any("検証方法がありません" in m for m in msgs) and any("用語ゆれ" in m for m in msgs)


def test_017_api_hidden_excluded_from_detection(client):
    iid = _create(client)
    data = _xlsx([["表の文です。"]], hidden_sheet=[["システムは必ず応答すること。"]])
    c = _upload(client, iid, "隠し.xlsx", data, "requirement", include_hidden="0").get_json()["documents"][0]
    assert c["hidden_passages"] == 1 and not any(x["rule"].endswith("UNIVERSAL") for x in c["candidates"])


def test_120_123_122_api_setcheck_flow(client):
    iid = _create(client)
    _upload(client, iid, "a.txt", "利用者はログインできること。\n".encode(), "requirement")
    r = client.post(f"/api/docverify/inspections/{iid}/r3/setcheck", json={}, headers=ORIGIN)
    sc = r.get_json()["set_check"]
    assert sc["pairs"] == [] and sc["empty_message"] == EMPTY and sc["coverage"] == "網羅性: 未検証"
    _upload(client, iid, "b.txt", "検索処理は5秒以内に応答すること。\n".encode(), "design")
    _upload(client, iid, "c.txt", "検索処理は10秒以内に応答すること。\n".encode(), "test")
    sc = client.post(f"/api/docverify/inspections/{iid}/r3/setcheck", json={}, headers=ORIGIN).get_json()["set_check"]
    assert [p["kind"] for p in sc["pairs"]] == ["contradiction"]
    pid = sc["pairs"][0]["id"]
    r = client.post(f"/api/docverify/inspections/{iid}/r3/setcheck/{pid}", json={"decision": "contradiction"}, headers=ORIGIN)
    assert r.status_code == 200 and r.get_json()["set_check"]["pairs"][0]["decision"] == "contradiction"
    findings = client.get(f"/api/docverify/inspections/{iid}").get_json()["inspection"]["findings"]
    assert [f["type"] for f in findings if f.get("origin") == "set"] == ["矛盾"]
    bad = client.post(f"/api/docverify/inspections/{iid}/r3/setcheck/{pid}", json={"decision": "x"}, headers=ORIGIN)
    assert bad.status_code == 400


def test_setcheck_without_documents_is_400(client):
    iid = _create(client)
    assert client.post(f"/api/docverify/inspections/{iid}/r3/setcheck", json={}, headers=ORIGIN).status_code == 400


def test_133_134_api_trace_format_and_xlsx(client):
    iid = _create(client)
    _upload(client, iid, "要件.txt", "REQ-001 と FR-001 を定める。\nREQ-002 を定める。\n".encode(), "requirement")
    _upload(client, iid, "設計.txt", "REQ-001 を満たす。\n".encode(), "design")
    t = client.post(f"/api/docverify/inspections/{iid}/r3/trace", json={"format": "FR"}, headers=ORIGIN).get_json()
    assert t["trace_format"] == "FR" and [r["id"] for r in t["trace"]["rows"]] == ["FR-001"]
    assert client.post(f"/api/docverify/inspections/{iid}/r3/trace", json={"format": "ZZ"}, headers=ORIGIN).status_code == 400
    client.post(f"/api/docverify/inspections/{iid}/r3/trace", json={"format": "REQ"}, headers=ORIGIN)
    x = client.get(f"/api/docverify/inspections/{iid}/r3/trace.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(x.data))["追跡表"]
    shown = client.get(f"/api/docverify/inspections/{iid}/r3").get_json()["trace"]["rows"]
    assert ws.max_row - 1 == len(shown) == 2


def test_delete_document(client):
    iid = _create(client)
    did = _upload(client, iid, "a.txt", "文です。\n".encode(), "test").get_json()["documents"][0]["id"]
    assert client.delete(f"/api/docverify/inspections/{iid}/r3/documents/{did}", headers=ORIGIN).get_json()["documents"] == []
    assert client.delete(f"/api/docverify/inspections/{iid}/r3/documents/{did}", headers=ORIGIN).status_code == 404
