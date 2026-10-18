"""検証観点（<base>/docverify/_viewpoints.json）。V 字の対で 3 区分: 要件定義書・基本設計書・対検査。"""
from __future__ import annotations

import csv
import io
import re
import zlib

from web.services.docverify import settings

SCOPES = (
    ("req", "要件定義書", "システムテスト",
     "要件定義書を単体で検査する観点です。システムテストで確かめられる書き方になっているかを見ます。"),
    ("bas", "基本設計書", "結合テスト",
     "基本設計書を単体で検査する観点です。結合テストで画面や機能のつながりを確かめられるかを見ます。"),
    ("pair", "対検査", "要件定義書 ↔ 基本設計書",
     "要件定義書と基本設計書を突き合わせる観点です。要件が設計に正しく落ちているかを見ます。"),
)
SEV = ("critical", "major", "minor")
SEV_JA = {"重大": "critical", "主要": "major", "軽微": "minor", "critical": "critical", "major": "major", "minor": "minor"}
SCOPE_JA = {"要件定義書": "req", "基本設計書": "bas", "対検査": "pair", "req": "req", "bas": "bas", "pair": "pair"}
MAX_ROWS = 500
MAX_BYTES = 2 * 1024 * 1024

_D = [  # (scope, name, desc, kind, severity, enabled)
    ("req", "目的と範囲の明記", "システムの目的と、対象外にする範囲が書かれているか。", "human", "critical", 1),
    ("req", "受入条件の有無", "各要件に、合否を判定できる条件が付いているか。", "human", "critical", 1),
    ("req", "曖昧な表現", "「適切に」「なるべく」など、人によって解釈が変わる語を見つける。", "auto", "major", 1),
    ("req", "検証できない要件", "数値や期待結果が無く、テストで確かめられない要件を見つける。", "auto", "critical", 1),
    ("req", "要件 ID の重複", "同じ ID が複数の要件に使われていないか。", "auto", "minor", 0),
    ("bas", "画面遷移の抜け", "遷移元と遷移先が対になっていない画面を見つける。", "auto", "major", 1),
    ("bas", "入出力項目の定義", "各機能の入力・出力項目に、型と桁数が書かれているか。", "human", "major", 1),
    ("bas", "エラー時の動作", "異常時にどう振る舞うかが、機能ごとに書かれているか。", "human", "critical", 1),
    ("bas", "用語ゆれ", "同じ意味の言葉が、別の表記で書かれていないか。", "auto", "minor", 1),
    ("pair", "反映漏れ", "要件にあるのに、設計のどこにも出てこないものを見つける。", "auto", "critical", 1),
    ("pair", "余分な設計", "設計にあるのに、対応する要件が無いものを見つける。", "auto", "major", 1),
    ("pair", "数値条件の矛盾", "制限時間や件数などの数値が、両書類で食い違っていないか。", "auto", "critical", 1),
    ("pair", "表の食い違い", "項目一覧などの表で、項目名や値が両書類で一致しているか。", "auto", "major", 1),
    ("pair", "用語ゆれ", "両書類で同じ意味の言葉が別の表記になっていないか。", "auto", "minor", 1),
    ("pair", "曖昧な表現", "設計側で、要件の条件があいまいに言い換えられていないか。", "human", "major", 1),
    ("pair", "検証できない要件", "設計を見ても、要件が満たされたか確かめられないものを見つける。", "human", "critical", 1),
]
DEFAULTS = [{"id": f"{s}-{i + 1}", "scope": s, "name": n, "desc": d, "kind": k, "severity": v,
             "enabled": bool(e), "source": "builtin"} for i, (s, n, d, k, v, e) in enumerate(_D)]


def load() -> list:
    return settings._read("_viewpoints.json", DEFAULTS)


def update(body) -> list:
    """変更分 [{id, enabled?, severity?}] を反映して保存する。"""
    if not isinstance(body, list):
        raise ValueError("変更の形が正しくありません。画面を再読み込みしてやり直してください。")
    rows = load()
    by_id = {r["id"]: r for r in rows}
    for ch in body:
        if not isinstance(ch, dict) or ch.get("id") not in by_id:
            raise ValueError("存在しない観点が含まれています。画面を再読み込みしてやり直してください。")
        r = by_id[ch["id"]]
        if "enabled" in ch:
            r["enabled"] = bool(ch["enabled"])
        if "severity" in ch:
            if ch["severity"] not in SEV:
                raise ValueError("重大度は「重大」「主要」「軽微」のいずれかにしてください。")
            r["severity"] = ch["severity"]
    settings._write("_viewpoints.json", rows)
    return rows


def _table(name: str, data: bytes) -> list[list[str]]:
    if len(data) > MAX_BYTES:
        raise ValueError("ファイルが大きすぎます（2MB まで）。行を分けて取り込んでください。")
    low = (name or "").lower()
    if low.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb["観点"] if "観点" in wb.sheetnames else wb.worksheets[0]
            return [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
        except Exception:
            raise ValueError("Excel を読み取れません。.xlsx か .csv で保存し直して、もう一度選んでください。")
    if low.endswith(".csv"):
        for enc in ("utf-8-sig", "cp932"):
            try:
                return [[c.strip() for c in row] for row in csv.reader(io.StringIO(data.decode(enc)))]
            except UnicodeDecodeError:
                continue
        raise ValueError("文字コードを判別できません。UTF-8 で保存し直してください。")
    raise ValueError("取り込めるのは .xlsx と .csv だけです。形式を変えて選び直してください。")


def parse(name: str, data: bytes) -> dict:
    """取り込みの確認結果（保存しない）。列: 観点名・説明・対象書類・重大度。"""
    grid = [r for r in _table(name, data) if any(r)]
    if not grid:
        raise ValueError("行がありません。見出し行（観点名・説明・対象書類・重大度）と観点を入れてください。")
    head = grid[0]
    need = {k: (head.index(k) if k in head else None) for k in ("観点名", "説明", "対象書類", "重大度")}
    if need["観点名"] is None or need["対象書類"] is None or need["重大度"] is None:
        raise ValueError("見出し行に「観点名」「対象書類」「重大度」が必要です（「説明」は空でも可）。見出しを直して選び直してください。")
    if len(grid) - 1 > MAX_ROWS:
        raise ValueError(f"行が多すぎます（{MAX_ROWS} 行まで）。ファイルを分けて取り込んでください。")
    existing = {(r["scope"], r["name"]) for r in load()}
    seen, items, issues = set(), [], []
    for n, row in enumerate(grid[1:], start=2):
        def cell(k):
            i = need[k]
            return row[i].strip() if i is not None and i < len(row) else ""
        nm, scope_t, sev_t = cell("観点名"), cell("対象書類"), cell("重大度")
        if not nm:
            issues.append({"row": n, "result": "error", "text": "観点名が空です。観点名を入れてください。"})
            continue
        if sev_t not in SEV_JA:
            issues.append({"row": n, "result": "error",
                           "text": "重大度が空です。「重大」「主要」「軽微」のいずれかを入れてください。" if not sev_t
                           else f"重大度「{sev_t[:20]}」を読み取れません。「重大」「主要」「軽微」のいずれかに直してください。"})
            continue
        if scope_t not in SCOPE_JA:
            issues.append({"row": n, "result": "error",
                           "text": "対象書類が読み取れません。「要件定義書」「基本設計書」「対検査」のいずれかに直してください。"})
            continue
        sc = SCOPE_JA[scope_t]
        if (sc, nm) in seen:
            issues.append({"row": n, "result": "error", "text": f"「{nm[:30]}」がファイル内で重複しています。どちらかを消してください。"})
            continue
        seen.add((sc, nm))
        item = {"scope": sc, "name": nm[:100], "desc": re.sub(r"\s+", " ", cell("説明"))[:300], "severity": SEV_JA[sev_t]}
        if (sc, nm) in existing:
            issues.append({"row": n, "result": "dup", "text": f"「{nm[:30]}」は既にあります。上書きするか、取り込まないかを選べます。"})
        items.append(item)
    errors = sum(1 for i in issues if i["result"] == "error")
    dups = sum(1 for i in issues if i["result"] == "dup")
    return {"items": items, "issues": issues, "summary": {"rows": len(items), "dup": dups, "error": errors}}


def commit(name: str, data: bytes, overwrite: bool) -> dict:
    res = parse(name, data)
    rows = load()
    idx = {(r["scope"], r["name"]): r for r in rows}
    added = updated = skipped = 0
    for it in res["items"]:
        cur = idx.get((it["scope"], it["name"]))
        if cur:
            if overwrite:
                cur.update(desc=it["desc"], severity=it["severity"], source="import")
                updated += 1
            else:
                skipped += 1
            continue
        key = zlib.crc32(f"{it['scope']}|{it['name']}".encode("utf-8"))
        row = {"id": f"imp-{key:08x}", "scope": it["scope"], "name": it["name"], "desc": it["desc"],
               "kind": "human", "severity": it["severity"], "enabled": True, "source": "import"}
        rows.append(row)
        idx[(it["scope"], it["name"])] = row
        added += 1
    settings._write("_viewpoints.json", rows)
    return {"added": added, "updated": updated, "skipped": skipped, "errors": res["summary"]["error"]}
