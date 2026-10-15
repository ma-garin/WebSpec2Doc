"""R3 の取り込み拡張（FR-06-016〜018）。ingest.ingest を包み、表の行・非表示の印・暗号化の判別を足す。

ingest.py は変更しない。契約: docs/sdlc/20_design/WS2D-DD-006_詳細設計書_System06_R3.md 5 章。
"""
from __future__ import annotations

import logging
from pathlib import Path

from web.services.docverify import ingest

log = logging.getLogger(__name__)
_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ENCRYPTED = "この文書はパスワードで保護されています。パスワードを外してから取り込んでください。"


def is_encrypted_office(data: bytes, suffix: str) -> bool:
    """暗号化 docx/xlsx は zip ではなく OLE 複合文書（EncryptedPackage）になる（FR-06-018）。"""
    return suffix in (".docx", ".xlsx") and data[:8] == _OLE


def _is_header(values: list) -> bool:
    texts = [v for v in values if v not in (None, "")]
    return len(texts) >= 2 and all(isinstance(v, str) and len(v.strip()) <= 30 for v in texts)


def _row_text(headers: list[str], values: list) -> str:
    pairs = [f"{h}: {'' if v is None else str(v).strip()}" for h, v in zip(headers, values) if h and v not in (None, "")]
    return "／".join(pairs)


def _sheet_hidden(ws) -> tuple[bool, set[str], set[int]]:
    from openpyxl.utils import column_index_from_string, get_column_letter
    cols: set[str] = set()
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            lo = dim.min or column_index_from_string(key)
            hi = dim.max or lo
            cols.update(get_column_letter(i) for i in range(lo, hi + 1))
    rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}
    return ws.sheet_state != "visible", cols, rows


def _xlsx_sheet(ws, start: int) -> tuple[list[dict], list[dict]]:
    """1 シートを断片に。1 行目が見出しなら各行を「見出し: 値／…」の 1 断片にする（FR-06-016）。"""
    sheet_hidden, hcols, hrows = _sheet_hidden(ws)
    rows = [list(r) for r in ws.iter_rows()]
    if not rows:
        return [], []
    head = [c.value for c in rows[0]]
    passages, tables = [], []
    if _is_header(head) and len(rows) > 1:
        headers = [str(h).strip() if h is not None else "" for h in head]
        tables.append({"sheet": ws.title, "headers": [h for h in headers if h], "location": f"{ws.title} 1行目",
                       "passage_id": f"P{start + 1}"})
        for r in rows[1:]:
            text = _row_text(headers, [c.value for c in r])
            if not text:
                continue
            rno = r[0].row
            hidden = sheet_hidden or rno in hrows or any(c.column_letter in hcols and c.value not in (None, "") for c in r)
            passages.append({"text": text, "location": f"{ws.title} {rno}行目", "hidden": hidden})
        return passages, tables
    for r in rows:
        for c in r:
            if c.value in (None, ""):
                continue
            hidden = sheet_hidden or c.row in hrows or c.column_letter in hcols
            passages.append({"text": str(c.value).strip(), "location": f"{ws.title} {c.coordinate}", "hidden": hidden})
    return passages, tables


def _xlsx(path: Path) -> tuple[list[dict], list[dict]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    passages: list[dict] = []
    tables: list[dict] = []
    for ws in wb.worksheets:
        ps, ts = _xlsx_sheet(ws, len(passages))
        passages.extend(ps)
        tables.extend(ts)
    return passages, tables


def _docx_tables(path: Path) -> tuple[list[dict], list[dict], set[str]]:
    """docx の表で 1 行目が見出しのものを行単位の断片に。置き換えたセルの文も返す。"""
    import docx
    document = docx.Document(str(path))
    passages, tables, cells = [], [], set()
    for t_no, table in enumerate(document.tables, start=1):
        grid = [[c.text.strip() for c in row.cells] for row in table.rows]
        if len(grid) < 2 or not _is_header(grid[0]):
            continue
        headers = grid[0]
        tables.append({"sheet": f"表{t_no}", "headers": [h for h in headers if h], "location": f"表{t_no} 1行目", "passage_id": None})
        cells.update(x for row in grid for x in row if x)
        for r_no, row in enumerate(grid[1:], start=2):
            text = _row_text(headers, row)
            if text:
                passages.append({"text": text, "location": f"表{t_no} {r_no}行目", "hidden": False})
    return passages, tables, cells


def _renumber(passages: list[dict]) -> list[dict]:
    return [{"id": f"P{i}", **{k: v for k, v in p.items() if k != "id"}} for i, p in enumerate(passages, start=1)]


def ingest_r3(filename: str, data: bytes, work_dir: Path) -> tuple[dict, list[dict]]:
    """ingest.ingest と同じ検査をしたうえで、表の行・非表示の印を付けた断片を返す。document に tables を足す。"""
    suffix = Path(filename or "").suffix.lower()
    if is_encrypted_office(data or b"", suffix):
        raise ingest.IngestError(ENCRYPTED)
    document, passages = ingest.ingest(filename, data, work_dir)
    path = Path(work_dir) / f"source{suffix}"
    tables: list[dict] = []
    try:
        if suffix == ".xlsx":
            new, tables = _xlsx(path)
            passages = _renumber(new) or passages
        elif suffix == ".docx":
            rows, tables, cells = _docx_tables(path)
            if rows:
                passages = _renumber([*[p for p in passages if p.get("text", "").strip() not in cells], *rows])
    except Exception:
        log.warning("表・非表示の読み取りに失敗（通常の断片で続ける）", exc_info=True)
        tables = []
    tables = [{**t, "passage_id": t.get("passage_id") or (passages[0]["id"] if passages else None)} for t in tables]
    hidden = sum(1 for p in passages if p.get("hidden"))
    return {**document, "passages": len(passages), "tables": tables, "hidden_passages": hidden}, passages
