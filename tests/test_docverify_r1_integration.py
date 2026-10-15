"""System 06 R1 結合テスト（Flask テストクライアント。WS2D-DD-006 12.2 章の API 側）。

画面側（帯・モーダル・キーボード）は Playwright MCP で別に確かめる。ここでは API の応答だけを見る。
"""
from __future__ import annotations

import io
import json
import re
from decimal import ROUND_HALF_UP, Decimal

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XlImage
from PIL import Image as PILImage

import app as appmod
from web import config
from web.services.docverify import home as home_mod
from web.services.docverify import metrics

ORIGIN = {"Origin": "http://localhost"}
M04 = "読めなかった要素は数えていません。文書を取り込み直すと数えます。"
M13 = "— 文書を取り込むと表示されます。"
M20 = "文書はまだ取り込まれていません。"
M40 = "ホームを読めませんでした。時間をおいて再読み込みしてください。"
M41 = "報告を作れませんでした。もう一度出力してください。"
LEAK = re.compile(r"Traceback|Exception|Error\b|\.py\b|/Users/|/tmp/|/private/|File \"")


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setenv("FLASK_TESTING", "1")
    monkeypatch.setenv("WEBSPEC2DOC_ALLOW_LOCAL", "1")
    monkeypatch.setenv("WEBSPEC2DOC_AUTH_MODE", "off")
    return appmod.app.test_client()


def _xlsx_bytes(tmp_path, *, images=3, comments=2, lines=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, text in enumerate(lines or [f"システムは、要件 {i} を満たすこと。" for i in range(1, 6)], start=1):
        ws.cell(row=r, column=1, value=text)
    for i in range(comments):
        ws.cell(row=i + 1, column=1).comment = Comment(f"コメント {i}", "作成者")
    png = tmp_path / "img.png"
    PILImage.new("RGB", (4, 4), (0, 128, 0)).save(png)
    for i in range(images):
        ws.add_image(XlImage(str(png)), f"D{i * 5 + 1}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _create(client, name="R1 結合") -> str:
    r = client.post("/api/docverify/inspections", json={"name": name, "goal": "確認"}, headers=ORIGIN)
    assert r.status_code == 201
    return r.get_json()["inspection"]["id"]


def _upload(client, iid, data: bytes, filename="要件.xlsx"):
    return client.post(f"/api/docverify/inspections/{iid}/document", headers=ORIGIN,
                       data={"file": (io.BytesIO(data), filename)}, content_type="multipart/form-data")


def _rows(rows) -> dict:
    return {r["label"]: r["value"] for r in rows}


def _round(values) -> int:
    """四捨五入（ROUND_HALF_UP）の平均。"""
    return int((Decimal(sum(values)) / Decimal(len(values))).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def test_upload_reports_unread_quality_scope(client, tmp_path):
    """IT-06R1-01 FR-06-015・FR-06-100・FR-06-075 G 画像 3・コメント 2 の xlsx W 取り込む
    T unread_lines に「画像 3・コメント 2 は読んでいません。」、document.unread が保存され、quality・scope が揃う。"""
    iid = _create(client)
    r = _upload(client, iid, _xlsx_bytes(tmp_path))
    assert r.status_code == 200
    body = r.get_json()
    unread = body["inspection"]["document"]["unread"]
    assert (unread["images"], unread["comments"], unread["scanned"]) == (3, 2, True)
    s = body["summary"]
    assert s["unread_lines"] == ["画像 3・コメント 2 は読んでいません。"]
    q = s["quality"]
    assert q["passages"] == len(body["inspection"]["passages"]) > 0
    assert isinstance(q["overall"], int) and 0 <= q["overall"] <= 100
    assert len(q["note"]) == 3 and all(item["label"] for item in q["items"].values())
    chk, unc = _rows(s["scope"]["checked"]), _rows(s["scope"]["unchecked"])
    assert chk["文書"] == "要件.xlsx（版 1・xlsx）"
    assert chk["読めなかった要素"] == "画像 3・コメント 2 は読んでいません。"
    assert unc["網羅性"].startswith("未検証。") and unc["規則の精度"].startswith("未測定。")
    again = client.get(f"/api/docverify/inspections/{iid}").get_json()
    assert again["inspection"]["document"]["unread"] == unread and "updated_at" in again["inspection"]


def test_second_upload_freezes_first_version(client, tmp_path):
    """IT-06R1-02 FR-06-104 G 版 1 を取り込み指摘を全て却下 W 版 2 を取り込む
    T versions[0].quality.final 真、versions[1].quality.final 偽、history に 2 版・版 2 だけ current。"""
    iid = _create(client)
    lines = ["システムは、適切に処理すること。", "システムは、速やかに応答すること。", "システムは、ログを保存すること。"]
    first = _upload(client, iid, _xlsx_bytes(tmp_path, images=0, comments=0, lines=lines)).get_json()
    for f in first["inspection"]["findings"]:
        r = client.post(f"/api/docverify/inspections/{iid}/findings/{f['id']}", json={"status": "rejected"}, headers=ORIGIN)
        assert r.status_code == 200
    before = client.get(f"/api/docverify/inspections/{iid}").get_json()["summary"]["quality"]["overall"]
    second = _upload(client, iid, _xlsx_bytes(tmp_path, images=0, comments=0, lines=lines + ["システムは、帳票を出すこと。"]))
    body = second.get_json()
    versions = body["inspection"]["versions"]
    assert len(versions) == 2
    assert versions[0]["quality"]["final"] is True and versions[0]["quality"]["overall"] == before
    assert versions[1]["quality"]["final"] is False
    hist = body["summary"]["history"]
    assert [h["version"] for h in hist] == [1, 2] and [h["current"] for h in hist] == [False, True]
    assert hist[0]["overall"] == before and hist[1]["overall"] == body["summary"]["quality"]["overall"]


def test_no_document_shows_dash_not_zero(client):
    """IT-06R1-05 FR-06-102・FR-06-075 G 文書なし W 検査を開く T quality.overall None（0 でない）、scope は M-20・M-13、unread_lines 空。"""
    iid = _create(client)
    s = client.get(f"/api/docverify/inspections/{iid}").get_json()["summary"]
    assert s["quality"]["overall"] is None
    chk = _rows(s["scope"]["checked"])
    assert chk["文書"] == M20 and chk["品質スコア"] == M13
    assert s["unread_lines"] == []


def test_home_average_and_items(client, tmp_path):
    """IT-06R1-04 FR-06-101・FR-06-102 G 進行中 2 件（片方は指摘を全て却下） W GET /api/docverify/home
    T quality.overall は 2 件の総合の四捨五入平均、特性別も同じ平均、n_reviews 2、active 2 行。"""
    lines = ["システムは、適切に処理すること。", "システムは、ログを保存すること。"]
    ids = [_create(client, "ホーム A"), _create(client, "ホーム B")]
    for iid in ids:
        assert _upload(client, iid, _xlsx_bytes(tmp_path, images=0, comments=0, lines=lines)).status_code == 200
    for f in client.get(f"/api/docverify/inspections/{ids[0]}").get_json()["inspection"]["findings"]:
        client.post(f"/api/docverify/inspections/{ids[0]}/findings/{f['id']}", json={"status": "rejected"}, headers=ORIGIN)
    qs = [client.get(f"/api/docverify/inspections/{i}").get_json()["summary"]["quality"] for i in ids]
    assert qs[0]["overall"] == 100 and qs[1]["overall"] < 100
    r = client.get("/api/docverify/home")
    assert r.status_code == 200
    h = r.get_json()
    assert h["quality"]["overall"] == _round([q["overall"] for q in qs])
    assert (h["quality"]["n_reviews"], h["quality"]["n_scored"]) == (2, 2)
    for key, item in h["quality"]["items"].items():
        scores = [q["items"][key]["score"] for q in qs if q["items"][key]["score"] is not None]
        assert item["score"] == (_round(scores) if scores else None)
        assert item["label"] == qs[0]["items"][key]["label"]
    assert {a["review_id"] for a in h["active"]} == set(ids)
    assert {a["quality_overall"] for a in h["active"]} == {q["overall"] for q in qs}
    assert {"mine", "active", "open_findings", "open_critical", "awaiting_answers"} <= set(h["counts"])
    assert h["next_scope"] in ("mine", "active") and len(h["next_actions"]) <= 5


def test_report_contains_scope_sheet(client, tmp_path):
    """IT-06R1-07 FR-06-075・NFR-06-08 G 取り込み済み W 報告を出力 T 「検査した範囲」シート、網羅性「未検証」、規則の精度「未測定」。"""
    iid = _create(client)
    _upload(client, iid, _xlsx_bytes(tmp_path))
    r = client.get(f"/api/docverify/inspections/{iid}/report.xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.data))["検査した範囲"]
    rows = {row[1]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["網羅性"][0] == "検査していないこと" and rows["網羅性"][2].startswith("未検証")
    assert rows["規則の精度"][2].startswith("未測定")


def test_report_failure_is_user_readable(client, tmp_path, monkeypatch):
    """IT-06R1-08 SCR-06-22・NFR-06-16 G to_xlsx が失敗 W 報告を出力 T 500 と M-41、例外名・パスが文に出ない。"""
    iid = _create(client)

    def boom(_insp):
        raise RuntimeError(f"disk full at {tmp_path}/secret.py")

    monkeypatch.setattr(metrics, "to_xlsx", boom)
    r = client.get(f"/api/docverify/inspections/{iid}/report.xlsx")
    assert r.status_code == 500
    assert r.get_json() == {"error": M41}
    assert not LEAK.search(r.get_data(as_text=True)) and str(tmp_path) not in r.get_data(as_text=True)


def test_home_failure_is_user_readable(client, monkeypatch):
    """SCR-06-01・NFR-06-16（M-40） G 集計が失敗 W ホームを開く T 500 と M-40、例外名・パスが文に出ない。"""
    def boom(*_a, **_k):
        raise KeyError("/Users/x/data/inspection.json")

    monkeypatch.setattr(home_mod, "build", boom)
    r = client.get("/api/docverify/home")
    assert r.status_code == 500
    assert r.get_json()["error"] == M40
    assert not LEAK.search(r.get_data(as_text=True))


def test_broken_docx_error_is_user_readable(client):
    """IT-06R1-14 FR-06-014・NFR-06-16 G 破損した docx W 取り込む T 400、文にパス・例外名が無く、文書は取り込まれない。"""
    iid = _create(client)
    r = _upload(client, iid, b"PK\x03\x04 broken", filename="壊れ.docx")
    assert r.status_code == 400
    err = r.get_json()["error"]
    assert err and not LEAK.search(err)
    assert client.get(f"/api/docverify/inspections/{iid}").get_json()["inspection"]["document"] is None


def test_old_json_compat(client, tmp_path):
    """IT-06R1-15 互換 G unread・quality・updated_at の無い旧 JSON W 検査とホームの API T 200、unread_lines は M-04、
    history は版 1 だけで current 真（現在値）、ホームの updated_at は created_at。"""
    iid = _create(client)
    _upload(client, iid, _xlsx_bytes(tmp_path))
    [path] = list(tmp_path.rglob(f"{iid}/inspection.json"))
    data = json.loads(path.read_text(encoding="utf-8"))
    data.pop("updated_at", None)
    data["document"].pop("unread", None)
    data["versions"] = [{k: v for k, v in ver.items() if k != "quality"} for ver in data["versions"]]
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    r = client.get(f"/api/docverify/inspections/{iid}")
    assert r.status_code == 200
    s = r.get_json()["summary"]
    assert s["unread_lines"] == [M04]
    assert len(s["history"]) == 1 and s["history"][0]["current"] is True
    h = client.get("/api/docverify/home")
    assert h.status_code == 200
    [act] = [a for a in h.get_json()["active"] if a["review_id"] == iid]
    assert act["updated_at"] == data["created_at"]
