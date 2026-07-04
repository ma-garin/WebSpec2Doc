"""Doc Fusion（多形式文書取り込み＋実測突合）のユニットテスト。"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from analyzer.html_analyzer import analyze_pages
from crawler.page_crawler import FieldData, FormData, PageData, SourceEvidence
from generator.fusion_reporter import (
    FUSION_JSON_NAME,
    FUSION_MD_NAME,
    save_fusion_outputs,
)
from ingest.data_reader import read_structured_data
from ingest.loader import load_reference_documents
from ingest.matcher import fuse
from ingest.models import DocumentBundle, DocumentedField, DocumentedScreen
from ingest.office_reader import read_docx, read_pptx_lines
from ingest.tables import (
    parse_max_length,
    parse_required,
    screens_from_lines,
)
from ingest.text_reader import read_markdown

# ---------- 共通フィクスチャ ----------


def _field(
    name: str,
    required: bool = False,
    maxlength: int | None = None,
    aria_label: str = "",
) -> FieldData:
    return FieldData(
        field_type="text",
        name=name,
        placeholder="",
        required=required,
        maxlength=maxlength,
        aria_label=aria_label,
        evidence=SourceEvidence(selector=f"[name='{name}']"),
    )


def _page(url: str, title: str, fields: tuple[FieldData, ...] = ()) -> PageData:
    forms = (FormData(action="/submit", method="post", fields=fields),) if fields else ()
    return PageData(
        url=url,
        title=title,
        headings=(title,),
        links=(),
        forms=forms,
        screenshot_path=None,
    )


# ---------- 値パース ----------


class TestValueParsing:
    def test_required_marks(self) -> None:
        assert parse_required("○") is True
        assert parse_required("必須") is True
        assert parse_required("Yes") is True
        assert parse_required("×") is False
        assert parse_required("任意") is False
        assert parse_required("") is None
        assert parse_required("わからない") is None

    def test_max_length_variants(self) -> None:
        assert parse_max_length("20") == 20
        assert parse_max_length("全角20") == 20
        assert parse_max_length("100桁") == 100

    def test_max_length_comma_separated(self) -> None:
        """3桁区切りのカンマ（"9,999"）は先頭桁のみでなく全体を数値として扱う。"""
        assert parse_max_length("9,999") == 9999
        assert parse_max_length("上限9,999円") == 9999
        assert parse_max_length("1,000,000") == 1000000
        assert parse_max_length("なし") is None


# ---------- Excel ----------


class TestExcelIngest:
    def _make_book(self, path: Path) -> None:
        book = Workbook()
        screens = book.active
        screens.title = "画面一覧"
        # ヘッダが先頭行でないケースを再現する
        screens.append(["画面設計書"])
        screens.append([])
        screens.append(["画面ID", "画面名", "URL", "備考"])
        screens.append(["GA-010", "ログイン画面", "/login.html", ""])
        screens.append(["GA-020", "お問い合わせ画面", "/contact.html", ""])
        fields = book.create_sheet("項目定義")
        fields.append(["画面名", "項目名", "物理名", "型", "必須", "桁数"])
        fields.append(["ログイン画面", "メールアドレス", "email", "文字列", "○", "100"])
        fields.append(["ログイン画面", "パスワード", "password", "文字列", "○", "20"])
        fields.append(["ログイン画面", "社員番号", "employee_no", "数値", "○", "8"])
        book.save(path)

    def test_excel_screens_and_fields(self, tmp_path: Path) -> None:
        book_path = tmp_path / "設計書.xlsx"
        self._make_book(book_path)
        bundle = load_reference_documents([book_path])
        assert [s.name for s in bundle.screens] == ["ログイン画面", "お問い合わせ画面"]
        assert bundle.screens[0].screen_id == "GA-010"
        assert bundle.screens[0].url_hint == "/login.html"
        email = next(f for f in bundle.fields if f.name == "メールアドレス")
        assert email.physical_name == "email"
        assert email.required is True
        assert email.max_length == 100
        assert email.evidence is not None
        assert "項目定義" in email.evidence.location


# ---------- Markdown ----------


class TestMarkdownIngest:
    def test_markdown_table_and_headings(self, tmp_path: Path) -> None:
        md = tmp_path / "spec.md"
        md.write_text(
            "\n".join(
                [
                    "# 基本設計書",
                    "## 3.1 ダッシュボード画面",
                    "",
                    "| 項目名 | 物理名 | 必須 | 桁数 |",
                    "|---|---|---|---|",
                    "| 氏名 | name | ○ | 40 |",
                    "| 部署 | dept | × | 20 |",
                ]
            ),
            encoding="utf-8",
        )
        tables, headings = read_markdown(md)
        assert len(tables) == 1
        assert len(tables[0].rows) == 2
        assert any("ダッシュボード画面" in text for _, text in headings)

        bundle = load_reference_documents([md])
        assert any(s.name == "ダッシュボード画面" for s in bundle.screens)
        dept = next(f for f in bundle.fields if f.name == "部署")
        assert dept.required is False
        assert dept.max_length == 20


# ---------- YAML / JSON ----------


class TestStructuredDataIngest:
    def test_yaml_nested_screens(self, tmp_path: Path) -> None:
        doc = tmp_path / "screens.yaml"
        doc.write_text(
            "\n".join(
                [
                    "画面一覧:",
                    "  - 画面名: ログイン画面",
                    "    url: /login.html",
                    "    項目:",
                    "      - 項目名: メールアドレス",
                    "        必須: true",
                    "        桁数: 100",
                ]
            ),
            encoding="utf-8",
        )
        screens, fields = read_structured_data(doc)
        assert screens[0].name == "ログイン画面"
        assert fields[0].screen_name == "ログイン画面"
        assert fields[0].required is True
        assert fields[0].max_length == 100

    def test_json_flat_screens(self, tmp_path: Path) -> None:
        doc = tmp_path / "screens.json"
        doc.write_text(
            json.dumps({"screens": [{"name": "設定画面", "path": "/settings"}]}),
            encoding="utf-8",
        )
        screens, _fields = read_structured_data(doc)
        assert screens[0].name == "設定画面"
        assert screens[0].url_hint == "/settings"
        assert screens[0].evidence is not None
        assert screens[0].evidence.location.startswith("$")


# ---------- Office（docx / pptx） ----------

_DOCX_XML = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
<w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>3.1 注文画面</w:t></w:r></w:p>
<w:tbl>
<w:tr><w:tc><w:p><w:r><w:t>項目名</w:t></w:r></w:p></w:tc>
<w:tc><w:p><w:r><w:t>必須</w:t></w:r></w:p></w:tc>
<w:tc><w:p><w:r><w:t>桁数</w:t></w:r></w:p></w:tc></w:tr>
<w:tr><w:tc><w:p><w:r><w:t>数量</w:t></w:r></w:p></w:tc>
<w:tc><w:p><w:r><w:t>○</w:t></w:r></w:p></w:tc>
<w:tc><w:p><w:r><w:t>3</w:t></w:r></w:p></w:tc></w:tr>
</w:tbl>
</w:body>
</w:document>
"""

_PPTX_SLIDE_XML = """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<p:cSld><p:spTree>
<p:sp><p:txBody><a:p><a:r><a:t>会員登録画面</a:t></a:r></a:p></p:txBody></p:sp>
</p:spTree></p:cSld>
</p:sld>
"""


class TestOfficeIngest:
    def test_docx_table_and_heading(self, tmp_path: Path) -> None:
        docx_path = tmp_path / "設計書.docx"
        with zipfile.ZipFile(docx_path, "w") as archive:
            archive.writestr("word/document.xml", _DOCX_XML)
        tables, headings = read_docx(docx_path)
        assert len(tables) == 1
        assert headings and "注文画面" in headings[0][1]

        bundle = load_reference_documents([docx_path])
        assert any(s.name == "注文画面" for s in bundle.screens)
        quantity = next(f for f in bundle.fields if f.name == "数量")
        assert quantity.required is True
        assert quantity.max_length == 3

    def test_pptx_lines(self, tmp_path: Path) -> None:
        pptx_path = tmp_path / "発表.pptx"
        with zipfile.ZipFile(pptx_path, "w") as archive:
            archive.writestr("ppt/slides/slide1.xml", _PPTX_SLIDE_XML)
        lines = read_pptx_lines(pptx_path)
        assert lines == [("slide 1", "会員登録画面")]
        bundle = load_reference_documents([pptx_path])
        assert any(s.name == "会員登録画面" for s in bundle.screens)


# ---------- テキスト行からの画面候補抽出 ----------


class TestScreenLineExtraction:
    def test_extracts_screen_names_from_short_lines(self) -> None:
        lines = [
            ("line 1", "3.2 ログイン画面"),
            ("line 2", "これは長い説明文で、画面という語を含むが対象外にしたい文章です。" * 2),
            ("line 3", "3.2 ログイン画面"),  # 重複
        ]
        screens = screens_from_lines(lines, "spec.txt")
        assert [s.name for s in screens] == ["ログイン画面"]
        assert screens[0].evidence is not None
        assert screens[0].evidence.location == "line 1"


# ---------- ローダのエラー処理 ----------


class TestLoaderErrors:
    def test_missing_file(self, tmp_path: Path) -> None:
        with pytest.raises(FileNotFoundError):
            load_reference_documents([tmp_path / "ない.xlsx"])

    def test_unsupported_suffix(self, tmp_path: Path) -> None:
        target = tmp_path / "spec.csv"
        target.write_text("a,b", encoding="utf-8")
        with pytest.raises(ValueError, match="未対応の文書形式"):
            load_reference_documents([target])

    def test_legacy_suffix_guides_conversion(self, tmp_path: Path) -> None:
        target = tmp_path / "old.xls"
        target.write_bytes(b"\x00")
        with pytest.raises(ValueError, match="変換してから"):
            load_reference_documents([target])


# ---------- 突合（matcher） ----------


def _bundle() -> DocumentBundle:
    login = DocumentedScreen(screen_id="GA-010", name="ログイン画面", url_hint="/login.html")
    legacy = DocumentedScreen(screen_id="GA-090", name="帳票出力画面", url_hint="/report.html")
    return DocumentBundle(
        screens=(login, legacy),
        fields=(
            DocumentedField(
                name="メールアドレス",
                physical_name="email",
                screen_name="ログイン画面",
                required=True,
                max_length=100,
            ),
            DocumentedField(
                name="パスワード",
                physical_name="password",
                screen_name="ログイン画面",
                required=True,
                max_length=20,
            ),
            DocumentedField(
                name="社員番号",
                physical_name="employee_no",
                screen_name="ログイン画面",
                required=True,
            ),
        ),
        source_files=("設計書.xlsx",),
    )


class TestFusion:
    def _fuse(self):  # noqa: ANN202
        pages = analyze_pages(
            [
                _page(
                    "https://example.com/login.html",
                    "ログイン",
                    fields=(
                        _field("email", required=True, maxlength=100),
                        # 文書では必須 20 桁 → 実測では任意 30 桁（矛盾を2種類再現）
                        _field("password", required=False, maxlength=30),
                        # 文書に記載のない実測項目（文書化漏れ）
                        _field("otp_code", required=False),
                    ),
                ),
                _page("https://example.com/dashboard.html", "ダッシュボード"),
            ]
        )
        return fuse(pages, _bundle())

    def test_screen_matching_by_url(self) -> None:
        result = self._fuse()
        match = next(m for m in result.screen_matches if m.page_url.endswith("login.html"))
        assert match.screen.name == "ログイン画面"
        assert match.method == "url"
        assert result.official_names[match.page_id] == "ログイン画面"

    def test_doc_only_and_crawl_only_screens(self) -> None:
        result = self._fuse()
        assert [s.name for s in result.doc_only_screens] == ["帳票出力画面"]
        assert len(result.crawl_only_page_ids) == 1  # ダッシュボード

    def test_field_gaps_three_kinds(self) -> None:
        result = self._fuse()
        kinds = {(g.kind, g.field_name) for g in result.field_gaps}
        # 矛盾: パスワードの必須区分と桁数
        mismatch_details = [g.detail for g in result.field_gaps if g.kind == "mismatch"]
        assert any("必須区分が矛盾" in d for d in mismatch_details)
        assert any("桁数が矛盾" in d for d in mismatch_details)
        # 文書のみ: 社員番号 / 実測のみ: otp_code
        assert ("doc_only", "社員番号") in kinds
        assert ("crawl_only", "otp_code") in kinds

    def test_matched_fields_do_not_appear_as_gaps(self) -> None:
        result = self._fuse()
        gap_names = {g.field_name for g in result.field_gaps if g.kind != "mismatch"}
        assert "メールアドレス" not in gap_names
        assert "email" not in gap_names


# ---------- レポート出力と用語注入 ----------


class TestFusionOutputs:
    def test_save_outputs_writes_json_and_md(self, tmp_path: Path) -> None:
        pages = analyze_pages(
            [
                _page(
                    "https://example.com/login.html",
                    "ログイン",
                    fields=(_field("password", required=False, maxlength=30),),
                )
            ]
        )
        bundle = _bundle()
        result = fuse(pages, bundle)
        save_fusion_outputs(result, bundle, tmp_path)
        data = json.loads((tmp_path / FUSION_JSON_NAME).read_text(encoding="utf-8"))
        assert data["meta"]["source_files"] == ["設計書.xlsx"]
        assert data["screen_matches"][0]["official_name"] == "ログイン画面"
        markdown = (tmp_path / FUSION_MD_NAME).read_text(encoding="utf-8")
        assert "文書×実測 突合レポート" in markdown
        assert "矛盾" in markdown

    def test_official_name_injected_into_report_json(self) -> None:
        from generator.json_reporter import generate_json_report
        from graph.transition_graph import build_graph

        pages = analyze_pages([_page("https://example.com/login.html", "ログイン")])
        graph = build_graph(pages)
        report = json.loads(
            generate_json_report(
                pages,
                graph,
                "https://example.com/login.html",
                official_names={"P001": "ログイン画面"},
            )
        )
        assert report["screens"][0]["official_name"] == "ログイン画面"

    def test_official_name_absent_without_reference_doc(self) -> None:
        from generator.json_reporter import generate_json_report
        from graph.transition_graph import build_graph

        pages = analyze_pages([_page("https://example.com/login.html", "ログイン")])
        graph = build_graph(pages)
        report = json.loads(generate_json_report(pages, graph, "https://example.com/login.html"))
        assert "official_name" not in report["screens"][0]
