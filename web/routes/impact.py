"""影響度分析（System 08）。処理の中身は web/services/impact/ にある。"""
from __future__ import annotations

import csv
import io
import logging

from flask import Blueprint, Response, jsonify, render_template, request

from web.services.impact import analyze, llm, store

bp = Blueprint("impact", __name__)
log = logging.getLogger(__name__)


def _err(message: str, status: int):
    return jsonify(error=message), status


def _load(jid: str):
    try:
        return store.load(jid), None
    except KeyError:
        return None, _err("分析が見つかりません。一覧から選び直してください。", 404)
    except (ValueError, OSError):
        log.exception("分析の読み込みに失敗 %s", jid)
        return None, _err("分析を読めませんでした。時間をおいて再読み込みしてください。", 500)


@bp.route("/impact")
def page() -> str:
    return render_template("impact_analysis.html")


@bp.route("/api/impact/jobs", methods=["GET"])
def list_jobs():
    try:
        return jsonify(jobs=store.list_all())
    except OSError:
        log.exception("分析一覧の取得に失敗")
        return _err("分析の一覧を読めませんでした。時間をおいて再読み込みしてください。", 500)


@bp.route("/api/impact/jobs", methods=["POST"])
def create_job():
    name = str((request.get_json(silent=True) or {}).get("name") or "").strip()
    if not name or len(name) > 100:
        return _err("分析の名前を 1〜100 文字で入れてください。", 400)
    return jsonify(job=store.save(store.new(name))), 201


@bp.route("/api/impact/jobs/<jid>", methods=["GET", "PATCH", "DELETE"])
def job(jid: str):
    job, err = _load(jid)
    if err:
        return err
    if request.method == "GET":
        return jsonify(job=job)
    if request.method == "DELETE":
        store.delete(job)
        return jsonify(ok=True)
    body = request.get_json(silent=True) or {}
    if "name" in body:
        name = str(body["name"] or "").strip()
        if not name or len(name) > 100:
            return _err("分析の名前を 1〜100 文字で入れてください。", 400)
        job = {**job, "name": name}
    if "order" in body:
        job = {**job, "order": str(body["order"] or "")[:2000]}
    return jsonify(job=store.save(job))


@bp.route("/api/impact/jobs/<jid>/docs", methods=["POST"])
def upload(jid: str):
    job, err = _load(jid)
    if err:
        return err
    uploads = request.files.getlist("files")
    if not uploads:
        return _err("文書が選ばれていません。設計書か画面仕様書を選んでください。", 400)
    notes = []
    for f in uploads:
        try:
            job = store.add_doc(job, f.filename or "", f.read())
        except store.StoreError as exc:
            notes.append(str(exc))
    job = store.save(job)
    if notes and len(notes) == len(uploads):
        return _err(" / ".join(notes), 400)
    return jsonify(job=job, notes=notes)


@bp.route("/api/impact/jobs/<jid>/docs/<path:name>", methods=["PATCH", "DELETE"])
def change_doc(jid: str, name: str):
    job, err = _load(jid)
    if err:
        return err
    try:
        if request.method == "DELETE":
            return jsonify(job=store.save(store.remove_doc(job, name)))
    except KeyError:
        return _err("その文書は取り込まれていません。一覧を再読み込みしてください。", 404)
    body = request.get_json(silent=True) or {}
    if not any(d["name"] == name for d in job["docs"]):
        return _err("その文書は取り込まれていません。一覧を再読み込みしてください。", 404)
    if ("role" in body and body["role"] not in store.ROLES) or ("store" in body and body["store"] not in store.STORES):
        return _err("選べない値です。一覧から選び直してください。", 400)
    docs = [{**d, **{k: body[k] for k in ("role", "store") if k in body}} if d["name"] == name else d for d in job["docs"]]
    return jsonify(job=store.save({**job, "docs": docs}))


@bp.route("/api/impact/jobs/<jid>/analyze", methods=["POST"])
def run_analysis(jid: str):
    job, err = _load(jid)
    if err:
        return err
    body = request.get_json(silent=True) or {}
    if "order" in body:
        job = {**job, "order": str(body["order"] or "")[:2000]}
    if not job["order"].strip() and not job["docs"]:
        return _err("変更の内容を書くか、変更後の仕様書を取り込んでください。", 400)
    lib_docs = [{**d, "role": "store", "library": True} for d in store.library()["docs"]
                if all(d["name"] != j["name"] for j in job["docs"])]
    lib_names = {d["name"] for d in lib_docs}

    def read(name: str) -> list[dict]:
        return store.library_passages(name) if name in lib_names else store.read_passages(job, name)

    use_llm = body.get("engine", "auto") != "rule"
    try:
        result = analyze.run(job["order"], job["docs"] + lib_docs, read, llm.judge if use_llm else None)
    except (OSError, ValueError):
        log.exception("影響度分析に失敗 %s", jid)
        return _err("分析できませんでした。文書を取り込み直してからもう一度お試しください。", 500)
    result["impacts"] = store.carry_reviews(job.get("impacts") or [], result["impacts"])
    return jsonify(job=store.save({**job, **result, "analyzed_at": store.now()}))


@bp.route("/api/impact/sample", methods=["POST"])
def create_sample():
    """見本の仕様書 4 本で分析を 1 件作り、分析まで済ませる。初めての人が 1 クリックで結果を見られるように。"""
    from pathlib import Path
    sample = Path(__file__).resolve().parents[1] / "services" / "impact" / "sample"
    job = store.new("見本: 電話番号を必須に")
    job = {**job, "order": "会員登録の電話番号を任意から必須にし、桁数を 10〜11 桁にする。"}
    try:
        for path in sorted(sample.glob("*.md")):
            job = store.add_doc(job, path.name, path.read_bytes())
        result = analyze.run(job["order"], job["docs"], lambda n: store.read_passages(job, n), llm.judge)
    except (store.StoreError, OSError, ValueError):
        log.exception("見本の分析に失敗")
        return _err("見本を作れませんでした。時間をおいてもう一度お試しください。", 500)
    return jsonify(job=store.save({**job, **result, "analyzed_at": store.now()})), 201


@bp.route("/api/impact/library", methods=["POST"])
def upload_library():
    uploads = request.files.getlist("files")
    if not uploads:
        return _err("文書が選ばれていません。設計書・画面仕様書・テストケースを選んでください。", 400)
    lib, notes = store.library(), []
    for f in uploads:
        try:
            lib = store.library_put(f.filename or "", f.read())
        except store.StoreError as exc:
            notes.append(str(exc))
    if notes and len(notes) == len(uploads):
        return _err(" / ".join(notes), 400)
    return jsonify(library=lib, notes=notes)


@bp.route("/api/impact/library/<path:name>", methods=["PATCH"])
def change_library_doc(name: str):
    kind = (request.get_json(silent=True) or {}).get("store")
    if kind not in store.STORES:
        return _err("選べない値です。一覧から選び直してください。", 400)
    try:
        return jsonify(library=store.library_set_store(name, kind))
    except KeyError:
        return _err("その文書は文書庫にありません。一覧を再読み込みしてください。", 404)


@bp.route("/api/impact/engine")
def engine():
    """判定に LLM を使えるか（接続先の有無だけを返す。鍵は返さない）。"""
    return jsonify(llm.endpoint_info())


@bp.route("/api/impact/library", methods=["GET"])
def get_library():
    try:
        return jsonify(library=store.library())
    except (OSError, ValueError):
        log.exception("文書庫の読み込みに失敗")
        return _err("文書庫を読めませんでした。時間をおいて再読み込みしてください。", 500)


@bp.route("/api/impact/library/<path:name>", methods=["DELETE"])
def remove_library_doc(name: str):
    try:
        return jsonify(library=store.library_remove(name))
    except KeyError:
        return _err("その文書は文書庫にありません。一覧を再読み込みしてください。", 404)


@bp.route("/api/impact/jobs/<jid>/docs/<path:name>/library", methods=["POST"])
def move_to_library(jid: str, name: str):
    job, err = _load(jid)
    if err:
        return err
    try:
        job, lib = store.library_add(job, name)
    except KeyError:
        return _err("その文書は取り込まれていません。一覧を再読み込みしてください。", 404)
    except store.StoreError as exc:
        return _err(str(exc), 400)
    return jsonify(job=store.save(job), library=lib)


@bp.route("/api/impact/jobs/<jid>/impacts/<iid>", methods=["PATCH"])
def review_impact(jid: str, iid: str):
    job, err = _load(jid)
    if err:
        return err
    body = request.get_json(silent=True) or {}
    pri = body.get("priority")
    try:
        job = store.set_review(job, iid, str(body.get("status") or ""), int(pri) if pri not in (None, "") else None,
                               str(body.get("note") or ""))
    except KeyError:
        return _err("その影響は見つかりません。画面を再読み込みしてください。", 404)
    except (store.StoreError, ValueError, TypeError):
        return _err("選べない値です。一覧から選び直してください。", 400)
    return jsonify(job=store.save(job))


_STATUS = {"accepted": "採用", "rejected": "除外", "": "未判断"}
_HEAD = ["ID", "優先度", "影響度", "判定", "人の判断", "メモ", "影響を受ける箇所", "元の変更", "判断の理由",
         "根拠の文書", "根拠の記述", "確かめるテスト"]


def _export_rows(job: dict) -> list[list]:
    rows = []
    for r in job["impacts"]:
        rv = r.get("review") or {}
        pri = rv.get("priority") or r["priority"]
        rows.append([r["id"], f"P{pri}", r["level_name"], "AI" if r.get("method") == "llm" else "規則",
                     _STATUS.get(rv.get("status", ""), "未判断"), rv.get("note", ""), r["where"], "・".join(r["changes"]),
                     r["why"], f"{r['doc']} {r['location']}".strip(), r["quote"], r["test"]])
    return rows


@bp.route("/api/impact/jobs/<jid>/export.xlsx")
def export_xlsx(jid: str):
    job, err = _load(jid)
    if err:
        return err
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "影響範囲"
    ws.append(_HEAD)
    for row in _export_rows(job):
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFEDD5")
    widths = [6, 7, 8, 6, 9, 24, 28, 9, 48, 28, 48, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    info = wb.create_sheet("変更箇所")
    info.append(["ID", "種類", "場所", "変更前", "変更後"])
    for c in job["changes"]:
        info.append([c["id"], c["kind"], c["location"], c["before"], c["after"]])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''impact.xlsx"})


@bp.route("/api/impact/jobs/<jid>/export.csv")
def export_csv(jid: str):
    job, err = _load(jid)
    if err:
        return err
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(_HEAD)
    w.writerows(_export_rows(job))
    body = "﻿" + buf.getvalue()
    return Response(body, mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''impact.csv"})
