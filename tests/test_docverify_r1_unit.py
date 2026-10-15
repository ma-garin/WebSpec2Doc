"""System 06 R1 単体テスト（WS2D-RD-006 の受入基準と WS2D-DD-006 12.1 章から導出）。

入力は要件定義書・詳細設計書だけ。テストデータは tmp_path に openpyxl / python-docx / PIL / pypdf で作る。
"""
from __future__ import annotations

import copy
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from docx.oxml import parse_xml
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XlImage
from PIL import Image as PILImage

from web import config
from web.services.docverify import checks, home, ingest, metrics, store, unread
from web.vendor.reverse_doc.docs_ingest import readers

# ---- 9 章の文（固定） ----
M01_EX = "画像 3・コメント 2 は読んでいません。"
M02_EX = "非表示シート 1・非表示の行・列 3 は読み取りました（画面では見えない内容です）。"
M03 = "読めなかった要素はありません。"
M04 = "読めなかった要素は数えていません。文書を取り込み直すと数えます。"
M05 = "読めなかった要素を数えられませんでした。画像やコメントが無いか、元の文書で確かめてください。"
M20 = "文書はまだ取り込まれていません。"
M13 = "— 文書を取り込むと表示されます。"
M30 = "未検証。見つかった指摘は検出規則と人の確認の範囲に限られ、ほかに欠陥が無いことは示しません。"
M31 = "未検証。記述が業務として正しいかは判断していません。"
M33 = "未測定。規則ごとの適合率・再現率はまだ測っていません。"

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
URI_WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
NS_WPS = URI_WPS
URI_DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"

ZERO = {"images": 0, "comments": 0, "hidden_sheets": 0, "hidden_cols": 0, "shapes": 0, "scanned": False}


# ---------------------------------------------------------------- 生成ヘルパ
def _png(tmp_path: Path, name: str = "p.png") -> Path:
    p = tmp_path / name
    PILImage.new("RGB", (4, 4), (200, 10, 10)).save(p)
    return p


def _rewrite_zip(path: Path, member: str, fn) -> None:
    src = zipfile.ZipFile(io.BytesIO(path.read_bytes()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info.filename)
            out.writestr(info.filename, fn(data) if info.filename == member else data)
    path.write_bytes(buf.getvalue())


def _xdr_sp(sp_id: int, text: str | None) -> str:
    body = (f"<xdr:txBody><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></xdr:txBody>" if text else "")
    pos = "<xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r}</xdr:row><xdr:rowOff>0</xdr:rowOff>"
    return (f'<xdr:twoCellAnchor xmlns:xdr="{NS_XDR}" xmlns:a="{NS_A}">'
            f"<xdr:from>{pos.format(r=sp_id)}</xdr:from><xdr:to>{pos.format(r=sp_id + 1)}</xdr:to>"
            f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{sp_id}" name="s{sp_id}"/><xdr:cNvSpPr/>'
            f"</xdr:nvSpPr><xdr:spPr/>{body}</xdr:sp><xdr:clientData/></xdr:twoCellAnchor>")


def _inject_shapes(path: Path, texts: list[str | None]) -> None:
    extra = "".join(_xdr_sp(60 + i, t) for i, t in enumerate(texts))

    def add(raw: bytes) -> bytes:
        s = raw.decode("utf-8")
        return re.sub(r"(</(?:\w+:)?wsDr>)\s*$", lambda m: extra + m.group(1), s).encode("utf-8")

    _rewrite_zip(path, "xl/drawings/drawing1.xml", add)


def build_xlsx(tmp_path: Path, *, images: int = 0, comments: int = 0, hidden: bool = False,
               chart: bool = False, shapes: list[str | None] | None = None, name: str = "要件.xlsx") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "要件"
    for r in range(1, 6):
        ws.cell(row=r, column=1, value=f"システムは、要件 {r} を満たすこと。")
        ws.cell(row=r, column=2, value=r)
        ws.cell(row=r, column=3, value=r * 2)
    for i in range(comments):
        ws.cell(row=i + 1, column=1).comment = Comment(f"コメント {i}", "作成者")
    png = _png(tmp_path)
    for i in range(images):
        ws.add_image(XlImage(str(png)), f"E{i * 5 + 1}")
    if chart or shapes:
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=5))
        ws.add_chart(ch, "H1")
    if hidden:
        wb.create_sheet("隠し").sheet_state = "hidden"
        wb["隠し"]["A1"] = "非表示シートの文"
        ws.column_dimensions["B"].hidden = True
        ws.column_dimensions["C"].hidden = True
        ws.row_dimensions[4].hidden = True
    path = tmp_path / name
    wb.save(path)
    if shapes:
        _inject_shapes(path, shapes)
    return path


def _inline_drawing(uri: str, inner: str) -> str:
    return (f'<w:r xmlns:w="{NS_W}" xmlns:wp="{NS_WP}" xmlns:a="{NS_A}" xmlns:wps="{NS_WPS}"><w:drawing>'
            f'<wp:inline><wp:extent cx="914400" cy="914400"/><wp:docPr id="{abs(hash(inner)) % 9000 + 100}" name="d"/>'
            f'<a:graphic><a:graphicData uri="{uri}">{inner}</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')


def _textbox_run(text: str) -> str:
    return _inline_drawing(URI_WPS, f"<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r>"
                                    "</w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp>")


def _smartart_run() -> str:
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    return _inline_drawing(URI_DGM, f'<dgm:relIds xmlns:dgm="{URI_DGM}" xmlns:r="{r}" r:dm="" r:lo="" r:qs="" r:cs=""/>')


def build_docx_unread(tmp_path: Path) -> Path:
    doc = Document()
    para = doc.add_paragraph("システムは、画面を表示すること。")
    doc.add_comment(para.runs, text="確認してください", author="レビュアー")
    png = _png(tmp_path)
    doc.add_picture(str(png))
    doc.add_picture(str(png))
    doc.add_paragraph("図の説明です。")._p.append(parse_xml(_smartart_run()))
    path = tmp_path / "設計.docx"
    doc.save(path)
    return path


def build_docx_structure(tmp_path: Path) -> Path:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "ヘッダーの文"
    doc.sections[0].footer.paragraphs[0].text = "フッターの文"
    doc.add_paragraph("前の段落の文です。")
    table = doc.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "結合セルの文"
    table.cell(1, 0).text = "左下の文"
    inner = table.cell(1, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "入れ子の文"
    doc.add_paragraph("後の段落の文です。")
    doc.add_paragraph("")._p.append(parse_xml(_textbox_run("ボックスの文")))
    path = tmp_path / "構造.docx"
    doc.save(path)
    return path


def _texts(path: Path) -> list[str]:
    return [p.text for p in readers.read_document(path).passages]


def _index_of(texts: list[str], needle: str) -> int:
    hits = [i for i, t in enumerate(texts) if needle in t]
    assert hits, f"{needle} が断片に無い: {texts}"
    return hits[0]


# ---------------------------------------------------------------- checks（FR-06-022）
def _p(text: str, pid: str = "P1") -> list[dict]:
    return [{"id": pid, "text": text, "location": "1 行"}]


def _rules(text: str) -> list[str]:
    return [f["rule"] for f in checks.candidates(_p(text))]


def test_incomplete_is_major_or_higher():
    """FR-06-022 G「詳細は別途定める」 W 検出 T REQ.INCOMPLETE が重大度「主要」以上で出る。"""
    found = [f for f in checks.candidates(_p("詳細は別途定める。")) if f["rule"] == "REQ.INCOMPLETE"]
    assert len(found) == 1
    assert found[0]["severity"] in ("critical", "major")
    assert found[0]["passage_id"] == "P1"


def test_incomplete_applies_to_non_requirement_sentence():
    """FR-06-022 G 義務表現の無い文に「TBD」 W 検出 T INCOMPLETE は全文対象なので出る。"""
    assert "REQ.INCOMPLETE" in _rules("上限値は TBD とする予定。")


def test_open_ended_on_requirement():
    """FR-06-022 G 「ほか」で閉じない列挙の要求文（曖昧語辞書に無い語） W 検出 T REQ.OPEN_ENDED が出る。"""
    assert "REQ.OPEN_ENDED" in _rules("システムは、氏名・住所ほかを保存すること。")


def test_open_ended_suppressed_when_ambiguous_term_already_flagged():
    """FR-06-022（契約: 曖昧語として指摘済みなら出さない） G「など」の要求文 W 検出 T AMBIGUOUS_TERM のみで OPEN_ENDED は出ない。"""
    rules = _rules("システムは、氏名・住所などを保存すること。")
    assert {"RD.AMBIGUOUS", "REQ.AMBIGUOUS_TERM"} & set(rules)  # 曖昧語の規則（辞書由来か追加語か）で指摘済み
    assert "REQ.OPEN_ENDED" not in rules


def test_universal_only_on_requirement():
    """FR-06-022 G 絶対表現「すべて」 W 要求文と説明文を検出 T 要求文だけ REQ.UNIVERSAL が出る。"""
    assert "REQ.UNIVERSAL" in _rules("システムは、すべての利用者に通知を送ること。")
    assert "REQ.UNIVERSAL" not in _rules("すべての画面は青色である。")


def test_and_or_on_requirement():
    """FR-06-022 G「および/または」の要求文 W 検出 T REQ.AND_OR が出る。"""
    assert "REQ.AND_OR" in _rules("システムは、氏名および/または住所を保存すること。")


def test_clean_requirement_has_no_29148_rules():
    """FR-06-022 G 未決定・列挙・絶対・および/または を含まない要求文 W 検出 T 4 規則はどれも出ない。"""
    rules = set(_rules("システムは、注文確定から 3 秒以内に確認画面を表示すること。"))
    assert not rules & {"REQ.INCOMPLETE", "REQ.OPEN_ENDED", "REQ.UNIVERSAL", "REQ.AND_OR"}


# ---------------------------------------------------------------- readers（FR-06-012・013）
def test_docx_reader_order_nested_merged_textbox_header(tmp_path):
    """FR-06-012 G 段落→表（結合セル・入れ子表）→段落、テキストボックス、ヘッダー/フッターを含む docx
    W 読む T 文書順に並び、入れ子・ボックス・ヘッダー・フッターの文が出て、結合セルの文は 1 回だけ。"""
    texts = _texts(build_docx_structure(tmp_path))
    before, merged, nested, after = (_index_of(texts, s) for s in ("前の段落の文", "結合セルの文", "入れ子の文", "後の段落の文"))
    assert before < merged <= nested < after
    assert sum(t.count("結合セルの文") for t in texts) == 1
    for needle in ("ボックスの文", "ヘッダーの文", "フッターの文"):
        _index_of(texts, needle)


def test_xlsx_reader_formula_without_cache_and_shape_text(tmp_path):
    """FR-06-013 G 保存値の無い式・図形内の文字・結合セルを含む xlsx W 読む T 式の文字列と図形の文が出て、結合セルの文は 1 回。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "システムは、合計を表示すること。"
    ws["A2"] = "=SUM(1,2)"
    ws["A3"] = "結合の文"
    ws.merge_cells("A3:C3")
    ws.add_chart(BarChart(), "E1")
    path = tmp_path / "式.xlsx"
    wb.save(path)
    _inject_shapes(path, ["図形の文"])
    texts = _texts(path)
    _index_of(texts, "SUM(1,2)")
    _index_of(texts, "図形の文")
    assert sum(t.count("結合の文") for t in texts) == 1


def test_docx_with_xml_entity_is_rejected(tmp_path):
    """FR-06-014・NFR-06-16 G DOCTYPE/ENTITY を埋めた docx W 取り込む T 実体を展開せず IngestError、文に例外名・パスが無い。"""
    path = build_docx_structure(tmp_path)

    def poison(raw: bytes) -> bytes:
        s = raw.decode("utf-8")
        s = re.sub(r"^(<\?xml[^>]*\?>)", r'\1<!DOCTYPE w:document [<!ENTITY xx "ENTITY_EXPANDED">]>', s)
        return s.replace("前の段落の文です。", "&xx;").encode("utf-8")

    _rewrite_zip(path, "word/document.xml", poison)
    with pytest.raises(ingest.IngestError) as ei:
        ingest.ingest("構造.docx", path.read_bytes(), tmp_path / "work")
    msg = str(ei.value)
    assert msg and "ENTITY_EXPANDED" not in msg
    assert str(tmp_path) not in msg and "Error" not in msg and "Traceback" not in msg


# ---------------------------------------------------------------- unread.scan（FR-06-015）
def test_scan_keys_constant():
    """FR-06-015 G — W KEYS を見る T 5 種が設計どおりの順で並ぶ。"""
    assert unread.KEYS == ("images", "comments", "hidden_sheets", "hidden_cols", "shapes")


def test_scan_xlsx_images_and_comments(tmp_path):
    """UT-06R1-01 FR-06-015 G 画像 3・セルコメント 2 の xlsx W scan T images 3・comments 2・scanned 真。"""
    u = unread.scan(build_xlsx(tmp_path, images=3, comments=2), ".xlsx")
    assert u == {"images": 3, "comments": 2, "hidden_sheets": 0, "hidden_cols": 0, "shapes": 0, "scanned": True}


def test_scan_xlsx_hidden(tmp_path):
    """UT-06R1-02 FR-06-015 G 非表示シート 1・非表示列 2・非表示行 1 W scan T hidden_sheets 1・hidden_cols 3。"""
    u = unread.scan(build_xlsx(tmp_path, hidden=True), ".xlsx")
    assert (u["hidden_sheets"], u["hidden_cols"], u["scanned"]) == (1, 3, True)


def test_scan_xlsx_shapes(tmp_path):
    """UT-06R1-03 FR-06-015 G グラフ 1・文字なし図形 1・文字あり図形 1 W scan T shapes 2。"""
    u = unread.scan(build_xlsx(tmp_path, chart=True, shapes=[None, "文字のある図形"]), ".xlsx")
    assert (u["shapes"], u["images"], u["scanned"]) == (2, 0, True)


def test_scan_docx(tmp_path):
    """UT-06R1-04 FR-06-015 G 画像 2・コメント 1・SmartArt 1 の docx W scan T images 2・comments 1・shapes 1。"""
    u = unread.scan(build_docx_unread(tmp_path), ".docx")
    assert u == {"images": 2, "comments": 1, "hidden_sheets": 0, "hidden_cols": 0, "shapes": 1, "scanned": True}


def test_scan_docx_textbox_with_text_is_not_a_shape(tmp_path):
    """FR-06-015 G 文字のあるテキストボックスだけの docx W scan T 読める図形なので shapes 0。"""
    u = unread.scan(build_docx_structure(tmp_path), ".docx")
    assert (u["shapes"], u["scanned"]) == (0, True)


def test_scan_pdf(tmp_path):
    """UT-06R1-05 FR-06-015 G 注釈 1・リンク 1・画像 1 の PDF W scan T comments 1（リンクは数えない）・images 1。"""
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import Link, Text

    raw = tmp_path / "raw.pdf"
    PILImage.new("RGB", (60, 60), (0, 0, 255)).save(raw, "PDF")
    w = PdfWriter()
    w.append(PdfReader(raw))
    w.add_annotation(0, Text(rect=(5, 5, 20, 20), text="注記"))
    w.add_annotation(0, Link(rect=(25, 25, 40, 40), url="https://example.com"))
    path = tmp_path / "注釈.pdf"
    with path.open("wb") as fh:
        w.write(fh)
    u = unread.scan(path, ".pdf")
    assert (u["images"], u["comments"], u["shapes"], u["scanned"]) == (1, 1, 0, True)


def test_scan_md_and_txt(tmp_path):
    """UT-06R1-06 FR-06-015 G `![a](b)` 2 個の md と同じ本文の txt W scan T md は images 2、txt は 0。"""
    body = "# 見出し\n![a](b) 本文 ![c](d.png)\n[リンク](e)\n"
    (tmp_path / "a.md").write_text(body, encoding="utf-8")
    (tmp_path / "a.txt").write_text(body, encoding="utf-8")
    assert unread.scan(tmp_path / "a.md", ".md")["images"] == 2
    assert unread.scan(tmp_path / "a.txt", ".txt") == {**ZERO, "scanned": True}


@pytest.mark.parametrize("suffix", [".xlsx", ".docx"])
def test_scan_broken_zip_returns_zero(tmp_path, suffix):
    """UT-06R1-07 FR-06-015 G 壊れた zip W scan T 例外なし・全 0・scanned 偽。"""
    path = tmp_path / f"壊れ{suffix}"
    path.write_bytes(b"PK\x03\x04 broken archive")
    assert unread.scan(path, suffix) == ZERO


def test_scan_doctype_xml_returns_zero(tmp_path):
    """UT-06R1-07 FR-06-015 G workbook.xml に DOCTYPE がある xlsx W scan T 例外なし・全 0・scanned 偽。"""
    path = build_xlsx(tmp_path, images=1, hidden=True)
    _rewrite_zip(path, "xl/workbook.xml",
                 lambda raw: b'<!DOCTYPE x [<!ENTITY e "a">]>' + raw)  # openpyxl は XML 宣言を書かないので先頭に置く
    assert unread.scan(path, ".xlsx") == ZERO


# ---------------------------------------------------------------- unread.lines・metrics.unread_lines
def _u(**kw) -> dict:
    return {**ZERO, "scanned": True, **kw}


def test_lines_images_comments():
    """UT-06R1-08 FR-06-015 G images 3・comments 2 W lines T M-01 の文 1 行（0 の種類は省く）。"""
    assert unread.lines(_u(images=3, comments=2)) == [M01_EX]


def test_lines_all_zero_and_failed_and_hidden():
    """UT-06R1-09 FR-06-015 G 全 0／scanned 偽／非表示のみ W lines T M-03／M-05／M-02。"""
    assert unread.lines(_u()) == [M03]
    assert unread.lines(dict(ZERO)) == [M05]
    assert unread.lines(_u(hidden_sheets=1, hidden_cols=3)) == [M02_EX]


def test_metrics_unread_lines_old_json_and_no_document():
    """UT-06R1-09 FR-06-015 G 文書なし／unread の無い旧 document W unread_lines T []／M-04。"""
    assert metrics.unread_lines({"document": None, "passages": [], "findings": []}) == []
    old = {"document": {"name": "a.txt", "format": "txt", "passages": 1, "version": 1}, "passages": [], "findings": []}
    assert metrics.unread_lines(old) == [M04]


# ---------------------------------------------------------------- ingest・store
def test_ingest_document_shape(tmp_path):
    """UT-06R1-10 FR-06-015 G 画像 1 の正常な xlsx W ingest T document のキーが 5 つ、unread に KEYS＋scanned。"""
    path = build_xlsx(tmp_path, images=1)
    doc, passages = ingest.ingest("要件.xlsx", path.read_bytes(), tmp_path / "work")
    assert set(doc) == {"name", "format", "passages", "version", "unread"}
    assert set(doc["unread"]) == set(unread.KEYS) | {"scanned"}
    assert doc["unread"]["images"] == 1 and doc["passages"] == len(passages) > 0


def test_store_save_adds_updated_at_without_mutating(tmp_path, monkeypatch):
    """UT-06R1-11 SCR-06-01 G 任意の insp W save T 戻り値と保存内容に updated_at、引数は不変。"""
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    insp = store.new("更新日時")
    before = copy.deepcopy(insp)
    saved = store.save(insp)
    assert insp == before and "updated_at" not in insp
    stamp = datetime.fromisoformat(saved["updated_at"].replace("Z", "+00:00"))
    assert stamp.utcoffset().total_seconds() == 0
    assert store.load(insp["id"])["updated_at"] == saved["updated_at"]


# ---------------------------------------------------------------- metrics.quality（FR-06-100〜103）
def _insp(n_passages: int, findings: list[dict], *, versions=None, name="要件.xlsx", status="reviewing") -> dict:
    passages = [{"id": f"P{i}", "text": f"システムは、要件 {i} を満たすこと。", "location": f"{i} 行"}
                for i in range(1, n_passages + 1)]
    doc = ({"name": name, "format": "xlsx", "passages": n_passages, "version": len(versions or []) or 1,
            "unread": _u(images=3, comments=2)} if n_passages else None)
    return {"id": "0123456789ab", "name": "検査", "created_at": "2026-10-01T00:00:00Z", "status": status,
            "goal": "", "exit_criteria": {"max_open_critical": 0, "max_open_warning": 5}, "document": doc,
            "passages": passages, "findings": findings, "versions": versions or [], "diff": None,
            "participants": [], "phase": "triage"}


def _f(fid: str, pid: str, status: str, *, severity="warning", rule="REQ.AMBIGUOUS_TERM", msg="曖昧語", quote="適切に") -> dict:
    return {"id": fid, "rule": rule, "severity": severity, "viewpoint": "曖昧さ", "message": msg,
            "location": "1 行", "passage_id": pid, "quote": quote, "certainty": "要確認",
            "status": status, "answer": None, "origin": "auto"}


def _ten_with_two_open() -> dict:
    return _insp(10, [_f("F1", "P1", "candidate"), _f("F2", "P1", "adopted"), _f("F3", "P2", "adopted"),
                      _f("F4", "P3", "rejected"), _f("F5", "P4", "fixed")])


def test_quality_overall_80_with_labels():
    """UT-06R1-12 FR-06-100 G 10 断片中 2 断片に未解決（候補・採用） W quality T overall 80、各 items に label。"""
    q = metrics.quality(_ten_with_two_open())
    assert q["overall"] == 80 and q["passages"] == 10
    assert q["items"]["unambiguous"]["label"] == "曖昧でない"
    for key, item in q["items"].items():
        assert item["label"] == metrics.ITEM_LABEL[key] and item["label"]


def test_quality_no_passages_is_none():
    """UT-06R1-13 FR-06-102 G 断片 0 W quality T overall None・全 score None（0 でない）。"""
    q = metrics.quality(_insp(0, []))
    assert q["overall"] is None
    assert q["items"] and all(item["score"] is None for item in q["items"].values())


def test_quality_note_three_sentences():
    """UT-06R1-14 FR-06-103 G — W note T 3 文で、式と「正しいこと」「網羅性」を含む。"""
    note = metrics.quality(_ten_with_two_open())["note"]
    assert note == list(metrics.QUALITY_NOTE) and len(note) == 3
    joined = "".join(note)
    assert "100 ×（1 − 未解決の指摘がある断片数 ÷ 断片数）" in joined
    assert "正しいこと" in joined and "網羅性" in joined


def test_quality_brief_shape():
    """FR-06-104 G 10 断片中 2 に未解決 W quality_brief(final 偽) T overall 80・items は特性別の値・final 偽。"""
    brief = metrics.quality_brief(_ten_with_two_open(), False)
    assert brief["overall"] == 80 and brief["final"] is False
    assert brief["items"]["unambiguous"] == metrics.quality(_ten_with_two_open())["items"]["unambiguous"]["score"]


# ---------------------------------------------------------------- history・freeze_current（FR-06-104）
def _v(n: int, quality=None) -> dict:
    v = {"version": n, "name": "要件.xlsx", "uploaded_at": f"2026-10-0{n}T00:00:00Z", "passages": 10}
    return {**v, "quality": quality} if quality is not None else v


def test_history_final_and_current():
    """UT-06R1-15 FR-06-104 G 版 1 確定 60・版 2 現在 80 W history T overall [60, 80]、版 2 だけ current 真。"""
    insp = {**_ten_with_two_open(), "versions": [_v(1, {"overall": 60, "items": {}, "final": True}),
                                                 _v(2, {"overall": 75, "items": {}, "final": False})]}
    h = metrics.history(insp)
    assert [x["overall"] for x in h] == [60, 80]
    assert [x["current"] for x in h] == [False, True]
    assert [x["final"] for x in h] == [True, False]


def test_history_old_version_without_quality():
    """UT-06R1-16 FR-06-104 G quality の無い旧版 1 と現在の版 2 W history T 版 1 の overall None。"""
    h = metrics.history({**_ten_with_two_open(), "versions": [_v(1), _v(2)]})
    assert h[0]["overall"] is None and h[1]["overall"] == 80


def test_freeze_current_marks_last_final_and_is_pure():
    """UT-06R1-17 FR-06-104 G 版 2 つ W freeze_current T 最後だけ final 真、入力は不変、空なら []。"""
    insp = {**_ten_with_two_open(), "versions": [_v(1, {"overall": 60, "items": {}, "final": True}), _v(2)]}
    before = copy.deepcopy(insp)
    out = metrics.freeze_current(insp)
    assert insp == before
    assert out[0] == insp["versions"][0]
    assert out[1]["quality"]["final"] is True and out[1]["quality"]["overall"] == 80
    assert metrics.freeze_current({**insp, "versions": []}) == []


# ---------------------------------------------------------------- scope（FR-06-075・NFR-06-08）
def _rows(rows: list[dict]) -> dict:
    return {r["label"]: r["value"] for r in rows}


def test_scope_labels_and_limits():
    """UT-06R1-18 FR-06-075・NFR-06-08 G 取り込み済み W scope T checked 7 行・unchecked に網羅性「未検証。」・規則の精度「未測定。」。"""
    s = metrics.scope(_ten_with_two_open())
    assert [r["label"] for r in s["checked"]] == ["文書", "断片数", "読めなかった要素", "自動検出", "人の確認",
                                                   "観点チェックリスト", "品質スコア"]
    assert [r["label"] for r in s["unchecked"]] == ["網羅性", "意味の正しさ", "読めなかった要素の中身", "スコアの限界", "規則の精度"]
    chk, unc = _rows(s["checked"]), _rows(s["unchecked"])
    assert chk["文書"] == "要件.xlsx（版 1・xlsx）" and chk["断片数"] == "10 断片"
    assert chk["読めなかった要素"] == M01_EX and chk["品質スコア"] == "総合 80"
    assert unc["網羅性"] == M30 and unc["意味の正しさ"] == M31 and unc["規則の精度"] == M33
    assert unc["スコアの限界"] == metrics.QUALITY_NOTE[2]
    assert unc["読めなかった要素の中身"] == "画像 3・コメント 2 の中身は検査していません。"


def test_scope_without_document():
    """FR-06-075・FR-06-102 G 未取り込み W scope T 文書は M-20、品質スコアは M-13、読めなかった中身は「該当なし」。"""
    chk, unc = (_rows(r) for r in metrics.scope(_insp(0, [])).values())
    assert chk["文書"] == M20 and chk["品質スコア"] == M13
    assert unc["読めなかった要素の中身"] == "該当なし"


def test_summary_keeps_old_keys_and_adds_new():
    """5.1 summary G 取り込み済み W summary T 既存キーを保ち history・unread_lines・scope を足す。"""
    s = metrics.summary(_ten_with_two_open())
    assert {"by_status", "by_severity", "density", "open_critical", "open_warning", "exit_ok", "reasons", "quality"} <= set(s)
    assert s["unread_lines"] == [M01_EX] and s["scope"] == metrics.scope(_ten_with_two_open())
    assert s["history"] == metrics.history(_ten_with_two_open())


# ---------------------------------------------------------------- to_xlsx（10 章）
def _book(insp: dict) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(metrics.to_xlsx(insp)))


def test_to_xlsx_scope_sheet():
    """UT-06R1-19 FR-06-075・104 G 版 2 つの検査 W to_xlsx T 3 枚目「検査した範囲」、網羅性の C 列が「未検証」で始まる、集計に版ごとの行。"""
    insp = {**_ten_with_two_open(), "versions": [_v(1, {"overall": 60, "items": {}, "final": True}), _v(2)]}
    wb = _book(insp)
    assert wb.sheetnames[2] == "検査した範囲"
    ws = wb["検査した範囲"]
    assert [c.value for c in ws[1]][:3] == ["区分", "項目", "内容"]
    rows = {r[1]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["網羅性"][0] == "検査していないこと" and rows["網羅性"][2].startswith("未検証")
    assert rows["規則の精度"][2].startswith("未測定") and rows["文書"][0] == "検査した範囲"
    assert ws.column_dimensions["A"].width == 18 and ws.column_dimensions["B"].width == 24
    assert ws.column_dimensions["C"].width == 80 and ws["C2"].alignment.wrap_text
    agg = [c for row in wb["集計"].iter_rows(values_only=True) for c in row if isinstance(c, str)]
    assert "版 1 の品質スコア" in agg and "版 2 の品質スコア" in agg


def test_to_xlsx_formula_injection_escaped():
    """UT-06R1-20 現行の回帰 G 先頭 `=` の名前・原文 W to_xlsx T 全シートで「'」が前置され `=` で始まるセルが無い。"""
    insp = {**_ten_with_two_open(), "document": {**_ten_with_two_open()["document"], "name": "=cmd.xlsx"}}
    insp["findings"] = [{**f, "quote": "=HYPERLINK(\"x\")", "message": "=1+1"} for f in insp["findings"]]
    wb = _book(insp)
    values = [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)]
    assert not [v for v in values if v.startswith("=")]
    assert "'=1+1" in values and any(v.startswith("'=cmd.xlsx") for v in values)


# ---------------------------------------------------------------- home（SCR-06-01・FR-06-101・102）
def _summ(overall, score) -> dict:
    return {"quality": {"overall": overall, "passages": 10 if overall is not None else 0,
                        "items": {"unambiguous": {"open": 1, "score": score, "label": "曖昧でない"}}}}


def test_average_quality_70():
    """UT-06R1-21 FR-06-101 G 進行中 2 件（80・60） W average_quality T overall 70、特性別も平均、n_reviews 2。"""
    q = home.average_quality([_summ(80, 80), _summ(60, 40)])
    assert (q["overall"], q["n_reviews"], q["n_scored"]) == (70, 2, 2)
    assert q["items"]["unambiguous"] == {"score": 60, "label": "曖昧でない"}


def test_average_quality_rounds_half_up_and_skips_none():
    """FR-06-101 G 80・61・算出不能 1 件 W average_quality T 70.5 は四捨五入で 71、None は分母に入れない。"""
    q = home.average_quality([_summ(80, 80), _summ(61, 61), _summ(None, None)])
    assert (q["overall"], q["n_reviews"], q["n_scored"]) == (71, 3, 2)


def test_average_quality_none_when_unscored():
    """UT-06R1-22 FR-06-102 G 算出できる 0 件 W average_quality T overall None・n_scored 0（0 でない）。"""
    q = home.average_quality([_summ(None, None)])
    assert q["overall"] is None and q["n_scored"] == 0
    assert home.average_quality([])["overall"] is None


def _ask(fid: str, status: str, answer=None) -> dict:
    return {**_f(fid, "P1", status, severity="info"), "certainty": "要確認", "answer": answer}


def test_awaiting_answers_counts_open_questions():
    """UT-06R1-23 SCR-06-01 G 要確認の未回答 2・回答済み 1・却下 1 W awaiting_answers T 2。終了済みは build で数えない。"""
    review = _insp(5, [_ask("F1", "adopted"), _ask("F2", "adopted"), _ask("F3", "adopted", "回答"),
                       _ask("F4", "rejected")])
    assert home.awaiting_answers(review) == 2
    built = home.build([{**review, "status": "closed"}], "山田")
    assert built["counts"]["awaiting_answers"] == 0


_PRIORITY = ["critical", "candidates", "answers", "my_review", "no_document", "ready_to_close"]


def _review(rid: str, findings: list[dict], *, due=None, docs=True, participants=None) -> dict:
    r = {**_insp(5 if docs else 0, findings), "id": rid, "name": f"レビュー{rid}", "due": due,
         "participants": participants or []}
    return r


def test_next_actions_priority_limit_and_scope():
    """UT-06R1-24 SCR-06-01 G 重大・候補・回答待ち・文書なしを含む 7 件・自分は不参加 W next_actions T 優先順に 5 件、同順位は期限順、scope active。"""
    crit = [_f("F1", "P1", "adopted", severity="critical"), _f("F2", "P2", "adopted", severity="critical")]
    reviews = [_review("a00000000001", [_ask("F1", "adopted")]),
               _review("a00000000002", crit, due="2026-10-20"),
               _review("a00000000003", [], docs=False),
               _review("a00000000004", crit, due=None),
               _review("a00000000005", [_f("F1", "P1", "candidate")]),
               _review("a00000000006", crit, due="2026-10-10"),
               _review("a00000000007", [_f("F1", "P1", "candidate")], due="2026-10-01")]
    items, scope = home.next_actions(reviews, "山田")
    assert scope == "active" and len(items) == 5
    kinds = [i["kind"] for i in items]
    assert kinds == sorted(kinds, key=_PRIORITY.index)
    assert [i["review_id"] for i in items[:3]] == ["a00000000006", "a00000000002", "a00000000004"]
    assert items[0] == {**items[0], "kind": "critical", "view": "triage", "label": "重大な指摘 2 件に対応する"}
    assert items[3]["kind"] == "candidates" and items[3]["review_id"] == "a00000000007"
    assert items[3]["label"] == "未検討の候補 1 件を仕分ける"


def test_next_actions_mine_scope_and_my_review():
    """SCR-06-01 G 自分がレビュアーで未完了のレビュー W next_actions T scope mine、kind my_review・view read。"""
    r = _review("b00000000001", [], participants=[{"name": "山田", "role": "reviewer", "done": False}])
    items, scope = home.next_actions([r], "山田")
    assert scope == "mine"
    mine = [i for i in items if i["kind"] == "my_review"]
    assert mine and mine[0]["view"] == "read" and mine[0]["label"] == "自分の確認を終える"


def test_build_keeps_existing_keys_and_adds_new():
    """UT-06R1-25 回帰 G 自分が参加する進行中 1 件（updated_at の無い旧 JSON） W build T mine・active・counts の既存キーが残り、追加キーが揃う。"""
    r = _review("c00000000001", [_f("F1", "P1", "adopted", severity="critical")],
                participants=[{"name": "山田", "role": "reviewer", "done": False}])
    out = home.build([r], "山田")
    assert out["me"] == "山田"
    assert {"review_id", "review_name", "role", "done", "phase", "due"} <= set(out["mine"][0])
    act = out["active"][0]
    assert {"review_id", "review_name", "phase", "due", "findings", "quality_overall", "open", "open_critical",
            "candidates", "updated_at", "versions", "document_name"} <= set(act)
    assert act["updated_at"] == r["created_at"] and act["document_name"] == "要件.xlsx"
    assert {"mine", "active", "open_findings", "open_critical", "awaiting_answers"} <= set(out["counts"])
    assert out["counts"]["open_critical"] == 1 and out["next_scope"] == "mine"
    assert out["quality"]["n_reviews"] == 1 and out["quality"]["overall"] == act["quality_overall"]
