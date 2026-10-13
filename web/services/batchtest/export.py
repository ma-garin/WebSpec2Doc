"""書き出し。ケース表（xlsx）・テストコード一式（zip）・要確認の質問票（CSV）。"""
from __future__ import annotations

import csv
import io
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from web.services.batchtest import store

_HEAD = ("ID", "観点", "技法", "前提・入力", "実行", "期待結果", "根拠", "状態")


def cases_xlsx(job: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ケース表"
    ws.append(_HEAD)
    for c in job["cases"]:
        ws.append((c["id"], c["obs"], c["tech"], c["pre"], c["run"], c["exp"], c["src"],
                   "要確認" if c["ask"] else "根拠あり"))
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="F1F5F9")
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    for col, width in zip("ABCDEFGH", (8, 12, 12, 32, 36, 44, 18, 10)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def code_zip(job: dict) -> bytes:
    if not job.get("sources"):
        raise ValueError("テストコードがまだありません。先にテストコードを生成してください。")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for s in job["sources"]:
            if not s["name"].endswith("/"):
                z.writestr(s["name"], s["body"])
        for f in job["files"]:
            if f["kind"] in ("script", "input", "expected"):
                path = store.file_path(job, f["name"])
                if path.is_file():
                    z.write(path, f"fixtures/{f['name']}")
    return buf.getvalue()


def questions_csv(job: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(("ID", "観点", "前提・入力", "実行", "確かめたいこと", "案"))
    for c in job["cases"]:
        if c["ask"]:
            proposal = c["exp"].removeprefix("仕様に記載なし（案: ").removesuffix("）")
            w.writerow((c["id"], c["obs"], c["pre"], c["run"], "この条件で実行したときの期待結果", proposal))
    return ("﻿" + buf.getvalue()).encode("utf-8")
