"""R4 レビュー工程管理: 種別・役割・入口／出口基準・工程の状態機械・工数・指標（WS2D-DD-006 R4 2 章）。

すべて純粋関数。引数の dict/list は変更せず、新しいものを返す。R4 の追加項目は検査 JSON の
"r4" に置く（v1・v2 の JSON は "r4" が無くても読める。NFR-06-17）。
"""
from __future__ import annotations

PHASES = ("plan", "individual", "triage", "rework", "followup", "closed")
PHASE_LABEL = {"plan": "計画", "individual": "個別レビュー", "triage": "指摘の検討", "rework": "修正",
               "followup": "修正版の確認", "closed": "判定と記録"}
ROLES = ("author", "moderator", "reviewer", "scribe", "approver")
ROLE_LABEL = {"author": "作成者", "moderator": "議長", "reviewer": "レビュアー", "scribe": "記録者",
              "approver": "承認者"}

# ISO/IEC 20246:2017 のレビュー種別を参照。必須役割・工程・既定の出口基準は種別で切り替える（FR-06-140）
REVIEW_TYPES = {
    "informal": {"label": "非公式レビュー", "required": ("author",),
                 "phases": ("plan", "individual", "closed"), "checklist": False,
                 "exit": {"require_all_verified": False, "require_all_individual_done": False}},
    "walkthrough": {"label": "ウォークスルー", "required": ("author", "reviewer"),
                    "phases": ("plan", "individual", "triage", "closed"), "checklist": False,
                    "exit": {"require_all_verified": False, "require_all_individual_done": False}},
    "technical": {"label": "技術レビュー", "required": ("author", "moderator", "reviewer"),
                  "phases": PHASES, "checklist": True,
                  "exit": {"require_all_verified": True, "require_all_individual_done": False}},
    "inspection": {"label": "インスペクション", "required": ("author", "moderator", "scribe", "reviewer"),
                   "phases": PHASES, "checklist": True,
                   "exit": {"require_all_verified": True, "require_all_individual_done": True}},
    "management": {"label": "マネジメントレビュー", "required": ("moderator", "reviewer", "approver"),
                   "phases": ("plan", "individual", "triage", "closed"), "checklist": False,
                   "exit": {"require_all_verified": False, "require_all_individual_done": False}},
    "audit": {"label": "監査", "required": ("moderator", "reviewer", "approver"),
              "phases": PHASES, "checklist": True,
              "exit": {"require_all_verified": True, "require_all_individual_done": True}},
}
LEGACY_LABEL = "未設定"
LEGACY_HINT = "レビューの種別を選ぶと、必須の役割と基準が決まります。"

KINDS = {"missing": "欠落", "wrong": "誤り", "ambiguous": "曖昧", "contradiction": "矛盾", "violation": "規約違反"}
VERDICTS = {"approved": "承認", "conditional": "条件付き承認", "rereview": "再レビュー"}

# 採用した規格・資料と扱い（FR-06-076）。「認証済み」「規格に準拠」と書かない
STANDARDS = (
    {"name": "ISO/IEC/IEEE 29148", "version": "2018",
     "scope": "要求の品質特性を観点チェックリストの項目として参照しています。"},
    {"name": "ISO/IEC 20246", "version": "2017",
     "scope": "レビューの種別・役割・工程・入口と出口の基準を参照しています。"},
    {"name": "IPA ソフトウェア開発データ白書", "version": "2018-2019",
     "scope": "指摘密度の参考範囲として表示しています。合否には使いません。"},
)
STANDARDS_NOTE = "上の規格・資料は設計の参考として使っています。認証や第三者の評価を受けたものではありません。"

# IPA 白書 2018-2019（E5 4 章）。P25〜P75 を範囲、中央値を目安として並べるだけ（FR-06-147）
REFERENCE = {
    "basic": {"label": "基本設計", "low": 0.12, "high": 0.76, "median": 0.28,
              "source": "IPA『ソフトウェア開発データ白書 2018-2019』図7-3-5（P25〜P75・中央値）"},
    "detail": {"label": "詳細設計", "low": 0.07, "high": 0.42, "median": 0.16,
               "source": "IPA『ソフトウェア開発データ白書 2018-2019』図7-3-10（P25〜P75・中央値）"},
}
REFERENCE_LIMIT = "他組織のプロジェクトの分布で、業種・規模・数え方の違いを含みます。合否の判定には使いません。"

M_NO_DOC = "文書を取り込んでください。"
M_NO_CHECKLIST = "観点チェックリストを選んでください。"
M_NO_PAGES = "頁数が無いため算出できません。指標の欄に文書の頁数を入れてください。"
M_NO_MEETING = "会議時間が無いため算出できません。工数の欄に会議時間を入れてください。"
M_APPROVAL_STALE = "承認後に文書が変わりました。承認は失効しています。再レビューしてください。"


def r4(insp: dict) -> dict:
    return dict(insp.get("r4") or {})


def with_r4(insp: dict, **changes) -> dict:
    return {**insp, "r4": {**r4(insp), **changes}}


def type_info(insp: dict) -> dict:
    t = r4(insp).get("review_type")
    if t not in REVIEW_TYPES:
        return {"type": None, "label": LEGACY_LABEL, "legacy": True, "hint": LEGACY_HINT,
                "required": [], "phases": list(PHASES)}
    spec = REVIEW_TYPES[t]
    return {"type": t, "label": spec["label"], "legacy": False, "hint": "",
            "required": list(spec["required"]), "phases": list(spec["phases"])}


def _names(participants: list, role: str) -> set:
    return {p.get("name") for p in participants or [] if p.get("role") == role}


def plan_errors(review_type, participants: list) -> list[str]:
    """計画の保存前に判定する。種別の必須役割（FR-06-140）と兼務制約（FR-06-141）。"""
    if review_type not in REVIEW_TYPES:
        return ["レビューの種別を選んでください。"]
    errs = []
    for role in REVIEW_TYPES[review_type]["required"]:
        if not _names(participants, role):
            errs.append(f"{REVIEW_TYPES[review_type]['label']}では{ROLE_LABEL[role]}が必須です。参加者に追加してください。")
    authors = _names(participants, "author")
    for role in ("moderator", "approver"):
        both = sorted(authors & _names(participants, role))
        if both:
            errs.append(f"作成者と{ROLE_LABEL[role]}は兼務できません（{'、'.join(both)}）。別の人を選んでください。")
    return errs


def next_phase(insp: dict):
    phases = type_info(insp)["phases"]
    cur = insp.get("phase") or "plan"
    if cur not in phases:
        later = [p for p in phases if PHASES.index(p) > PHASES.index(cur)] if cur in PHASES else []
        return later[0] if later else None
    i = phases.index(cur)
    return phases[i + 1] if i + 1 < len(phases) else None


def entry_errors(insp: dict, target: str) -> list[str]:
    """工程 target に入る前の入口基準（FR-06-142）。空なら前進できる。"""
    info = type_info(insp)
    errs = []
    if target != "plan" and not insp.get("document"):
        errs.append(M_NO_DOC)
    if target == "individual" and not info["legacy"]:
        missing = [r for r in info["required"] if not _names(insp.get("participants"), r)]
        errs += [f"{ROLE_LABEL[r]}が割り当てられていません。計画の画面で参加者に追加してください。" for r in missing]
        if REVIEW_TYPES[info["type"]]["checklist"] and not (insp.get("checklist_id") or r4(insp).get("preset_id")):
            errs.append(M_NO_CHECKLIST)
    if target == "closed":
        errs += exit_errors(insp)
    return errs


def exit_options(insp: dict) -> dict:
    t = r4(insp).get("review_type")
    base = dict(REVIEW_TYPES[t]["exit"]) if t in REVIEW_TYPES else {
        "require_all_verified": False, "require_all_individual_done": False}
    return {**base, **(r4(insp).get("exit_options") or {})}


def exit_errors(insp: dict) -> list[str]:
    """追加の出口基準（FR-06-143）。現行の未対応件数の判定（metrics.summary）とは別に足す。"""
    opt = exit_options(insp)
    errs = []
    if opt.get("require_all_verified"):
        waiting = sum(1 for f in insp.get("findings") or [] if f.get("status") == "fixed")
        if waiting:
            errs.append(f"修正済みで確認待ちの指摘が {waiting} 件あります。修正版の確認で確認済みにしてください。")
    if opt.get("require_all_individual_done"):
        left = [p.get("name") for p in insp.get("participants") or [] if p.get("role") == "reviewer" and not p.get("done")]
        if left:
            errs.append(f"個別レビューを終えていない人がいます（{'、'.join(left)}）。全員が終えてから進めてください。")
    return errs


def advance(insp: dict, target: str, actor: str, now: str) -> dict:
    """工程を進め、実施記録（FR-06-144）の前工程を閉じて新工程を開く。"""
    log = [dict(e) for e in r4(insp).get("phase_log") or []]
    if log and not log[-1].get("ended_at"):
        log[-1] = {**log[-1], "ended_at": now, "ended_by": actor}
    if target != "closed":
        log.append({"phase": target, "started_at": now, "started_by": actor, "ended_at": None, "ended_by": None})
    else:
        log.append({"phase": "closed", "started_at": now, "started_by": actor, "ended_at": now, "ended_by": actor})
    out = with_r4({**insp, "phase": target}, phase_log=log)
    return {**out, "status": "closed"} if target == "closed" else out


def effort_errors(prep, meeting) -> list[str]:
    if not isinstance(prep, dict) or not all(isinstance(v, int) and not isinstance(v, bool) and v >= 0 for v in prep.values()):
        return ["準備時間は 0 以上の整数（分）で入力してください。"]
    if not (isinstance(meeting, int) and not isinstance(meeting, bool) and meeting >= 0):
        return ["会議時間は 0 以上の整数（分）で入力してください。"]
    return []


def effort(insp: dict) -> dict:
    """工数（FR-06-145）。合計 = 参加者ごとの準備時間の和 + 会議時間。"""
    e = r4(insp).get("effort") or {}
    prep = {str(k): int(v) for k, v in (e.get("prep") or {}).items()}
    meeting = int(e.get("meeting_minutes") or 0)
    return {"prep": prep, "meeting_minutes": meeting, "total_minutes": sum(prep.values()) + meeting}


_ADOPTED = {"adopted", "fixed", "verified"}


def _round(x: float, n: int = 2) -> float:
    from decimal import ROUND_HALF_UP, Decimal
    return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP))


def compare(density, doc_kind, override: dict | None = None) -> dict | None:
    """参考範囲との比較（FR-06-147・148）。合否には使わない（affects_verdict は常に False）。"""
    if doc_kind not in REFERENCE or density is None:
        return None
    ref = {**REFERENCE[doc_kind]}
    ov = (override or {}).get(doc_kind)
    if ov:
        ref = {**ref, "low": ov["low"], "high": ov["high"], "median": None, "source": "組織の基準値（テナントで設定）"}
    if density < ref["low"]:
        note = "参考範囲より低い"
    elif density > ref["high"]:
        note = "参考範囲より高い"
    else:
        note = "参考範囲内"
    return {**ref, "note": note, "attention": note != "参考範囲内", "limit": REFERENCE_LIMIT,
            "affects_verdict": False}


def indicators(insp: dict, override: dict | None = None) -> dict:
    """指標（FR-06-146）。頁数は r4.pages（人が入れる。要確認-13）。"""
    st = r4(insp)
    pages = st.get("pages")
    adopted = [f for f in insp.get("findings") or [] if f.get("status") in _ADOPTED]
    by_sev = {s: sum(1 for f in adopted if f.get("severity") == s) for s in ("critical", "major", "minor")}
    done = sum(1 for f in adopted if f.get("status") in {"fixed", "verified"})
    eff = effort(insp)
    out = {"pages": pages, "adopted": len(adopted), "by_severity": by_sev,
           "fix_rate": _round(done / len(adopted)) if adopted else None,
           "density": None, "speed": None, "effort_per_page": None, "messages": []}
    if not (isinstance(pages, (int, float)) and pages > 0):
        return {**out, "messages": [M_NO_PAGES]}
    out = {**out, "density": _round(len(adopted) / pages),
           "effort_per_page": _round(eff["total_minutes"] / 60 / pages)}
    if eff["meeting_minutes"] > 0:
        out = {**out, "speed": _round(pages / (eff["meeting_minutes"] / 60))}
    else:
        out = {**out, "messages": [M_NO_MEETING]}
    return {**out, "reference": compare(out["density"], st.get("doc_kind"), override)}


def visible_findings(insp: dict, viewer: str) -> list[dict]:
    """個別レビューの独立（FR-06-034）。設定オン・個別レビュー中・本人が未終了なら他者の指摘を隠す。"""
    findings = list(insp.get("findings") or [])
    if not r4(insp).get("independent") or insp.get("phase") != "individual":
        return findings
    me = [p for p in insp.get("participants") or [] if p.get("name") == viewer]
    if me and all(p.get("done") for p in me):
        return findings
    return [f for f in findings if f.get("origin", "auto") == "auto" or f.get("raised_by") == viewer]


def unchecked(items: list, checks: dict) -> list[dict]:
    """判定画面に出す未確認のチェック項目（FR-06-162）。"""
    return [dict(i) for i in items or [] if not (checks or {}).get(i.get("id"))]


def approval_state(insp: dict) -> dict:
    """承認後の変更検知（FR-06-054）。承認時の版より新しい版が取り込まれていれば失効。"""
    ap = insp.get("approval") or {}
    cur = (insp.get("document") or {}).get("version")
    at = ap.get("doc_version")
    stale = bool(ap.get("approver")) and at is not None and cur is not None and cur > at
    return {"approved": bool(ap.get("approver")), "stale": stale, "message": M_APPROVAL_STALE if stale else "",
            "verdict": r4(insp).get("verdict"), "verdict_label": VERDICTS.get(r4(insp).get("verdict"), "")}


# --- 回答の取り込み（FR-06-043 xlsx・FR-06-044 反映できない行） ---------------------------

M_ANSWER_HEADER = "質問票に「No.」と「回答」の列がありません。出力した質問票に回答を書いて取り込んでください。"


def _table(rows: list[list]) -> list[tuple[int, str, str]]:
    """1 行目を見出しとして (行番号, No., 回答) を返す。行番号は表計算ソフトの表示と同じ 1 始まり。"""
    if not rows:
        raise ValueError(M_ANSWER_HEADER)
    head = [str(c or "").strip().lstrip("﻿") for c in rows[0]]
    if "No." not in head or "回答" not in head:
        raise ValueError(M_ANSWER_HEADER)
    ino, ians = head.index("No."), head.index("回答")
    out = []
    for n, r in enumerate(rows[1:], start=2):
        cells = list(r) + [None] * (max(ino, ians) + 1 - len(r))
        no, ans = str(cells[ino] or "").strip().lstrip("'"), str(cells[ians] or "").strip()
        if no or ans:
            out.append((n, no, ans))
    return out


def answer_rows_csv(text: str) -> list[tuple[int, str, str]]:
    import csv
    import io
    return _table(list(csv.reader(io.StringIO(text.lstrip("﻿")))))


def answer_rows_xlsx(data: bytes) -> list[tuple[int, str, str]]:
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return _table([list(r) for r in wb.worksheets[0].iter_rows(values_only=True)])
    finally:
        wb.close()


def apply_answers(insp: dict, rows: list[tuple[int, str, str]]) -> tuple[dict, int, list[dict]]:
    """回答を反映した新しい検査・反映件数・反映できない行を返す（引数は変更しない）。"""
    ids = {f.get("id") for f in insp.get("findings") or []}
    answers, skipped = {}, []
    for n, no, ans in rows:
        if no not in ids:
            skipped.append({"row": n, "reason": "不明な No."})
        elif not ans:
            skipped.append({"row": n, "reason": "回答が空"})
        else:
            answers[no] = ans
    findings = [{**f, "answer": answers[f["id"]], "certainty": "事実"} if f.get("id") in answers else dict(f)
                for f in insp.get("findings") or []]
    return {**insp, "findings": findings}, len(answers), skipped


def skipped_message(skipped: list[dict]) -> str:
    if not skipped:
        return ""
    return "反映できない行: " + "、".join(f"{s['row']} 行目" for s in skipped) + "（不明な No. または空の回答）"


# --- レビュー記録（FR-06-074）と規格の扱い（FR-06-076） -----------------------------------

def _minutes(a, b):
    from datetime import datetime
    try:
        return round((datetime.fromisoformat(b) - datetime.fromisoformat(a)).total_seconds() / 60)
    except (TypeError, ValueError):
        return None


def record_rows(insp: dict) -> list[list]:
    """報告に入れるレビュー記録の行 [項目, 値]。"""
    doc = insp.get("document") or {}
    rows = [["種別", type_info(insp)["label"]]]
    for p in insp.get("participants") or []:
        rows.append(["参加者と役割", f"{p.get('name')}（{ROLE_LABEL.get(p.get('role'), p.get('role'))}）"])
    total = 0
    for e in r4(insp).get("phase_log") or []:
        m = _minutes(e.get("started_at"), e.get("ended_at"))
        total += m or 0
        rows.append([f"工程: {PHASE_LABEL.get(e.get('phase'), e.get('phase'))}",
                     f"{e.get('started_at') or '—'} 〜 {e.get('ended_at') or '進行中'}（担当 {e.get('started_by') or '—'}）"])
    rows.append(["所要時間（工程の合計）", f"{total} 分"])
    rows.append(["工数（準備＋会議）", f"{effort(insp)['total_minutes']} 分"])
    rows.append(["対象版", str(doc.get("version") or "—")])
    rows.append(["対象のハッシュ（SHA-256）", doc.get("sha256") or "記録なし（ハッシュを記録する前に取り込んだ文書）"])
    st = approval_state(insp)
    rows.append(["判定", st["verdict_label"] or "未判定"])
    if st["stale"]:
        rows.append(["承認の状態", st["message"]])
    return rows


def record_xlsx(insp: dict) -> bytes:
    import io
    from openpyxl import Workbook
    from web.services.docverify.metrics import _safe
    wb = Workbook()
    ws = wb.active
    ws.title = "レビュー記録"
    ws.append(["項目", "値"])
    for r in record_rows(insp):
        ws.append([_safe(v) for v in r])
    ws2 = wb.create_sheet("指摘一覧")
    ws2.append(["No.", "種別", "重大度", "指摘", "場所", "状態"])
    for f in insp.get("findings") or []:
        ws2.append([_safe(v) for v in (f.get("id"), KINDS.get(f.get("kind"), "—"), f.get("severity"),
                                       f.get("message"), f.get("location"), f.get("status"))])
    ws3 = wb.create_sheet("規格の扱い")
    ws3.append(["規格・資料", "版", "扱い"])
    for s in STANDARDS:
        ws3.append([s["name"], s["version"], s["scope"]])
    ws3.append(["注記", "", STANDARDS_NOTE])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
