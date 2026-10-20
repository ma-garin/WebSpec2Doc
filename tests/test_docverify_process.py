"""System 06 R4 のテスト（WS2D-DD-006 R4 9 章）。受入基準（RD-006 の G/W/T）を 1 件ずつ確かめる。

Blueprint の登録（web/__init__.py）は wire で行うため、ここでは R4 の Blueprint だけを載せた
Flask アプリで API を叩く。画面（docverify-process.js）は Playwright MCP で別に確かめる（本テストの範囲外）。
"""
from __future__ import annotations

import io
import json
import re

import openpyxl
import pytest
from flask import Flask

from web import config
from web.routes import doc_verify_process as r4route
from web.services.docverify import perm, presets, process, store

P = "/api/docverify/r4"
LEAK = re.compile(r"Traceback|Exception|Error\b|\.py\b|/Users/|/tmp/|/private/|File \"")


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setenv("WEBSPEC2DOC_AUTH_MODE", "off")
    a = Flask(__name__)
    a.config["TESTING"] = True
    a.register_blueprint(r4route.bp)
    return a


@pytest.fixture
def client(app):
    return app.test_client()


def as_user(monkeypatch, name, admin=False, auth=True):
    monkeypatch.setattr(perm, "_context", lambda: (name, admin, auth))


def seed(app, **over):
    with app.test_request_context():
        insp = {**store.new("R4 テスト"), "rev": 0, **over}
        store.save(insp)
        return insp["id"]


def load(app, iid):
    with app.test_request_context():
        return store.load(iid)


DOC = {"name": "a.docx", "format": "docx", "passages": 3, "version": 1}
PS = [{"name": "A", "role": "author"}, {"name": "M", "role": "moderator"},
      {"name": "S", "role": "scribe"}, {"name": "R", "role": "reviewer"}]


def fnd(fid, status="adopted", severity="major", origin="manual", raised_by="R"):
    return {"id": fid, "status": status, "severity": severity, "origin": origin, "raised_by": raised_by,
            "message": "m", "location": "1", "certainty": "要確認"}


# --- FR-06-140 種別 ---------------------------------------------------------------

def test_create_requires_type(client):
    r = client.post(f"{P}/inspections", json={"name": "x"})
    assert r.status_code == 400 and "種別" in r.get_json()["error"]
    r = client.post(f"{P}/inspections", json={"name": "x", "review_type": "inspection"})
    assert r.status_code == 201 and r.get_json()["type"]["label"] == "インスペクション"


def test_inspection_requires_moderator_and_scribe(client, app):
    iid = seed(app)
    r = client.put(f"{P}/inspections/{iid}/plan", json={"rev": 0, "review_type": "inspection",
                                                        "participants": [{"name": "A", "role": "author"}]})
    errs = r.get_json()["errors"]
    assert r.status_code == 400
    assert any("議長が必須" in e for e in errs) and any("記録者が必須" in e for e in errs)
    r = client.put(f"{P}/inspections/{iid}/plan", json={"rev": 0, "review_type": "inspection", "participants": PS})
    assert r.status_code == 200 and r.get_json()["rev"] == 1


def test_legacy_without_type_opens(client, app):
    iid = seed(app)
    body = client.get(f"{P}/inspections/{iid}").get_json()
    assert body["type"]["label"] == "未設定（旧形式）" and body["type"]["hint"]


# --- FR-06-141 兼務 ---------------------------------------------------------------

def test_author_cannot_be_moderator(client, app):
    iid = seed(app)
    ps = [{"name": "A", "role": "author"}, {"name": "A", "role": "moderator"}, {"name": "R", "role": "reviewer"}]
    r = client.put(f"{P}/inspections/{iid}/plan", json={"rev": 0, "review_type": "technical", "participants": ps})
    assert r.status_code == 400 and "兼務できません" in r.get_json()["error"]
    assert load(app, iid).get("participants") == [] and perm.rev_of(load(app, iid)) == 0


# --- FR-06-142〜144 入口・出口・実施記録 --------------------------------------------------

def test_entry_needs_document(client, app):
    iid = seed(app, participants=PS, r4={"review_type": "walkthrough"})
    r = client.post(f"{P}/inspections/{iid}/advance", json={"rev": 0})
    assert r.status_code == 409 and r.get_json()["error"] == "文書を取り込んでください。"
    assert load(app, iid)["phase"] == "plan"


def test_exit_requires_all_verified(client, app):
    iid = seed(app, participants=PS, document=DOC, phase="followup", findings=[fnd("F1", status="fixed")],
               r4={"review_type": "technical", "exit_options": {"require_all_verified": True}})
    r = client.post(f"{P}/inspections/{iid}/advance", json={"rev": 0})
    assert r.status_code == 409 and "確認待ち" in r.get_json()["error"]
    assert load(app, iid)["phase"] == "followup"


def test_phase_log_two_steps(client, app):
    iid = seed(app, participants=PS, document=DOC, r4={"review_type": "walkthrough"})
    assert client.post(f"{P}/inspections/{iid}/advance", json={"rev": 0}).status_code == 200
    body = client.post(f"{P}/inspections/{iid}/advance", json={"rev": 1}).get_json()
    log = body["phase_log"]
    assert [e["phase"] for e in log] == ["individual", "triage"]
    assert log[0]["started_at"] and log[0]["ended_at"] and log[0]["started_by"] and log[1]["started_at"]


# --- FR-06-145〜148 工数・指標・参考範囲 ----------------------------------------------------

def test_effort_total(client, app):
    iid = seed(app)
    r = client.put(f"{P}/inspections/{iid}/effort", json={"rev": 0, "prep": {"A": 60, "B": 60}, "meeting_minutes": 30})
    assert r.get_json()["effort"]["total_minutes"] == 150


def test_indicators_density_and_speed():
    insp = {"findings": [fnd(f"F{i}") for i in range(6)],
            "r4": {"pages": 20, "effort": {"prep": {}, "meeting_minutes": 60}}}
    ind = process.indicators(insp)
    assert ind["density"] == 0.30 and ind["speed"] == 20.0
    assert process.indicators({"findings": []})["messages"] == [process.M_NO_PAGES]
    assert "頁数が無いため算出できません" in process.M_NO_PAGES


def test_reference_low_does_not_change_verdict():
    c = process.compare(0.05, "basic")
    assert c["note"] == "参考範囲より低い" and "IPA" in c["source"] and c["limit"] and c["affects_verdict"] is False
    assert (c["low"], c["high"], c["median"]) == (0.12, 0.76, 0.28)
    d = process.REFERENCE["detail"]
    assert (d["low"], d["high"], d["median"]) == (0.07, 0.42, 0.16)


def test_reference_override(client, app):
    assert client.put(f"{P}/reference", json={"basic": {"low": 0.2, "high": 0.5}}).status_code == 200
    iid = seed(app, findings=[fnd(f"F{i}") for i in range(3)],
               r4={"pages": 20, "doc_kind": "basic", "effort": {"prep": {}, "meeting_minutes": 60}})
    ref = client.get(f"{P}/inspections/{iid}").get_json()["indicators"]["reference"]
    assert (ref["low"], ref["high"]) == (0.2, 0.5) and ref["note"] == "参考範囲より低い"


# --- FR-06-149 改ざん検知 -----------------------------------------------------------

def test_audit_tamper_detected(app, tmp_path):
    with app.test_request_context():
        for i in range(10):
            perm.record("finding", f"x/{i}", {"n": i}, actor="A")
        assert perm.verify_log()["ok"] is True
        path = perm._audit_path()
        lines = path.read_text(encoding="utf-8").splitlines()
        e = json.loads(lines[4])
        lines[4] = json.dumps({**e, "actor": "Z"}, ensure_ascii=False)
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        res = perm.verify_log()
    assert res["ok"] is False and res["line"] == 5 and "改ざんの疑い" in res["message"]


# --- FR-06-160〜162・165 プリセット ---------------------------------------------------

def test_preset_select_screen(client, app):
    ps = client.get(f"{P}/presets").get_json()["presets"]
    names = {p["name"] for p in ps}
    assert "画面設計書" in names and any("29148" in n for n in names) and "AI 時代の完了条件 4 項目" in names
    assert len(next(p for p in ps if p["id"] == "dod-ai")["items"]) == 4
    iid = seed(app)
    body = client.post(f"{P}/inspections/{iid}/preset", json={"rev": 0, "preset_id": "screen-design"}).get_json()
    assert len(body["unchecked"]) == len(presets.get("screen-design")["items"])


def test_preset_copy_keeps_original(client):
    orig = presets.BUILTIN[0]["items"]
    c = client.post(f"{P}/presets/req-29148/copy", json={"name": "自社版"}).get_json()["preset"]
    items = [*c["items"], {"id": "extra", "text": "追加の観点"}]
    r = client.put(f"{P}/presets/{c['id']}", json={"name": "自社版", "items": items})
    assert r.status_code == 200 and len(r.get_json()["preset"]["items"]) == len(orig) + 1
    got = {p["id"]: p for p in client.get(f"{P}/presets").get_json()["presets"]}
    assert len(got["req-29148"]["items"]) == len(orig)
    assert client.put(f"{P}/presets/req-29148", json={"name": "x", "items": items}).status_code == 409


def test_unchecked_items():
    items = [{"id": f"i{n}", "text": "t"} for n in range(7)]
    assert len(process.unchecked(items, {f"i{n}": True for n in range(5)})) == 2


def test_preset_in_use_cannot_delete(client, app):
    c = client.post(f"{P}/presets/basic-design/copy", json={}).get_json()["preset"]
    seed(app, r4={"preset_id": c["id"]})
    r = client.delete(f"{P}/presets/{c['id']}")
    assert r.status_code == 409 and "使っているレビュー" in r.get_json()["error"]
    assert any(p["id"] == c["id"] for p in client.get(f"{P}/presets").get_json()["presets"])


# --- FR-06-034・035 個別レビューの独立・指摘の種別 ---------------------------------------------

def test_independent_individual_review(client, app, monkeypatch):
    ps = [{"name": "A", "role": "reviewer", "done": False}, {"name": "B", "role": "reviewer", "done": True}]
    iid = seed(app, participants=ps, phase="individual", r4={"independent": True},
               findings=[fnd("F1", raised_by="B"), fnd("F2", origin="auto", raised_by="自動候補")])
    as_user(monkeypatch, "A", auth=False)
    assert [f["id"] for f in client.get(f"{P}/inspections/{iid}").get_json()["findings"]] == ["F2"]
    with app.test_request_context():
        store.save({**store.load(iid), "participants": [{**ps[0], "done": True}, ps[1]]})
    assert {f["id"] for f in client.get(f"{P}/inspections/{iid}").get_json()["findings"]} == {"F1", "F2"}


def test_finding_kind_in_list_and_report(client, app):
    iid = seed(app, findings=[fnd("F1")])
    body = client.post(f"{P}/inspections/{iid}/findings/F1/kind", json={"rev": 0, "kind": "missing"}).get_json()
    assert body["findings"][0]["kind"] == "missing"
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"{P}/inspections/{iid}/record.xlsx").data))
    rows = list(wb["指摘一覧"].iter_rows(values_only=True))
    assert rows[0][1] == "種別" and rows[1][1] == "欠落"


# --- FR-06-043・044 回答の取り込み ------------------------------------------------------

def _xlsx(rows):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_answers_xlsx_same_as_csv(client, app):
    rows = [["No.", "場所", "質問", "回答"], ["F1", "1", "q", "はい"], ["F2", "2", "q", "いいえ"]]
    csv_text = "﻿" + "\r\n".join(",".join(r) for r in rows) + "\r\n"
    iid1 = seed(app, findings=[fnd("F1"), fnd("F2")])
    iid2 = seed(app, findings=[fnd("F1"), fnd("F2")])
    a = client.post(f"{P}/inspections/{iid1}/answers", data={"rev": "0", "file": (io.BytesIO(csv_text.encode()), "q.csv")})
    b = client.post(f"{P}/inspections/{iid2}/answers", data={"rev": "0", "file": (io.BytesIO(_xlsx(rows)), "q.xlsx")})
    assert a.status_code == b.status_code == 200
    assert a.get_json()["applied"] == b.get_json()["applied"] == 2


def test_answers_unknown_row_reported(client, app):
    rows = [["No.", "場所", "質問", "回答"], ["F1", "", "", "a"], ["F2", "", "", ""], ["F1", "", "", "b"], ["F9", "", "", "c"]]
    iid = seed(app, findings=[fnd("F1"), fnd("F2")])
    body = client.post(f"{P}/inspections/{iid}/answers", data={"rev": "0", "file": (io.BytesIO(_xlsx(rows)), "q.xlsx")}).get_json()
    assert "5 行目" in body["message"] and body["message"].startswith("反映できない行:")
    assert [s["row"] for s in body["skipped"]] == [3, 5]


# --- FR-06-054・074・076 承認の失効・記録・規格の扱い ----------------------------------------------

def test_approval_stale_after_new_version(client, app):
    iid = seed(app, document=DOC)
    r = client.post(f"{P}/inspections/{iid}/verdict",
                    json={"rev": 0, "verdict": "approved", "approver": "承認者X", "decided_at": "2026-10-15 10:00"})
    assert r.status_code == 200 and r.get_json()["approval"]["stale"] is False
    with app.test_request_context():
        store.save({**store.load(iid), "document": {**DOC, "version": 2}})
    ap = client.get(f"{P}/inspections/{iid}").get_json()["approval"]
    assert ap["stale"] is True and "承認後に文書が変わりました" in ap["message"]


def test_record_has_all_items(client, app):
    iid = seed(app, participants=PS, document={**DOC, "sha256": "ab" * 32}, r4={"review_type": "walkthrough"})
    client.post(f"{P}/inspections/{iid}/advance", json={"rev": 0})
    client.post(f"{P}/inspections/{iid}/verdict",
                json={"rev": 1, "verdict": "conditional", "approver": "X", "decided_at": "2026-10-15"})
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"{P}/inspections/{iid}/record.xlsx").data))
    labels = {r[0]: r[1] for r in wb["レビュー記録"].iter_rows(min_row=2, values_only=True)}
    for key in ("種別", "参加者と役割", "工程: 個別レビュー", "所要時間（工程の合計）", "対象版", "対象のハッシュ（SHA-256）", "判定"):
        assert key in labels
    assert labels["判定"] == "条件付き承認" and labels["対象のハッシュ（SHA-256）"] == "ab" * 32


def test_no_certification_claims(client, app):
    iid = seed(app)
    texts = [client.get(f"{P}/meta").get_data(as_text=True)]
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"{P}/inspections/{iid}/record.xlsx").data))
    texts += [str(c) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for c in r]
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    texts += [(root / p).read_text(encoding="utf-8") for p in ("templates/docverify/_process.html", "static/js/docverify-process.js")]
    joined = json.loads(texts[0]).__repr__() + "".join(texts[1:])
    assert "認証済み" not in joined and "規格に準拠" not in joined
    assert "ISO/IEC 20246" in joined and "2017" in joined


# --- NFR-06-21 楽観ロック・NFR-06-22 権限 ---------------------------------------------------

def test_optimistic_lock(client, app):
    iid = seed(app)
    a = client.put(f"{P}/inspections/{iid}/effort", json={"rev": 0, "prep": {"A": 10}, "meeting_minutes": 5})
    b = client.put(f"{P}/inspections/{iid}/effort", json={"rev": 0, "prep": {"B": 99}, "meeting_minutes": 1})
    assert a.status_code == 200 and b.status_code == 409
    assert b.get_json()["error"] == "他の人が先に更新しました。最新を読み込んでから操作してください。"
    assert process.effort(load(app, iid))["prep"] == {"A": 10}


def test_reviewer_forbidden(client, app, monkeypatch):
    iid = seed(app, participants=[{"name": "R", "role": "reviewer"}], document=DOC)
    as_user(monkeypatch, "R")
    reqs = [client.post(f"{P}/inspections/{iid}/verdict", json={"rev": 0, "verdict": "approved", "approver": "R", "decided_at": "d"}),
            client.put(f"{P}/reference", json={"basic": {"low": 0, "high": 1}}),
            client.delete(f"{P}/presets/c-0000000000")]
    for r in reqs:
        assert r.status_code == 403
        msg = r.get_json()["error"]
        assert "権限がありません" in msg and not LEAK.search(msg)
    assert not perm.can("delete_review", {"reviewer"}) and perm.can("delete_review", {"author"})
    assert perm.can("approve", {"reviewer"}, auth_on=False)
    assert load(app, iid)["approval"]["approver"] is None
