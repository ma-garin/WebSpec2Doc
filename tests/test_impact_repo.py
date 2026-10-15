"""System 09 影響度分析のリポジトリ層（DD-001/002/003/009/010）の単体テスト。"""
from __future__ import annotations

import io
import json
import sqlite3

import pytest

from web import config
from web.services.impact import db, migrate, report, roles
from web.services.impact.repo import (actuals, approvals, audit, baselines, crs, documents, impacts, links,
                                      systems)
from web.services.impact.repo.errors import ImpactError

SCREEN = "# 会員登録画面\n\nパスワードは 8 文字以上とする。\n\nメールアドレスは必須とする。\n".encode()
TESTS = "# テストケース\n\nSCR-0001 のパスワードが 8 文字以上なら登録できる。\n\nメール未入力でエラー。\n".encode()


@pytest.fixture()
def conn(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setenv("WEBSPEC2DOC_AUTH_MODE", "required")
    c = db.connect()
    yield c
    c.close()


@pytest.fixture()
def sysdocs(conn, tmp_path):
    s = systems.create(conn, "会員管理", "説明", "alice")
    scr = documents.add(conn, s["id"], "画面仕様書_会員登録.md", SCREEN, "alice", tmp_path / "w1")
    tc = documents.add(conn, s["id"], "テストケース_会員登録.md", TESTS, "alice", tmp_path / "w2")
    return s["id"], scr, tc


def _actions(conn, sid):
    return [a["action"] for a in audit.query(conn, sid)]


def _impact(conn, cr_id, item, risk=6, priority=2):
    iid = db.new_id()
    conn.execute("INSERT INTO impacts(id, cr_id, change_ids, item_id, item_code, document_name, text, depth, evidence,"
                 " impact_level, likelihood, risk, priority, tests, test_gap, method) "
                 "VALUES (?,?,?,?,?,?,?,1,?,3,2,?,?,?,0,'rule')",
                 (iid, cr_id, '["C1"]', item["id"], item["code"], item["document_name"], item["text"],
                  json.dumps([{"type": "link", "detail": "x", "score": 1}]), risk, priority, '["TC-0001"]'))
    conn.commit()
    return iid


def _analyzed(conn, sid, code):
    cr = crs.get(conn, sid, code)
    conn.execute("INSERT INTO jobs(id, cr_id, status, started_at, engine) VALUES (?,?,'done',?,?)",
                 (db.new_id(), cr["id"], db.now(), '{"mode": "rule"}'))
    crs.set_status(conn, cr, "analyzing", "bob")
    conn.commit()
    return cr


# ---- DD-001 ----
def test_db_schema_and_audit_append_only(conn):
    """Given 新しい DB / When 監査に書いて UPDATE・DELETE する / Then 追記のみで例外になる。"""
    assert db.db_path().name == "impact.db" and db.db_path().is_file()
    assert conn.execute("PRAGMA foreign_keys").fetchone()[0] == 1
    assert conn.execute("SELECT version FROM schema_version").fetchone()[0] == db.SCHEMA_VERSION
    audit.write(conn, "s", None, "a", "x", "t", {"a": 1}, None)
    conn.commit()
    with pytest.raises(sqlite3.DatabaseError, match="append-only"):
        conn.execute("UPDATE audit SET actor='z'")
    with pytest.raises(sqlite3.DatabaseError, match="append-only"):
        conn.execute("DELETE FROM audit")
    db.migrate(conn)
    assert conn.execute("SELECT COUNT(*) FROM schema_version").fetchone()[0] == 1


# ---- DD-003 ----
def test_roles_table_and_require(conn, monkeypatch):
    """Given 役割表 / When 各役割で可否を見る / Then 表どおり、メンバー外は 404、不可は 403。"""
    assert roles.can("admin", "edit_system") and not roles.can("analyst", "edit_system")
    assert roles.can("analyst", "override") and not roles.can("approver", "override")
    assert roles.can("approver", "approve") and roles.can("approver", "audit_view")
    assert not roles.can("viewer", "documents") and roles.can("viewer", "view") and not roles.can(None, "view")
    sid = systems.create(conn, "S", "", "alice")["id"]
    systems.set_member(conn, sid, "vic", "viewer", "alice")
    assert roles.require(conn, sid, "alice", "members") == "admin"
    with pytest.raises(ImpactError) as e:
        roles.require(conn, sid, "vic", "documents")
    assert e.value.status == 403
    with pytest.raises(ImpactError) as e:
        roles.require(conn, sid, "mallory", "view")
    assert e.value.status == 404
    monkeypatch.setenv("WEBSPEC2DOC_AUTH_MODE", "off")
    assert roles.effective_actor("") == "local" and roles.require(conn, sid, "local", "members") == "admin"
    with pytest.raises(ImpactError) as e:  # 認証オフでも名前のある利用者はメンバー表で判定する
        roles.require(conn, sid, "vic", "documents")
    assert e.value.status == 403
    assert [s["role"] for s in systems.list_for(conn, "local")] == ["admin"]


# ---- systems ----
def test_systems_crud_members_and_audit(conn):
    """Given 作成 / When 設定変更・メンバー変更・削除 / Then 値が変わり監査に残る。不正は 400/404/409。"""
    s = systems.create(conn, "S", "d", "alice")
    assert s["members"] == [{"actor": "alice", "role": "admin"}] and s["counts"]["documents"] == 0
    assert [x["id"] for x in systems.list_for(conn, "alice")] == [s["id"]] and systems.list_for(conn, "bob") == []
    u = systems.update(conn, s["id"], {"policy": "local", "p1_min": 9, "p2_min": 3}, "alice")
    assert (u["policy"], u["p1_min"], u["p2_min"]) == ("local", 9, 3)
    for bad in ({"p1_min": 4, "p2_min": 4}, {"policy": "x"}, {"depth_default": 5}, {"name": ""}):
        with pytest.raises(ImpactError) as e:
            systems.update(conn, s["id"], bad, "alice")
        assert e.value.status == 400
    systems.set_member(conn, s["id"], "bob", "approver", "alice")
    assert systems.role_of(conn, s["id"], "bob") == "approver"
    with pytest.raises(ImpactError) as e:
        systems.remove_member(conn, s["id"], "alice", "alice")
    assert e.value.status == 409
    systems.remove_member(conn, s["id"], "bob", "alice")
    with pytest.raises(ImpactError) as e:
        systems.remove_member(conn, s["id"], "bob", "alice")
    assert e.value.status == 404
    systems.delete(conn, s["id"], "alice")
    with pytest.raises(ImpactError) as e:
        systems.get(conn, s["id"])
    assert e.value.status == 404
    assert _actions(conn, s["id"]) == ["system.delete", "member.remove", "member.set", "system.update",
                                      "system.create"]


# ---- documents ----
def test_documents_ingest_replace_kind_remove(conn, sysdocs, tmp_path):
    """Given 2 文書 / When 置換・種類変更・削除 / Then 項目コードが採番され、同じ本文の項目は ID を保つ。"""
    sid, scr, tc = sysdocs
    assert scr["kind"] == "screen" and tc["kind"] == "testcase" and documents.guess_kind("要件定義.docx") == "requirement"
    its = documents.items(conn, sid, doc_id=scr["id"])
    assert its[0]["code"] == "SCR-0001" and its[0]["document_name"] == scr["name"] and its[0]["kind"] == "screen"
    keep = {i["text"]: i["id"] for i in its}
    new = SCREEN.replace("8 文字".encode(), "10 文字".encode())
    documents.add(conn, sid, scr["name"], new, "alice", tmp_path / "w3")
    after = {i["text"]: i for i in documents.items(conn, sid, doc_id=scr["id"])}
    same = [t for t in keep if t in after]
    assert same and all(after[t]["id"] == keep[t] for t in same)
    assert any(i["code"].startswith("SCR-") and "10" in t for t, i in after.items())
    assert documents.items(conn, sid, q="メール")
    documents.set_kind(conn, sid, scr["id"], "design", "alice")
    assert all(i["code"].startswith("DSN-") for i in documents.items(conn, sid, doc_id=scr["id"]))
    with pytest.raises(ImpactError) as e:
        documents.set_kind(conn, sid, scr["id"], "x", "alice")
    assert e.value.status == 400
    with pytest.raises(ImpactError) as e:
        documents.add(conn, sid, "a.exe", b"x", "alice", tmp_path / "w4")
    assert e.value.status == 400
    documents.remove(conn, sid, scr["id"], "alice")
    with pytest.raises(ImpactError) as e:
        documents.remove(conn, sid, scr["id"], "alice")
    assert e.value.status == 404
    assert [d["name"] for d in documents.list(conn, sid)] == [tc["name"]]
    assert {"document.add", "document.replace", "document.kind", "document.remove"} <= set(_actions(conn, sid))


# ---- baselines ----
def test_baseline_frozen_and_diff(conn, sysdocs, tmp_path):
    """Given ベースライン v1 / When 文書を差し替えて v2 / Then v1 は変わらず、差分に変更・追加・削除が出る。"""
    sid, scr, _ = sysdocs
    b1 = baselines.create(conn, sid, "v1", "", "alice")
    before = baselines.items(conn, b1["id"])
    new = SCREEN.replace("8 文字".encode(), "10 文字".encode()).replace(
        "メールアドレスは必須とする。".encode(), "電話番号を任意で入力できる。".encode())
    documents.add(conn, sid, scr["name"], new + "生年月日を入力する欄を新しく設ける。\n".encode(), "alice", tmp_path / "w5")
    b2 = baselines.create(conn, sid, "v2", "", "alice")
    assert baselines.items(conn, b1["id"]) == before
    d = baselines.diff(conn, b1["id"], b2["id"])
    assert any("10" in m["new"]["text"] and "8" in m["old"]["text"] for m in d["modified"])
    assert d["added"] and d["removed"]
    assert [b["name"] for b in baselines.list(conn, sid)] == ["v1", "v2"]
    with pytest.raises(ImpactError) as e:
        baselines.create(conn, sid, "v1", "", "alice")
    assert e.value.status == 409
    with pytest.raises(ImpactError) as e:
        baselines.diff(conn, b1["id"], "nope")
    assert e.value.status == 404


# ---- links ----
class _Emb:
    available, mode = True, "ngram"

    def vectors(self, texts):
        return [set(t) for t in texts]

    def similar(self, q, corpus, top_k):
        scores = [(i, len(q & v) / max(1, len(q | v))) for i, v in enumerate(corpus)]
        return sorted(scores, key=lambda x: -x[1])[:top_k]


def test_links_add_remove_suggest_decide(conn, sysdocs):
    """Given 2 文書の項目 / When つなぐ・候補を作り採用する / Then 正規化・重複 409・同一 400・監査あり。"""
    sid, scr, tc = sysdocs
    a = documents.items(conn, sid, doc_id=scr["id"])[1]
    b = documents.items(conn, sid, doc_id=tc["id"])[2]
    lk = links.add(conn, sid, b["id"], a["id"], "alice")
    assert lk["a_item"] < lk["b_item"] and {lk["a_code"], lk["b_code"]} == {a["code"], b["code"]}
    for x, y, st in ((a["id"], b["id"], 409), (a["id"], a["id"], 400), (a["id"], "nope", 404)):
        with pytest.raises(ImpactError) as e:
            links.add(conn, sid, x, y, "alice")
        assert e.value.status == st
    assert [n["id"] for n in links.neighbors(conn, sid, a["id"])] == [b["id"]]
    assert len(links.list(conn, sid)) == 1
    made = links.suggest(conn, sid, _Emb())
    assert made >= 1 and links.suggest(conn, sid, _Emb()) == 0
    sug = links.suggestions(conn, sid)
    assert any("SCR-0001" in s["reason"] for s in sug)
    links.decide(conn, sid, sug[0]["id"], True, "alice")
    with pytest.raises(ImpactError) as e:
        links.decide(conn, sid, sug[0]["id"], False, "alice")
    assert e.value.status == 409
    assert any(r["origin"] == "suggested" for r in links.list(conn, sid))
    links.remove(conn, sid, lk["id"], "alice")
    assert {"link.add", "link.remove", "suggestion.accept"} <= set(_actions(conn, sid))


# ---- crs / impacts ----
def test_cr_lifecycle_and_review_conditions(conn, sysdocs):
    """Given 起票 / When 遷移する / Then DD-007 の表どおり。分析前・要確認ありは review へ 409。"""
    sid, scr, _ = sysdocs
    cr = crs.create(conn, sid, "桁数変更", "法改正", "8 文字を 10 文字に", None, None, None, "bob")
    assert cr["code"] == "CR-001" and cr["status"] == "draft" and cr["depth"] == 2 and cr["changes"] == []
    assert crs.create(conn, sid, "二件目", "", "", None, None, 1, "bob")["code"] == "CR-002"
    for args, st in ((("", "", "", None, None, None), 400), (("t", "", "", "nobl", None, None), 404),
                     (("t", "", "", None, None, 9), 400)):
        with pytest.raises(ImpactError) as e:
            crs.create(conn, sid, *args, "bob")
        assert e.value.status == st
    with pytest.raises(ImpactError) as e:
        crs.transition(conn, sid, "CR-001", "verified", "bob")
    assert e.value.status == 409
    with pytest.raises(ImpactError) as e:
        crs.transition(conn, sid, "CR-001", "rejected", "bob")
    assert e.value.status == 400
    crs.transition(conn, sid, "CR-001", "analyzing", "bob")
    with pytest.raises(ImpactError) as e:
        crs.transition(conn, sid, "CR-001", "review", "bob")
    assert e.value.status == 409 and "分析" in e.value.message
    cr = _analyzed(conn, sid, "CR-001")
    iid = _impact(conn, cr["id"], documents.items(conn, sid, doc_id=scr["id"])[1])
    with pytest.raises(ImpactError) as e:
        crs.transition(conn, sid, "CR-001", "review", "bob")
    assert e.value.status == 409 and "要確認" in e.value.message
    impacts.resolve(conn, iid, "required", "桁数の検証を直す", "bob")
    assert crs.transition(conn, sid, "CR-001", "review", "bob")["status"] == "review"
    with pytest.raises(ImpactError) as e:
        crs.update(conn, sid, "CR-001", {"title": "x"}, "bob")
    assert e.value.status == 409
    assert crs.update(conn, sid, "CR-002", {"title": "改題"}, "bob")["title"] == "改題"
    assert crs.transition(conn, sid, "CR-002", "rejected", "bob", "不要")["status"] == "rejected"
    assert [c["code"] for c in crs.list(conn, sid, status="review")] == ["CR-001"]
    with pytest.raises(ImpactError) as e:
        crs.get(conn, sid, "CR-999")
    assert e.value.status == 404
    assert _actions(conn, sid).count("cr.transition") >= 4


def test_impacts_resolve_override_summary(conn, sysdocs):
    """Given 影響 1 件 / When 解消・修正 / Then 理由必須、リスクと優先度を閾値で再計算、集計と監査が合う。"""
    sid, scr, _ = sysdocs
    cr = _analyzed(conn, sid, crs.create(conn, sid, "t", "", "", None, None, None, "bob")["code"])
    iid = _impact(conn, cr["id"], documents.items(conn, sid, doc_id=scr["id"])[1])
    got = impacts.get(conn, iid)
    assert got["evidence"][0]["type"] == "link" and got["change_ids"] == ["C1"] and got["tests"] == ["TC-0001"]
    for args in (("required", ""), ("open", "x")):
        with pytest.raises(ImpactError) as e:
            impacts.resolve(conn, iid, *args, "bob")
        assert e.value.status == 400
    for args in ((4, 3, ""), (5, 1, "x"), (1, 0, "x")):
        with pytest.raises(ImpactError) as e:
            impacts.override(conn, iid, *args, "bob")
        assert e.value.status == 400
    o = impacts.override(conn, iid, 4, 3, "決済に関わる", "bob")
    assert (o["risk"], o["priority"], o["effective_impact"]) == (12, 1, 4)
    assert impacts.override(conn, iid, 1, 3, "軽微", "bob")["priority"] == 3
    assert impacts.list(conn, cr["id"], priority=3)[0]["id"] == iid and impacts.list(conn, cr["id"], state="required") == []
    assert impacts.summary(conn, cr["id"]) == {"total": 1, "p1": 0, "p2": 0, "p3": 1, "open": 1, "test_gap": 0,
                                               "retest": 1}
    impacts.resolve(conn, iid, "not_required", "表示のみ", "bob")
    assert impacts.summary(conn, cr["id"])["open"] == 0
    with pytest.raises(ImpactError) as e:
        impacts.get(conn, "nope")
    assert e.value.status == 404
    assert {"impact.override", "impact.resolve"} <= set(_actions(conn, sid))


def _to_review(conn, sid, scr, resolver="bob"):
    code = crs.create(conn, sid, "t", "", "", None, None, None, "bob")["code"]
    cr = _analyzed(conn, sid, code)
    iid = _impact(conn, cr["id"], documents.items(conn, sid, doc_id=scr["id"])[1])
    impacts.resolve(conn, iid, "required", "直す", resolver)
    crs.transition(conn, sid, code, "review", "bob")
    return code


def test_approvals_self_forbidden_and_decide(conn, sysdocs):
    """Given レビュー待ち / When 承認・差し戻し / Then 承認者のみ・起票者/解消者は 403・差し戻しは理由必須。"""
    sid, scr, _ = sysdocs
    for actor, role in (("bob", "analyst"), ("carol", "approver"), ("dave", "approver"), ("vic", "viewer")):
        systems.set_member(conn, sid, actor, role, "alice")
    code = _to_review(conn, sid, scr, resolver="dave")
    for actor, st in (("vic", 403), ("bob", 403), ("dave", 403), ("mallory", 404)):
        with pytest.raises(ImpactError) as e:
            approvals.decide(conn, sid, code, "approved", "", actor)
        assert e.value.status == st
    with pytest.raises(ImpactError) as e:
        approvals.decide(conn, sid, code, "returned", "", "carol")
    assert e.value.status == 400
    r = approvals.decide(conn, sid, code, "returned", "根拠不足", "carol")
    assert r["status"] == "returned" and crs.get(conn, sid, code)["status"] == "returned"
    crs.transition(conn, sid, code, "review", "bob")
    assert crs.transition(conn, sid, code, "approved", "carol")["status"] == "approved"
    with pytest.raises(ImpactError) as e:
        approvals.decide(conn, sid, code, "approved", "", "carol")
    assert e.value.status == 409
    assert [a["decision"] for a in crs.get(conn, sid, code)["approvals"]] == ["returned", "approved"]
    assert {"cr.approve", "cr.return"} <= set(_actions(conn, sid))


def test_actuals_review_trend(conn, sysdocs):
    """Given 承認済み CR / When 実際の影響を記録 / Then 一致・見逃し・再現率が出る。重複 409・未知 404・承認前 409。"""
    sid, scr, tc = sysdocs
    systems.set_member(conn, sid, "carol", "approver", "alice")
    code = _to_review(conn, sid, scr)
    with pytest.raises(ImpactError) as e:
        actuals.record(conn, sid, code, "SCR-0002", "", "bob")
    assert e.value.status == 409
    approvals.decide(conn, sid, code, "approved", "", "carol")
    hit_code = documents.items(conn, sid, doc_id=scr["id"])[1]["code"]
    miss_code = documents.items(conn, sid, doc_id=tc["id"])[1]["code"]
    actuals.record(conn, sid, code, hit_code, "直した", "bob")
    actuals.record(conn, sid, code, miss_code, "", "bob")
    for item, st in ((hit_code, 409), ("ZZZ-9", 404)):
        with pytest.raises(ImpactError) as e:
            actuals.record(conn, sid, code, item, "", "bob")
        assert e.value.status == st
    r = actuals.review(conn, sid, code)
    assert (r["hit"], r["missed"], r["false_alarm"], r["recall"]) == (1, [miss_code], 0, 0.5)
    t = actuals.trend(conn, sid)
    assert t[0]["code"] == code and t[0]["adoption_rate"] == 1.0
    assert "actual.record" in _actions(conn, sid)


def test_audit_query_filters(conn, sysdocs):
    """Given 複数の操作 / When 絞り込む / Then 操作者・CR で絞れ、cr_code が付く。"""
    sid, _, _ = sysdocs
    crs.create(conn, sid, "t", "", "", None, None, None, "bob")
    rows = audit.query(conn, sid, cr_code="CR-001")
    assert [r["action"] for r in rows] == ["cr.create"] and rows[0]["cr_code"] == "CR-001"
    assert all(r["actor"] == "alice" for r in audit.query(conn, sid, actor="alice"))
    assert audit.query(conn, sid, since="2999-01-01") == []
    assert audit.query(conn, sid, actor="alice")[0]["cr_code"] == ""


# ---- DD-009 ----
def test_report_has_five_sheets(conn, sysdocs):
    """Given 承認済みの CR / When レポートを作る / Then 5 シートに影響・テスト・承認・監査が入る。"""
    from openpyxl import load_workbook
    sid, scr, _ = sysdocs
    systems.set_member(conn, sid, "carol", "approver", "alice")
    code = _to_review(conn, sid, scr)
    approvals.decide(conn, sid, code, "approved", "", "carol")
    wb = load_workbook(io.BytesIO(report.build_xlsx(conn, sid, code)))
    assert wb.sheetnames == ["概要", "影響", "テスト", "承認", "監査"]
    assert wb["影響"].max_row == 2 and "×" in wb["影響"]["C2"].value
    assert wb["テスト"]["C2"].value == "TC-0001" and wb["承認"]["A2"].value == "承認"
    assert wb["監査"].max_row > 3 and wb["概要"]["B2"].value == code


# ---- DD-010 ----
def test_migrate_idempotent(conn):
    """Given 旧版の分析 1 件 / When 移行を 2 回 / Then 1 回目だけ取り込み、判断が解消の記録に移る。"""
    from web.services.impact import store
    job = store.new("旧分析")
    job = store.add_doc(job, "画面仕様書_会員登録.md", SCREEN)
    job = {**job, "order": "8 文字を 10 文字に",
           "changes": [{"id": "C1", "kind": "変更", "location": "x", "before": "8", "after": "10"}],
           "impacts": [{"id": "I1", "level": 3, "priority": 2, "doc": "画面仕様書_会員登録.md", "location": "",
                        "changes": ["C1"], "why": "同じ語", "quote": "パスワード", "test": "確かめる", "score": 1,
                        "method": "rule", "review": {"status": "accepted", "note": "直す"}},
                       {"id": "I2", "level": 1, "priority": 3, "doc": "", "location": "", "changes": ["C1"],
                        "why": "", "quote": "", "test": "", "score": 0, "method": "rule",
                        "review": {"status": "rejected", "note": ""}}]}
    store.save(job)
    first = migrate.run(conn, "alice")
    assert first == {"systems": 1, "crs": 1, "impacts": 2}
    assert migrate.run(conn, "alice") == {"systems": 0, "crs": 0, "impacts": 0}
    sid = conn.execute("SELECT id FROM systems WHERE name='移行データ'").fetchone()["id"]
    cr = crs.get(conn, sid, "CR-001")
    assert cr["status"] == "analyzing" and cr["changes"][0]["kind"] == "modify"
    rows = {r["state"]: r for r in impacts.list(conn, cr["id"])}
    assert rows["required"]["resolve_reason"] == "直す" and "not_required" in rows
    assert documents.list(conn, sid)[0]["item_count"] >= 1
    assert "migrate.job" in _actions(conn, sid)
