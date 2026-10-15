"""書き出し。報告書（Excel 6 シート・Markdown）とテストコード一式（zip）。

事実（実行で得た値）と判断（機械判定・人の判定）は別の列に分ける。
実行していないもの・判定していないものは「未実行」「未判定」のまま載せる（空欄にしない）。
"""
from __future__ import annotations

import csv
import io
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from web.services.backendtest import cases as cases_mod
from web.services.backendtest import evidence, store
from web.services.backendtest.store import StoreError

LEVEL_LABELS = {"integration_internal": "結合テスト（内部）", "integration_external": "結合テスト（外部）",
                "system": "システムテスト", "acceptance": "受け入れテスト"}
STAGE_NAMES = {"1": "目的・範囲", "2": "対象の読み取り", "3": "観点表", "4": "ケース表",
               "5": "テストデータ仕様", "6": "テストコード", "7": "実行", "8": "判定・報告"}
STATUS_LABELS = {"draft": "未作成", "generated": "承認待ち", "approved": "承認済み", "rejected": "差し戻し"}
STATE_LABELS = {"pass": "合格", "fail": "不合格", "undetermined": "未判定", "not_run": "未実行", "error": "実行エラー"}
METHOD_LABELS = {"db_diff": "DB 前後比較", "alt_sql": "別経路計算", "golden": "ゴールデン", "schema": "スキーマ検査"}
ORACLE_STATE = {"unconfirmed": "未確定", "auto": "機械確定", "human": "人が確定"}
DECISION_LABELS = {"emit": "採用", "excluded": "除外", "adopted": "追加採用", "edited": "修正", "discarded": "破棄"}
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _s(value) -> str:
    """セルに入れる文字列。None は空、リストは改行区切り、制御文字は除く。"""
    if value is None:
        return ""
    if isinstance(value, dict):
        value = "、".join(f"{k}={v}" for k, v in value.items())
    elif isinstance(value, (list, tuple)):
        value = "\n".join(_s(v) for v in value)
    return _CTRL.sub("", str(value))


def _join(value) -> str:
    if isinstance(value, (list, tuple)):
        return "、".join(_s(v) for v in value)
    return _s(value)


def _overview(job: dict) -> tuple[tuple, list[list]]:
    stages = job.get("stages") or {}
    rows = [
        ["案件名", job.get("name")], ["対象の種類", "バッチジョブ" if job.get("target_kind") == "batch" else job.get("target_kind")],
        ["テストレベル", LEVEL_LABELS.get(job.get("level"), job.get("level"))],
        ["目的", job.get("purpose") or "未記入"], ["範囲", job.get("scope") or "未記入"],
        ["作成日時", job.get("created_at")], ["更新日時", job.get("updated_at")],
        ["観点の数", len(job.get("viewpoints") or [])], ["ケースの数", len(job.get("cases") or [])],
        ["実行結果の数", len(job.get("results") or [])], ["不具合の数", len(job.get("defects") or [])],
    ]
    for no in store.STAGE_NOS:
        st = stages.get(no) or {}
        who = f"（{st['approved_by']}・{st.get('approved_at') or ''}）" if st.get("approved_by") else ""
        rows.append([f"段階 {no} {STAGE_NAMES[no]}", f"{STATUS_LABELS.get(st.get('status'), '未作成')}{who}"])
    rows.append(["未検証・未判定の件数", _unverified_note(job)])
    return ("項目", "内容"), rows


def _result_index(job: dict) -> dict:
    index: dict = {}
    for r in job.get("results") or []:
        index.setdefault(r.get("case"), []).append(r)
    return index


def _unverified_note(job: dict) -> str:
    cases = job.get("cases") or []
    index = _result_index(job)
    not_run = sum(1 for c in cases if c.get("id") not in index)
    undetermined = sum(1 for rs in index.values() for r in rs if r.get("state") == "undetermined")
    unconfirmed = sum(1 for c in cases if (c.get("oracle") or {}).get("state") in (None, "unconfirmed"))
    return f"未実行のケース {not_run} 件／未判定の結果 {undetermined} 件／期待値が未確定のケース {unconfirmed} 件"


def _viewpoints(job: dict) -> tuple[tuple, list[list]]:
    head = ("ID", "分類", "観点", "技法", "扱い", "理由", "根拠", "未検証", "リリース")
    rows = [[v.get("id"), v.get("cat"), v.get("name"), v.get("technique"),
             DECISION_LABELS.get(v.get("decision"), v.get("decision")), v.get("reason"), v.get("basis"),
             "未検証" if v.get("unverified") else "", v.get("release")] for v in job.get("viewpoints") or []]
    return head, rows


def _cases(job: dict) -> tuple[tuple, list[list]]:
    head = ("ID", "観点", "技法", "レベル", "要件 ID", "前提", "手順", "期待結果", "期待値方式", "期待値の状態", "対象表", "根拠")
    rows = []
    for c in job.get("cases") or []:
        o = c.get("oracle") or {}
        rows.append([c.get("id"), c.get("viewpoint"), c.get("technique"), LEVEL_LABELS.get(c.get("level"), c.get("level")),
                     _join(c.get("req_ids")), c.get("pre"), c.get("steps"), c.get("expected"),
                     METHOD_LABELS.get(o.get("method"), o.get("method") or "未設定"),
                     ORACLE_STATE.get(o.get("state"), "未確定"), _join(o.get("tables")), c.get("basis")])
    return head, rows


def _data_specs(job: dict) -> tuple[tuple, list[list]]:
    head = ("ID", "ケース", "テーブル", "条件", "行数", "文字コード", "伏せる列", "根拠")
    rows = [[d.get("id"), d.get("case"), d.get("table"), d.get("conditions"), d.get("rows"),
             d.get("encoding"), _join(d.get("mask_cols")), d.get("basis")] for d in job.get("data_specs") or []]
    return head, rows


def _fact_cell(r: dict) -> tuple[str, str]:
    fact = r.get("fact") or {}
    return ("" if fact.get("rc") is None else str(fact.get("rc"))), _s(fact.get("diff"))


def _results(job: dict) -> tuple[tuple, list[list]]:
    head = ("ケース", "実行 ID", "【事実】戻りコード", "【事実】差分の証跡", "【判断】機械判定", "【判断】適用した規則",
            "【判断】人の判定", "判定理由", "判定者", "判定日時", "状態", "原因")
    index, rows = _result_index(job), []
    for c in job.get("cases") or []:
        if c.get("id") not in index:
            rows.append([c.get("id"), "", "", "", "", "", "", "未実行（実行結果がありません）", "", "", "未実行", ""])
    for rs in index.values():
        for r in rs:
            rc, diff = _fact_cell(r)
            auto, human = r.get("auto") or {}, r.get("human") or {}
            rows.append([r.get("case"), r.get("run"), rc, diff, auto.get("verdict") or "未判定", auto.get("rule"),
                         human.get("verdict") or "未判定", human.get("reason"), human.get("by"), human.get("at"),
                         STATE_LABELS.get(r.get("state"), r.get("state") or "未判定"), r.get("cause")])
    return head, rows


def _defects(job: dict) -> tuple[tuple, list[list]]:
    head = ("ID", "実行 ID", "ケース", "件名", "重大度（ISTQB）", "優先度", "再現の入力", "再現の手順", "差分")
    rows = []
    for d in job.get("defects") or []:
        rp = d.get("repro") or {}
        rows.append([d.get("id"), d.get("run"), d.get("case"), d.get("title"), d.get("severity"),
                     d.get("priority"), rp.get("input"), rp.get("steps"), rp.get("diff")])
    return head, rows


# (シート名, 組み立て関数, 列幅, 空のときの文言)
_SHEETS = (
    ("概要", _overview, (22, 70), ""),
    ("観点表", _viewpoints, (10, 10, 36, 24, 10, 30, 28, 8, 10), "観点がまだありません。"),
    ("ケース表", _cases, (8, 10, 18, 14, 14, 30, 36, 36, 14, 12, 18, 24), "ケースがまだありません。"),
    ("テストデータ仕様", _data_specs, (8, 8, 18, 36, 8, 12, 18, 24), "テストデータ仕様がまだありません。"),
    ("実行結果", _results, (8, 16, 12, 28, 14, 18, 14, 30, 12, 20, 10, 24), "実行結果がまだありません。"),
    ("不具合", _defects, (10, 16, 8, 30, 14, 10, 30, 36, 30), "不具合はありません。"),
)


def _put(ws, row: int, values) -> None:
    for col, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=_s(v))
        if cell.value.startswith(("=", "+", "-", "@")):
            cell.data_type = "s"  # 数式として解釈させない
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _write_sheet(ws, head, rows, widths, empty: str) -> None:
    _put(ws, 1, head)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="F1F5F9")
    for i, row in enumerate(rows, 2):
        _put(ws, i, row)
    if not rows and empty:
        _put(ws, 2, [empty])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def report_xlsx(job: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, build, widths, empty in _SHEETS:
        head, rows = build(job)
        _write_sheet(wb.create_sheet(title), head, rows, widths, empty)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _md(value) -> str:
    return _s(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def report_md(job: dict) -> bytes:
    out = [f"# テスト報告書：{_md(job.get('name'))}", ""]
    for title, build, _widths, empty in _SHEETS:
        head, rows = build(job)
        out += [f"## {title}", ""]
        if not rows:
            out += [empty or "（なし）", ""]
            continue
        out += ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
        out += ["| " + " | ".join(_md(v) for v in row) + " |" for row in rows]
        out.append("")
    return "\n".join(out).encode("utf-8")


def code_zip(job: dict) -> bytes:
    """生成コード（code/ 配下）を code/ 付きで固める。まだ無ければ理由つきで失敗。"""
    root = store.job_dir(job["id"]) / "code"
    paths = sorted(p for p in root.rglob("*") if p.is_file()) if root.is_dir() else []
    if not paths:
        raise StoreError("テストコードがまだありません。段階 6 でテストコードを作ってから書き出してください。")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            if "__pycache__" not in p.parts:
                z.write(p, "code/" + p.relative_to(root).as_posix())
    return buf.getvalue()


def trace_csv(job: dict) -> bytes:
    """追跡表: 観点 → ケース → 根拠の事実 ID → 期待値方式・最新の結果。どのケースにも無い業務規則も行に出す。"""
    vps = {v.get("id"): v for v in job.get("viewpoints") or []}
    latest = _result_index(job)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["観点ID", "観点", "ケースID", "根拠の事実ID", "期待値方式", "期待値の状態", "最新の結果"])
    for c in job.get("cases") or []:
        o, r = c.get("oracle") or {}, (latest.get(c.get("id")) or [{}])[-1]
        state = (r.get("judgement") or {}).get("state") if r else "not_run"
        w.writerow([c.get("viewpoint", ""), (vps.get(c.get("viewpoint")) or {}).get("name", ""), c.get("id", ""),
                    " ".join(c.get("req_ids") or []), METHOD_LABELS.get(o.get("method"), _s(o.get("method"))),
                    ORACLE_STATE.get(o.get("state"), _s(o.get("state"))), STATE_LABELS.get(state, _s(state))])
    for fid in cases_mod.trace(job.get("cases"), job.get("facts"))["uncovered_reqs"]:
        w.writerow(["", "", "", fid, "", "", "どのケースにも結び付いていません"])
    return buf.getvalue().encode("utf-8-sig")


def evidence_zip(job: dict, run_id: str) -> bytes:
    """1 回の実行の証跡フォルダを evidence/<run_id>/ 付きで固める。"""
    root = evidence.run_dir(job["id"], run_id)
    paths = sorted(p for p in root.rglob("*") if p.is_file()) if root.is_dir() else []
    if not paths:
        raise KeyError(run_id)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            if "__pycache__" not in p.parts and not p.name.startswith(".tmp-"):
                z.write(p, f"evidence/{run_id}/" + p.relative_to(root).as_posix())
    return buf.getvalue()
